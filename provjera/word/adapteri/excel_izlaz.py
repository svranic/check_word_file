"""Excel zbirna tablica (priča 4.3, FR27, AD-12): samo zapisivanje i spremanje.

Sadržaj računa ``jezgra.zbirna``. Svaki tekst zapisuje se kao tekst (nikad formula), a
postojeća datoteka se nikad ne prepisuje.
"""

import errno
import io
import logging
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from provjera.word.jezgra.zbirna import ZbirnaTablica

log = logging.getLogger(__name__)

PREFIKS = "rezultati_"
FORMAT_OZNAKE = "%Y%m%d-%H%M%S"
NAJVISE_POKUSAJA = 100
REDAK_STUPACA = 3  # 1: zaglavlje vježbe, 2: prazno, 3: naslovi stupaca
NAJVECA_SIRINA = 60
NAJVECA_DULJINA_CELIJE = 32_767  # granica Excela; dulji tekst Excel prijavljuje kao oštećenje
ZNAKOVI_FORMULE = ("=", "+", "-", "@", "\t", "\r")


def _upisi(list_, redak: int, stupac: int, vrijednost: str | int, podebljano: bool = False) -> None:
    celija = list_.cell(row=redak, column=stupac)
    if isinstance(vrijednost, str):
        # Kontrolni znakovi (iz naziva datoteke ili isječka) nisu dopušteni u .xlsx.
        vrijednost = ILLEGAL_CHARACTERS_RE.sub("", vrijednost)[:NAJVECA_DULJINA_CELIJE]
    celija.value = vrijednost
    if isinstance(vrijednost, str):
        celija.data_type = "s"  # „=1+1“, „@SUM(A1)“ … ostaju tekst, nikad formula
        if vrijednost.startswith(ZNAKOVI_FORMULE):
            celija.quotePrefix = True  # ostaje tekst i nakon uređivanja ćelije u Excelu
    if podebljano:
        celija.font = Font(bold=True)


def izradi(tablica: ZbirnaTablica) -> bytes:
    knjiga = Workbook()
    knjiga.remove(knjiga.active)
    for opis in tablica.listovi:
        list_ = knjiga.create_sheet(opis.naziv)
        for i, vrijednost in enumerate(opis.zaglavlje, start=1):
            _upisi(list_, 1, i, vrijednost, podebljano=i % 2 == 1)
        for i, naslov in enumerate(opis.stupci, start=1):
            _upisi(list_, REDAK_STUPACA, i, naslov, podebljano=True)
        for r, redak in enumerate(opis.retci, start=REDAK_STUPACA + 1):
            for i, vrijednost in enumerate(redak, start=1):
                _upisi(list_, r, i, vrijednost)
        list_.freeze_panes = list_.cell(row=REDAK_STUPACA + 1, column=1)
        zadnji = get_column_letter(max(len(opis.stupci), 1))
        list_.auto_filter.ref = f"A{REDAK_STUPACA}:{zadnji}{REDAK_STUPACA + len(opis.retci)}"
        for i in range(1, max(len(opis.stupci), len(opis.zaglavlje)) + 1):
            duljine = [len(str(redak[i - 1])) for redak in opis.retci if i <= len(redak)]
            duljine += [len(str(v[i - 1])) for v in (opis.stupci, opis.zaglavlje) if i <= len(v)]
            list_.column_dimensions[get_column_letter(i)].width = min(max(duljine) + 2, NAJVECA_SIRINA)
    izlaz = io.BytesIO()
    knjiga.save(izlaz)
    return izlaz.getvalue()


def _kandidati(mapa_vjezbe: Path, id_vjezbe: str, vrijeme: datetime):
    yield mapa_vjezbe / f"{PREFIKS}{id_vjezbe}.xlsx"
    oznaka = vrijeme.strftime(FORMAT_OZNAKE)
    yield mapa_vjezbe / f"{PREFIKS}{id_vjezbe}_{oznaka}.xlsx"
    for broj in range(2, NAJVISE_POKUSAJA):
        yield mapa_vjezbe / f"{PREFIKS}{id_vjezbe}_{oznaka}_{broj}.xlsx"


def spremi(sadrzaj: bytes, mapa_vjezbe: Path, id_vjezbe: str, vrijeme: datetime) -> Path:
    """Ekskluzivno stvara prvu slobodnu datoteku; postojeća ili zaključana se preskače."""
    mapa_vjezbe.mkdir(parents=True, exist_ok=True)
    zakljucano: PermissionError | None = None
    for putanja in _kandidati(mapa_vjezbe, id_vjezbe, vrijeme):
        try:
            datoteka = open(putanja, "xb")  # noqa: SIM115 — zatvara se odmah ispod
        except FileExistsError:
            continue
        except PermissionError as greska:
            zakljucano = greska
            continue
        try:
            with datoteka:
                datoteka.write(sadrzaj)
                datoteka.flush()
                os.fsync(datoteka.fileno())
        except BaseException:
            putanja.unlink(missing_ok=True)  # nedovršena datoteka koju je stvorio ovaj poziv
            raise
        log.info("zbirna tablica spremljena")
        return putanja
    if zakljucano is not None:
        raise zakljucano  # mapa nije zapisiva: pravi razlog, ne „nema imena“
    raise FileExistsError(errno.EEXIST, "nema slobodnog imena za zbirnu tablicu")

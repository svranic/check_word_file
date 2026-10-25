"""Izrada Excel predloška vježbe (openpyxl): obrazac za svaki zadatak profila.

Ulazne ćelije su plave, ćelije za formule žute, a stupci od M nadalje su slobodno
područje za pomoćne izračune. List je zaštićen lozinkom tako da student ne može
pomaknuti ćelije obrasca ni promijeniti zaglavlje (retci 1–4: ime, prezime i JMBAG); žute
ćelije i slobodno područje ostaju otključani. Zaštićena je i struktura radne knjige
(listovi se ne mogu dodavati, brisati ni preimenovati). Osobni predložak ima upisane i zaključane ulazne podatke studenta te
JMBAG u prilagođenim svojstvima datoteke. Kod raspodjele "redom" ili "nasumicno" osobni
predložak ima samo list dodijeljenog zadatka.
"""

import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.worksheet import Worksheet

from provjera.excel.adapteri.rjesenja import RJESENJA
from provjera.excel.jezgra import dodjela
from provjera.excel.jezgra.profil import ProfilExcel
from provjera.excel.jezgra.provjera import (
    ADRESA_JMBAG,
    ADRESA_STUDENTA,
    SVOJSTVO_JMBAG,
    SVOJSTVO_STUDENT,
    SVOJSTVO_VJEZBA,
    SVOJSTVO_ZADATAK,
)
from provjera.excel.jezgra.zadaci import ZADACI, Vrsta, Zadatak, osobni_podaci
from provjera.zajednicko import naziv

PLAVA = PatternFill("solid", fgColor="DDEBF7")
ZUTA = PatternFill("solid", fgColor="FFF2CC")
SIVA = PatternFill("solid", fgColor="EDEDED")
_TANKO = Side(style="thin", color="808080")
OKVIR = Border(left=_TANKO, right=_TANKO, top=_TANKO, bottom=_TANKO)
SREDINA = Alignment(horizontal="center", vertical="center", wrap_text=True)
SLOBODNO_OD, SLOBODNO_DO = "M", "Z"
SLOBODNO_REDAKA = 60


@dataclass(frozen=True)
class StudentPredloska:
    ime: str
    prezime: str
    jmbag: str


# Oznake obrazaca: adresa → tekst; spojene ćelije; širine stupaca.
_OZNAKE: dict[str, dict[str, str]] = {
    "sinus": {
        "B6": "Trokut", "C6": "Mjereni kutovi", "F6": "Izjednačeni kutovi", "I6": "sin",
        "J6": "b·cos γ′\nc·cos β′", "K6": "Duljine", "L6": "Opis duljine",
        "C7": "°", "D7": "′", "E7": "″", "F7": "°", "G7": "′", "H7": "″", "I7": "sin α′, β′, γ′",
        "K7": "m", "B8": "α", "B9": "β", "B10": "γ", "B11": "Σ",
        "L8": "a – mjerena duljina", "L9": "b = m · sin β′", "L10": "c = m · sin γ′",
        "L13": "a = b·cos γ′ + c·cos β′ (kontrola)",
        "C12": "fβ = 180° − Σ  (″)", "C13": "vβ = fβ / 3  (″)", "H12": "m =",
        "F12": "Σv = fβ (cijele ″)",
        "I13": "← b·cos γ′ / c·cos β′ →",
    },
    "kosinus": {
        "B6": "Kut", "C6": "Kutovi", "F6": "Duljine", "I6": "Veličina", "J6": "Vrijednost",
        "C7": "°", "D7": "′", "E7": "″", "G7": "m",
        "B8": "α (mjereno)", "B9": "β", "B10": "γ", "B11": "Σ = 180°",
        "F8": "b", "F9": "c", "F13": "a = (b/sin β)·sin α", "F14": "a = (c/sin γ)·sin α",
        "I8": "b²", "I9": "c²", "I10": "cos α", "I11": "2bc · cos α", "I12": "a²", "I13": "a",
        "I14": "cos β", "I15": "cos γ",
    },
    "tangens": {
        "B6": "Kut", "C6": "+/−", "D6": "Kutovi", "G6": "Duljine", "J6": "Veličina", "K6": "Vrijednost",
        "D7": "°", "E7": "′", "F7": "″", "H7": "m",
        "B8": "α (mjereno)", "B9": "β", "B10": "γ", "B11": "Σ = 180°", "B12": "(β + γ) / 2",
        "B13": "(β − γ) / 2",
        "G8": "b", "G9": "c", "G10": "a", "G13": "Kontrola a",
        "J8": "α / 2  (°)", "J9": "B = cot(α / 2)", "J10": "b − c", "J11": "b + c",
        "J12": "A = (b − c) / (b + c)", "J13": "arctan(A · B)  (°)",
    },
}
_SPOJENO: dict[str, tuple[str, ...]] = {
    "sinus": ("C6:E6", "F6:H6", "C12:D12", "C13:D13", "F12:G12"),
    "kosinus": ("C6:E6", "F6:G6"),
    "tangens": ("D6:F6", "G6:H6"),
}
_SIRINE: dict[str, dict[str, float]] = {
    "sinus": {"A": 3, "B": 10, "I": 13, "J": 14, "K": 12, "L": 34},
    "kosinus": {"A": 3, "B": 14, "F": 20, "G": 12, "H": 3, "I": 14, "J": 16},
    "tangens": {"A": 3, "B": 14, "C": 6, "G": 12, "H": 12, "I": 3, "J": 22, "K": 16},
}
_UPUTA = (
    "Upute: u plave ćelije upisani su (ili upiši) mjereni podaci. U svaku žutu ćeliju upiši FORMULU "
    "(ne broj) koja računa traženu veličinu iz ulaznih ćelija. Pomoćne izračune (decimalni stupnjevi, "
    "radijani…) radi u slobodnom području od stupca M nadalje. Provjera ponovno izračunava tvoje formule "
    "i s drugim ulaznim podacima, pa formula mora vrijediti općenito. Sinusov poučak: odstupanje fβ "
    "rasporedi na tri kuta u cijelim sekundama (npr. fβ = 20″ → 7″, 7″, 6″) tako da zbroj izjednačenih "
    "kutova bude točno 180° 00′ 00″."
)


def _format(decimale: int) -> str:
    return "0" if decimale == 0 else "0." + "0" * decimale


def _celija(ws: Worksheet, adresa: str, vrijednost=None, *, ispuna=None, zakljucano=True,
            format_broja: str | None = None, okvir=True, podebljano=False) -> None:
    c = ws[adresa]
    if vrijednost is not None:
        c.value = vrijednost
    if ispuna is not None:
        c.fill = ispuna
    if okvir:
        c.border = OKVIR
    if format_broja:
        c.number_format = format_broja
    c.alignment = SREDINA
    if podebljano:
        c.font = Font(bold=True)
    c.protection = Protection(locked=zakljucano)


def _crtaj_zadatak(
    ws: Worksheet, profil: ProfilExcel, zadatak: Zadatak, student: StudentPredloska | None,
    rjesenje: bool, lozinka: str,
) -> None:
    ws["A1"] = zadatak.obrazac
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = profil.vjezba.naziv
    ws["A3"], ws["A4"] = "Ime i prezime:", "JMBAG:"
    for a, v in ((ADRESA_STUDENTA, f"{student.ime} {student.prezime}" if student else None),
                 (ADRESA_JMBAG, student.jmbag if student else None)):
        ws.merge_cells(f"{a}:H{a[1:]}")
        _celija(ws, a, v, ispuna=None if student else PLAVA, zakljucano=student is not None, okvir=False,
                podebljano=True)
        ws[a].alignment = Alignment(horizontal="left")
    for adresa, tekst in _OZNAKE[zadatak.id].items():
        _celija(ws, adresa, tekst, ispuna=SIVA, podebljano=True)
    for raspon in _SPOJENO[zadatak.id]:
        ws.merge_cells(raspon)
    for stupac, sirina in _SIRINE[zadatak.id].items():
        ws.column_dimensions[stupac].width = sirina
    for stupac in "CDEFGH":
        if stupac not in _SIRINE[zadatak.id]:
            ws.column_dimensions[stupac].width = 7

    podaci = osobni_podaci(zadatak, profil.zadaci.sjeme, student.jmbag, profil.zadaci.korak_mjerenja) if (
        student and profil.zadaci.osobni_podaci) else None
    for u in zadatak.ulazi:
        vrijednost = podaci[u.kljuc] if podaci else None
        if vrijednost is not None and u.decimale == 0:
            vrijednost = int(vrijednost)
        _celija(ws, u.adresa, vrijednost, ispuna=PLAVA, zakljucano=podaci is not None,
                format_broja=_format(u.decimale))
    for r in zadatak.rezultati:
        adrese = r.adrese + ((r.predznak,) if r.predznak else ())
        for i, a in enumerate(adrese):
            decimale = 0 if r.vrsta is Vrsta.KUT and i < 2 else (1 if r.vrsta is Vrsta.KUT else r.decimale)
            _celija(ws, a, ispuna=ZUTA, zakljucano=False, format_broja=_format(decimale))

    # Popis formula ispod obrasca
    redak = 18
    ws.cell(redak, 2, "Što se računa u žutim ćelijama").font = Font(bold=True)
    for r in zadatak.rezultati:
        redak += 1
        adr = r.adrese[0] if len(r.adrese) == 1 else f"{r.adrese[0]}:{r.adrese[-1]}"
        ws.cell(redak, 2, adr)
        ws.cell(redak, 3, f"{r.opis}:  {r.formula}")
    ws.cell(redak + 2, 2, _UPUTA).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=redak + 2, start_column=2, end_row=redak + 6, end_column=12)

    # Slobodno područje
    ws[f"{SLOBODNO_OD}6"] = "Pomoćni izračuni (slobodno područje)"
    ws[f"{SLOBODNO_OD}6"].font = Font(bold=True, italic=True)
    for stupac in range(column_index_from_string(SLOBODNO_OD), column_index_from_string(SLOBODNO_DO) + 1):
        for r_ in range(7, SLOBODNO_REDAKA + 1):
            ws.cell(r_, stupac).protection = Protection(locked=False)
        ws.column_dimensions[get_column_letter(stupac)].width = max(
            ws.column_dimensions[get_column_letter(stupac)].width or 0, 12)

    if rjesenje:
        for adresa, formula in RJESENJA[zadatak.id]().items():
            ws[adresa] = formula

    ws.protection.sheet = True
    if lozinka:
        ws.protection.password = lozinka
    # Dopušteno: širina stupaca, visina redaka i oblikovanje; zabranjeno: umetanje i brisanje.
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.formatCells = False


def izradi(profil: ProfilExcel, student: StudentPredloska | None = None, *, rjesenje: bool = False,
           redni_broj: int | None = None, lozinka: str | None = None) -> bytes:
    """Predložak vježbe; ``student`` daje osobni predložak, ``rjesenje`` upisuje formule.

    ``redni_broj`` (položaj studenta u popisu) određuje zadatak kod raspodjele "redom".
    Prazan predložak i rješenje imaju sve zadatke. ``lozinka`` zaštite listova i strukture
    (bez nje lozinka iz profila); ne sprema se u repozitorij.
    """
    lozinka = profil.zadaci.lozinka if lozinka is None else lozinka
    vrste = (dodjela.dodijeljeni(profil, student.jmbag, redni_broj) if student else None) or profil.zadaci.vrste
    wb = Workbook()
    wb.remove(wb.active)
    for vrsta in vrste:
        zadatak = ZADACI[vrsta]
        _crtaj_zadatak(wb.create_sheet(zadatak.list), profil, zadatak, student, rjesenje, lozinka)
    if lozinka:
        wb.security = WorkbookProtection(workbookPassword=lozinka, lockStructure=True)
    wb.properties.creator = "Provjera vježbi"
    wb.properties.title = profil.vjezba.naziv
    wb.custom_doc_props.append(StringProperty(name=SVOJSTVO_VJEZBA, value=profil.vjezba.id))
    if student:
        wb.custom_doc_props.append(StringProperty(name=SVOJSTVO_JMBAG, value=student.jmbag))
        wb.custom_doc_props.append(StringProperty(name=SVOJSTVO_STUDENT, value=f"{student.ime} {student.prezime}"))
        wb.custom_doc_props.append(StringProperty(name=SVOJSTVO_ZADATAK, value=",".join(vrste)))
    spremnik = io.BytesIO()
    wb.save(spremnik)
    return spremnik.getvalue()


def naziv_predloska(profil: ProfilExcel, student: StudentPredloska) -> str:
    return naziv.kanonski(student.ime, student.prezime, profil)


def izradi_za_studente(
    profil: ProfilExcel, studenti: Iterable[StudentPredloska], mapa: Path, lozinka: str | None = None
) -> Sequence[Path]:
    """Osobni predložak za svakog studenta, već s ispravnim nazivom za predaju."""
    mapa.mkdir(parents=True, exist_ok=True)
    putanje = []
    for i, student in enumerate(studenti):
        putanja = mapa / naziv_predloska(profil, student)
        putanja.write_bytes(izradi(profil, student, redni_broj=i, lozinka=lozinka))
        putanje.append(putanja)
    return putanje

"""Skupna provjera mape predanih Excel vježbi i zbirni izvještaj (.xlsx).

Predaja se uparuje sa studentom iz ISVU popisa po nazivu datoteke, a ako to ne uspije,
po JMBAG-u koji je osobni predložak upisao u svojstva datoteke. Izvještaj ima tri lista:
pregled po studentu, sve stavke s porukama i parove predaja sličnih formula.
"""

import io
from collections import defaultdict
from collections.abc import Sequence
from dataclasses import dataclass
from itertools import combinations
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from provjera.excel.adapteri.predlozak import StudentPredloska
from provjera.excel.jezgra.profil import ProfilExcel
from provjera.excel.jezgra.provjera import IzvjestajExcel, provjeri, slicnost
from provjera.zajednicko import naziv
from provjera.zajednicko.status import Status

PRAG_SLICNOSTI = 0.6
_NAZIVI_STATUSA = {Status.PROSLO: "prošlo", Status.NE_PROLAZI: "ne prolazi", Status.NEPROVJERLJIVO: "neprovjerljivo"}
_BOJE = {Status.PROSLO: "C6EFCE", Status.NE_PROLAZI: "FFC7CE", Status.NEPROVJERLJIVO: "FFEB9C"}


@dataclass(frozen=True)
class Predaja:
    putanja: Path
    student: StudentPredloska | None
    izvjestaj: IzvjestajExcel


def upari(putanja: Path, profil: ProfilExcel, studenti: Sequence[StudentPredloska],
          jmbag_iz_datoteke: str | None) -> StudentPredloska | None:
    kljuc = naziv.kljuc_uparivanja(putanja.name)
    for s in studenti:
        if naziv.kljuc_uparivanja(naziv.kanonski(s.ime, s.prezime, profil)) == kljuc:
            return s
    if jmbag_iz_datoteke:
        return next((s for s in studenti if s.jmbag == jmbag_iz_datoteke), None)
    return None


def provjeri_mapu(mapa: Path, profil: ProfilExcel, studenti: Sequence[StudentPredloska] = ()) -> list[Predaja]:
    predaje = []
    for putanja in sorted(p for p in mapa.rglob("*") if p.is_file() and not p.name.startswith("~$")):
        if putanja.suffix.lower() not in (".xlsx", ".xls", ".xlsm"):
            continue
        sadrzaj = putanja.read_bytes()
        prvi = provjeri(putanja.name, sadrzaj, profil)
        student = upari(putanja, profil, studenti, prvi.jmbag_iz_datoteke) if studenti else None
        izvjestaj = provjeri(putanja.name, sadrzaj, profil, student.jmbag, studenti.index(student))             if student else prvi
        predaje.append(Predaja(putanja, student, izvjestaj))
    return predaje


def slicni_parovi(predaje: Sequence[Predaja], prag: float = PRAG_SLICNOSTI) -> list[tuple[float, Predaja, Predaja]]:
    parovi = [
        (slicnost(a.izvjestaj, b.izvjestaj), a, b)
        for a, b in combinations(predaje, 2)
        if a.izvjestaj.odbijeno is None and b.izvjestaj.odbijeno is None
    ]
    return sorted((p for p in parovi if p[0] >= prag), key=lambda p: -p[0])


def isti_jmbag_u_vise_datoteka(predaje: Sequence[Predaja]) -> dict[str, list[Predaja]]:
    """Ista kopija osobnog predloška predana pod više imena."""
    po_jmbagu: dict[str, list[Predaja]] = defaultdict(list)
    for p in predaje:
        if p.izvjestaj.jmbag_iz_datoteke:
            po_jmbagu[p.izvjestaj.jmbag_iz_datoteke].append(p)
    return {j: ps for j, ps in po_jmbagu.items() if len(ps) > 1}


def _ime(p: Predaja) -> str:
    return f"{p.student.ime} {p.student.prezime}" if p.student else ""


def _zaglavlje(ws, stupci: Sequence[str]) -> None:
    ws.append(list(stupci))
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _sirine(ws) -> None:
    for i, stupac in enumerate(ws.columns, start=1):
        najdulje = max((len(str(c.value)) for c in stupac if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, najdulje + 2), 80)


def zbirni_izvjestaj(predaje: Sequence[Predaja]) -> bytes:
    wb = Workbook()
    pregled = wb.active
    pregled.title = "Pregled"
    _zaglavlje(pregled, ["Datoteka", "Student", "JMBAG", "Ukupno", "Prošlo", "Ne prolazi",
                         "Neprovjerljivo", "Zadnji spremio", "Napomena"])
    for p in predaje:
        iz = p.izvjestaj
        napomena = f"odbijeno: {iz.odbijeno.value}" if iz.odbijeno else ("" if p.student else "nije uparen s popisom")
        pregled.append([
            p.putanja.name, _ime(p), p.student.jmbag if p.student else (iz.jmbag_iz_datoteke or ""),
            _NAZIVI_STATUSA[iz.status], iz.broj(Status.PROSLO), iz.broj(Status.NE_PROLAZI),
            iz.broj(Status.NEPROVJERLJIVO), iz.svojstva.zadnji_izmijenio if iz.svojstva else "", napomena,
        ])
        pregled.cell(pregled.max_row, 4).fill = PatternFill("solid", fgColor=_BOJE[iz.status])

    detalji = wb.create_sheet("Detalji")
    _zaglavlje(detalji, ["Datoteka", "Student", "Zadatak", "Stavka", "Ćelije", "Formula", "Vrijednost",
                         "Općenito", "Poruke"])
    for p in predaje:
        for s in p.izvjestaj.stavke:
            detalji.append([
                p.putanja.name, _ime(p), s.zadatak, s.opis, s.adresa,
                _NAZIVI_STATUSA[s.formula], _NAZIVI_STATUSA[s.vrijednost], _NAZIVI_STATUSA[s.opcenito],
                " ".join(s.poruke),
            ])
            for stupac, status in ((6, s.formula), (7, s.vrijednost), (8, s.opcenito)):
                if status is not Status.PROSLO:
                    detalji.cell(detalji.max_row, stupac).fill = PatternFill("solid", fgColor=_BOJE[status])

    slicnosti = wb.create_sheet("Sličnost")
    _zaglavlje(slicnosti, ["Sličnost formula", "Datoteka A", "Student A", "Datoteka B", "Student B", "Napomena"])
    for jmbag, ps in isti_jmbag_u_vise_datoteka(predaje).items():
        slicnosti.append(["isti predložak", ps[0].putanja.name, _ime(ps[0]), ps[1].putanja.name, _ime(ps[1]),
                          f"oba nose JMBAG {jmbag} iz osobnog predloška"])
    for koef, a, b in slicni_parovi(predaje):
        isti_spremio = (a.izvjestaj.svojstva and b.izvjestaj.svojstva
                        and a.izvjestaj.svojstva.zadnji_izmijenio
                        and a.izvjestaj.svojstva.zadnji_izmijenio == b.izvjestaj.svojstva.zadnji_izmijenio)
        slicnosti.append([round(koef, 2), a.putanja.name, _ime(a), b.putanja.name, _ime(b),
                          "isti korisnik zadnji spremio datoteke" if isti_spremio else ""])
    for ws in (pregled, detalji, slicnosti):
        _sirine(ws)
    spremnik = io.BytesIO()
    wb.save(spremnik)
    return spremnik.getvalue()


def tekst_izvjestaja(izvjestaj: IzvjestajExcel) -> str:
    """Čitljiv izvještaj jedne datoteke za naredbeni redak."""
    redovi = [f"{izvjestaj.naziv_datoteke}: {_NAZIVI_STATUSA[izvjestaj.status].upper()}"]
    if izvjestaj.odbijeno:
        return redovi[0] + f" (odbijeno: {izvjestaj.odbijeno.value})"
    redovi.append(
        f"  prošlo {izvjestaj.broj(Status.PROSLO)}, ne prolazi {izvjestaj.broj(Status.NE_PROLAZI)}, "
        f"neprovjerljivo {izvjestaj.broj(Status.NEPROVJERLJIVO)}"
    )
    zadatak = None
    for s in izvjestaj.stavke:
        if s.zadatak != zadatak:
            zadatak = s.zadatak
            redovi.append(f"  [{zadatak or 'datoteka'}]")
        oznaka = {Status.PROSLO: "✓", Status.NE_PROLAZI: "✗", Status.NEPROVJERLJIVO: "?"}[s.status]
        redovi.append(f"    {oznaka} {s.opis}" + (f" ({s.adresa})" if s.adresa else ""))
        redovi += [f"        {poruka}" for poruka in s.poruke]
    return "\n".join(redovi)

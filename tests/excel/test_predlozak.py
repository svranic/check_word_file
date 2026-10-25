import io
from pathlib import Path

import openpyxl
import pytest

from provjera.excel.__main__ import main
from provjera.excel.adapteri import skupno
from provjera.excel.adapteri.predlozak import StudentPredloska, izradi, izradi_za_studente
from provjera.excel.jezgra import dodjela
from provjera.excel.jezgra import profil as profil_modul
from provjera.excel.jezgra.provjera import provjeri
from provjera.excel.jezgra.zadaci import ZADACI, osobni_podaci

KORIJEN = Path(__file__).resolve().parents[2]
PUTANJA_PROFILA = KORIJEN / "profili" / "excel" / "vjezba_2.toml"
PROFIL = profil_modul.ucitaj(PUTANJA_PROFILA.read_bytes())
ANA = StudentPredloska("Ana", "Horvat", "0036000001")


def test_osobni_predlozak_ima_zakljucane_podatke_i_svojstva():
    wb = openpyxl.load_workbook(io.BytesIO(izradi(PROFIL, ANA)))
    assert wb.sheetnames == [ZADACI[v].list for v in PROFIL.zadaci.vrste]
    zadatak = ZADACI["sinus"]
    ws = wb[zadatak.list]
    podaci = osobni_podaci(zadatak, PROFIL.zadaci.sjeme, ANA.jmbag, PROFIL.zadaci.korak_mjerenja)
    for u in zadatak.ulazi:
        assert ws[u.adresa].value == pytest.approx(podaci[u.kljuc])
        assert ws[u.adresa].protection.locked
    for r in zadatak.rezultati:
        assert all(not ws[a].protection.locked for a in r.adrese)
        assert ws[r.adrese[0]].value is None
    assert ws.protection.sheet
    assert not ws["M7"].protection.locked  # slobodno područje
    svojstva = {p.name: p.value for p in wb.custom_doc_props.props}
    assert svojstva["provjera_jmbag"] == ANA.jmbag


def test_prazan_predlozak_ima_otkljucane_ulaze():
    ws = openpyxl.load_workbook(io.BytesIO(izradi(PROFIL)))["Kosinusov poučak"]
    assert ws["G8"].value is None and not ws["G8"].protection.locked


def test_osobni_predlosci_imaju_naziv_za_predaju(tmp_path):
    putanje = izradi_za_studente(PROFIL, [ANA], tmp_path)
    assert [p.name for p in putanje] == ["Ana_Horvat_GiG_OG_Vjezba_2.xlsx"]


def profil_s(**zadaci) -> profil_modul.ProfilExcel:
    return PROFIL.model_copy(update={"zadaci": PROFIL.zadaci.model_copy(update=zadaci)})


def test_raspodjela_redom_daje_zadatke_naizmjence():
    p = profil_s(raspodjela="redom")
    assert [dodjela.dodijeljeni(p, "x", i) for i in range(4)] == [
        ("sinus",), ("kosinus",), ("tangens",), ("sinus",)]
    assert dodjela.dodijeljeni(p, "x") is None  # bez položaja u popisu


def test_raspodjela_nasumicno_ovisi_samo_o_jmbagu_i_sjemenu():
    p = profil_s(raspodjela="nasumicno")
    dodijeljeno = [dodjela.dodijeljeni(p, f"00360{i:05d}") for i in range(60)]
    assert dodijeljeno == [dodjela.dodijeljeni(p, f"00360{i:05d}", 7) for i in range(60)]
    assert {d[0] for d in dodijeljeno} == set(p.zadaci.vrste)
    assert dodjela.dodijeljeni(profil_s(vrste=("sinus", "kosinus", "tangens"), raspodjela="svi"), "x") == (
        "sinus", "kosinus", "tangens")


def test_osobni_predlozak_ima_samo_dodijeljeni_zadatak(tmp_path):
    p = profil_s(raspodjela="redom")
    ivo = StudentPredloska("Ivo", "Kovač", "0036000002")
    izradi_za_studente(p, [ANA, ivo], tmp_path)
    wb = openpyxl.load_workbook(tmp_path / "Ivo_Kovač_GiG_OG_Vjezba_2.xlsx")
    assert wb.sheetnames == ["Kosinusov poučak"]
    assert {q.name: q.value for q in wb.custom_doc_props.props}["provjera_zadatak"] == "kosinus"


def test_redom_vrijedi_zadatak_iz_osobnog_predloska():
    # Ana je pri izradi bila druga na popisu (kosinus); u kasnijem popisu je prva.
    p = profil_s(raspodjela="redom")
    rjesenje = izradi(p, ANA, rjesenje=True, redni_broj=1)
    izvjestaj = provjeri("Ana_Horvat_GiG_OG_Vjezba_2.xlsx", rjesenje, p, ANA.jmbag, 0)
    assert {s.zadatak for s in izvjestaj.stavke} == {"", "kosinus"}
    assert izvjestaj.status.value == "PROSLO"


def test_predan_list_pogresnog_zadatka():
    p = profil_s(raspodjela="nasumicno")
    dodijeljen = dodjela.dodijeljeni(p, ANA.jmbag)[0]
    drugi = next(v for v in p.zadaci.vrste if v != dodijeljen)
    # Ana je ispunila prazan predložak i ostavila samo list drugog zadatka.
    wb = openpyxl.load_workbook(io.BytesIO(izradi(p, rjesenje=True)))
    for ime in [ZADACI[v].list for v in p.zadaci.vrste if v != drugi]:
        del wb[ime]
    spremnik = io.BytesIO()
    wb.save(spremnik)
    izvjestaj = provjeri("Ana_Horvat_GiG_OG_Vjezba_2.xlsx", spremnik.getvalue(), p, ANA.jmbag)
    s = next(s for s in izvjestaj.stavke if s.kljuc == "list")
    assert s.zadatak == dodijeljen and "dodijeljeni zadatak" in s.poruke[0]


def test_profil_odbija_korak_koji_ne_dijeli_minutu():
    tekst = PUTANJA_PROFILA.read_text(encoding="utf-8").replace("korak_mjerenja = 10", "korak_mjerenja = 7")
    with pytest.raises(profil_modul.GreskaProfila):
        profil_modul.ucitaj(tekst.encode("utf-8"))


def test_profil_odbija_nepoznat_zadatak():
    tekst = PUTANJA_PROFILA.read_text(encoding="utf-8").replace('"tangens"]', '"kotangens"]')
    with pytest.raises(profil_modul.GreskaProfila):
        profil_modul.ucitaj(tekst.encode("utf-8"))


def test_skupna_provjera_i_izvjestaj(tmp_path):
    mapa = tmp_path / "predaje"
    mapa.mkdir()
    ivo = StudentPredloska("Ivo", "Kovač", "0036000002")
    (mapa / "Ana_Horvat_GiG_OG_Vjezba_2.xlsx").write_bytes(izradi(PROFIL, ANA, rjesenje=True))
    # Ivo je predao Anin riješeni predložak pod svojim imenom.
    (mapa / "Ivo_Kovač_GiG_OG_Vjezba_2.xlsx").write_bytes(izradi(PROFIL, ANA, rjesenje=True))
    predaje = skupno.provjeri_mapu(mapa, PROFIL, [ANA, ivo])
    assert [p.student for p in predaje] == [ANA, ivo]
    assert predaje[0].izvjestaj.status.value == "PROSLO"
    assert predaje[1].izvjestaj.status.value == "NE_PROLAZI"
    assert skupno.isti_jmbag_u_vise_datoteka(predaje)
    assert skupno.slicni_parovi(predaje)[0][0] == 1.0
    listovi = openpyxl.load_workbook(io.BytesIO(skupno.zbirni_izvjestaj(predaje))).sheetnames
    assert listovi == ["Pregled", "Detalji", "Sličnost"]


def test_naredbeni_redak(tmp_path, capsys):
    assert main(["predlozak", str(PUTANJA_PROFILA), str(tmp_path), "--rjesenje"]) == 0
    rjesenje = tmp_path / "RJESENJE_vjezba_2.xlsx"
    assert rjesenje.exists()
    assert main(["provjeri", str(PUTANJA_PROFILA), str(rjesenje)]) == 0
    izlaz = capsys.readouterr().out
    # Prazan predložak s rješenjem nema ulaznih podataka pa rezultati nisu provjerljivi.
    assert "RJESENJE_vjezba_2.xlsx: NE PROLAZI" in izlaz
    assert "Nedostaju ulazni podaci" in izlaz

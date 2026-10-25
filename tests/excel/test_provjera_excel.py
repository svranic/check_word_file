"""Provjera predaje na predlošcima koje izrađuje sama aplikacija (bez Excela).

openpyxl ne sprema izračunate vrijednosti, pa se ovdje vrijednosti računaju iz formula;
``s_vrijednostima`` dodaje spremljene vrijednosti kakve bi upisao Excel.
"""

import io
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from provjera.excel.adapteri.predlozak import StudentPredloska, izradi
from provjera.excel.jezgra import profil as profil_modul
from provjera.excel.jezgra.provjera import IzvjestajExcel, provjeri, slicnost
from provjera.zajednicko.status import Status
from provjera.zajednicko.ulaz import RazlogOdbijanja

KORIJEN = Path(__file__).resolve().parents[2]
PROFIL = profil_modul.ucitaj((KORIJEN / "profili" / "excel" / "vjezba_2.toml").read_bytes())
ANA = StudentPredloska("Ana", "Horvat", "0036000001")
NAZIV = "Ana_Horvat_GiG_OG_Vjezba_2.xlsx"


def izmijeni(sadrzaj: bytes, izmjene: dict[tuple[str, str], object]) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(sadrzaj))
    for (list_, adresa), vrijednost in izmjene.items():
        wb[list_][adresa] = vrijednost
    spremnik = io.BytesIO()
    wb.save(spremnik)
    return spremnik.getvalue()


def s_vrijednostima(sadrzaj: bytes, list_br: int, vrijednosti: dict[str, float]) -> bytes:
    """Upisuje spremljenu vrijednost (``<v>``) u ćelije s formulom, kao da je Excel spremio."""
    ulaz, izlaz = zipfile.ZipFile(io.BytesIO(sadrzaj)), io.BytesIO()
    with zipfile.ZipFile(izlaz, "w") as z:
        for stavka in ulaz.infolist():
            podaci = ulaz.read(stavka)
            if stavka.filename == f"xl/worksheets/sheet{list_br}.xml":
                tekst = podaci.decode("utf-8")
                for adresa, v in vrijednosti.items():
                    tekst, n = re.subn(rf'(<c r="{adresa}"[^>]*>)(<f>[^<]*</f>)(<v>[^<]*</v>)?', rf"\1\2<v>{v}</v>", tekst)
                    assert n == 1, adresa
                podaci = tekst.encode("utf-8")
            z.writestr(stavka, podaci)
    return izlaz.getvalue()


def stavka(izvjestaj: IzvjestajExcel, zadatak: str, kljuc: str):
    return next(s for s in izvjestaj.stavke if s.zadatak == zadatak and s.kljuc == kljuc)


@pytest.fixture(scope="module")
def rjesenje() -> bytes:
    return izradi(PROFIL, ANA, rjesenje=True)


def test_rjesenje_prolazi_sve(rjesenje):
    izvjestaj = provjeri(NAZIV, rjesenje, PROFIL)
    neprosle = [(s.zadatak, s.kljuc, s.poruke) for s in izvjestaj.stavke if s.status is not Status.PROSLO]
    assert neprosle == []
    assert izvjestaj.jmbag_iz_datoteke == ANA.jmbag


def test_prazan_predlozak_ne_prolazi():
    izvjestaj = provjeri(NAZIV, izradi(PROFIL, ANA), PROFIL)
    s = stavka(izvjestaj, "sinus", "m")
    assert s.status is Status.NE_PROLAZI and s.poruke == ("Ćelija je prazna.",)
    assert stavka(izvjestaj, "sinus", "ulazi").status is Status.PROSLO


def test_upisan_broj_umjesto_formule(rjesenje):
    b = provjeri(NAZIV, izmijeni(rjesenje, {("Sinusov poučak", "K9"): 120.64}), PROFIL)
    s = stavka(b, "sinus", "b")
    assert s.formula is Status.NE_PROLAZI
    assert s.opcenito is Status.NE_PROLAZI
    assert "Upisan broj umjesto formule u K9." in s.poruke


def test_prepisana_vrijednost_u_formuli_prolazi_vrijednost_ali_ne_opcenitost(rjesenje):
    wb = openpyxl.load_workbook(io.BytesIO(rjesenje))
    a = wb["Sinusov poučak"]["K8"].value
    izvjestaj = provjeri(NAZIV, izmijeni(rjesenje, {("Sinusov poučak", "I12"): f"={a}/I8"}), PROFIL)
    s = stavka(izvjestaj, "sinus", "m")
    assert (s.formula, s.vrijednost, s.opcenito) == (Status.PROSLO, Status.PROSLO, Status.NE_PROLAZI)
    assert any("prepisanu vrijednost" in p for p in s.poruke)
    # b ovisi o m: greška se označava kao moguća posljedica
    assert any("posljedica greške u I12" in p for p in stavka(izvjestaj, "sinus", "b").poruke)


def test_upisan_broj_u_pomocnoj_celiji(rjesenje):
    wb = openpyxl.load_workbook(io.BytesIO(rjesenje), data_only=False)
    ws = wb["Kosinusov poučak"]
    alfa = ws["C8"].value + ws["D8"].value / 60 + ws["E8"].value / 3600
    izvjestaj = provjeri(NAZIV, izmijeni(rjesenje, {("Kosinusov poučak", "N8"): alfa}), PROFIL)
    s = stavka(izvjestaj, "kosinus", "cos_alfa")
    assert s.vrijednost is Status.PROSLO and s.opcenito is Status.NE_PROLAZI
    assert any(p.startswith("Ćelija N8 sadrži upisan broj") for p in s.poruke)


def test_pogresna_formula(rjesenje):
    izvjestaj = provjeri(NAZIV, izmijeni(rjesenje, {("Kosinusov poučak", "J14"): "=(J12+J8-J9)/(2*J13*G9)"}), PROFIL)
    s = stavka(izvjestaj, "kosinus", "cos_beta")
    assert s.vrijednost is Status.NE_PROLAZI and s.opcenito is Status.NE_PROLAZI


def test_rucno_upisan_predznak_ne_vrijedi_opcenito(rjesenje):
    # Kontrolni podaci imaju i b > c i b < c, pa ručni „+“ ili „-“ ne može proći.
    for znak in ("+", "-"):
        izvjestaj = provjeri(NAZIV, izmijeni(rjesenje, {("Tangensov poučak", "C13"): znak}), PROFIL)
        assert stavka(izvjestaj, "tangens", "razlika_pola").opcenito is Status.NE_PROLAZI


def test_nerasporedeni_popravak_ne_prolazi(rjesenje):
    # α′ = α + fβ/3 bez raspoređivanja: sekunde nisu cijele (osim kad je fβ djeljiv s 3).
    izmjene = {("Sinusov poučak", f"O{r}"): f"=N{r}+$E$13/3600" for r in (8, 9, 10)}
    izvjestaj = provjeri(NAZIV, izmijeni(rjesenje, izmjene), PROFIL)
    s = stavka(izvjestaj, "sinus", "alfa_izj")
    assert s.opcenito is Status.NE_PROLAZI
    assert any("cijelim sekundama" in p for p in s.poruke)


def test_svaki_kut_zaokruzen_posebno_ne_daje_180(rjesenje):
    izmjene = {("Sinusov poučak", f"R{r}"): "=ROUND($E$12/3,0)" for r in (8, 9, 10)}
    izvjestaj = provjeri(NAZIV, izmijeni(rjesenje, izmjene), PROFIL)
    assert stavka(izvjestaj, "sinus", "alfa_izj").status is Status.PROSLO
    assert stavka(izvjestaj, "sinus", "zbroj_izj").opcenito is Status.NE_PROLAZI


def test_ostatak_popravka_smije_dobiti_bilo_koji_kut(rjesenje):
    izmjene = {("Sinusov poučak", "R8"): "=$E$12-R9-R10", ("Sinusov poučak", "R10"): "=ROUND($E$12/3,0)"}
    izvjestaj = provjeri(NAZIV, izmijeni(rjesenje, izmjene), PROFIL)
    assert izvjestaj.status is Status.PROSLO


def test_promijenjeni_ulazni_podaci(rjesenje):
    izvjestaj = provjeri(NAZIV, izmijeni(rjesenje, {("Kosinusov poučak", "G8"): 100.0}), PROFIL)
    s = stavka(izvjestaj, "kosinus", "ulazi")
    assert s.vrijednost is Status.NE_PROLAZI
    assert "Mjerena duljina b (G8): 100 umjesto" in s.poruke[0]
    # Rezultati se i dalje provjeravaju prema podacima u datoteci.
    assert stavka(izvjestaj, "kosinus", "a").status is Status.PROSLO


def test_predlozak_drugog_studenta(rjesenje):
    izvjestaj = provjeri("Ivo_Kovac_GiG_OG_Vjezba_2.xlsx", rjesenje, PROFIL, jmbag="0036000002")
    assert stavka(izvjestaj, "", "izdano").status is Status.NE_PROLAZI
    assert stavka(izvjestaj, "sinus", "ulazi").vrijednost is Status.NE_PROLAZI


def test_prazan_predlozak_bez_studenta_ne_uspoređuje_ulaze():
    izvjestaj = provjeri(NAZIV, izradi(PROFIL, rjesenje=True), PROFIL)
    s = stavka(izvjestaj, "sinus", "ulazi")
    assert s.formula is Status.NE_PROLAZI  # bez upisanih podataka nema ni rezultata
    assert stavka(izvjestaj, "sinus", "m").status is Status.NEPROVJERLJIVO


def test_spremljena_vrijednost_ima_prednost(rjesenje):
    # Excel je spremio pogrešnu vrijednost (npr. ručni izračun isključen): vrijednost ne
    # prolazi, a općenitost je neprovjerljiva jer formula i spremljeno nisu usklađeni.
    izvjestaj = provjeri(NAZIV, s_vrijednostima(rjesenje, 1, {"I12": 1.0}), PROFIL)
    s = stavka(izvjestaj, "sinus", "m")
    assert s.vrijednost is Status.NE_PROLAZI and s.opcenito is Status.NEPROVJERLJIVO


def test_preimenovan_list_se_uparuje_po_redoslijedu(rjesenje):
    wb = openpyxl.load_workbook(io.BytesIO(rjesenje))
    wb["Sinusov poučak"].title = "List1"
    spremnik = io.BytesIO()
    wb.save(spremnik)
    assert provjeri(NAZIV, spremnik.getvalue(), PROFIL).status is Status.PROSLO


def test_naziv_datoteke(rjesenje):
    izvjestaj = provjeri("zadaca.xlsx", rjesenje, PROFIL)
    assert stavka(izvjestaj, "", "naziv").status is Status.NE_PROLAZI


@pytest.mark.parametrize(("naziv", "sadrzaj", "razlog"), [
    ("rad.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 10, RazlogOdbijanja.STARI_DOC),
    ("rad.xlsx", b"", RazlogOdbijanja.PRAZNA),
    ("rad.docx", b"PK\x03\x04", RazlogOdbijanja.NEISPRAVAN_FORMAT),
])
def test_odbijene_datoteke(naziv, sadrzaj, razlog):
    izvjestaj = provjeri(naziv, sadrzaj, PROFIL)
    assert izvjestaj.odbijeno is razlog and izvjestaj.status is Status.NE_PROLAZI


def test_slicnost(rjesenje):
    a = provjeri(NAZIV, rjesenje, PROFIL)
    b = provjeri(NAZIV, izmijeni(rjesenje, {("Sinusov poučak", "M30"): "=1+1"}), PROFIL)
    assert slicnost(a, a) == 1.0
    assert 0.9 < slicnost(a, b) < 1.0


def test_promijenjen_jmbag_u_zaglavlju(rjesenje):
    izvjestaj = provjeri(NAZIV, izmijeni(rjesenje, {("Sinusov poučak", "D4"): "0036999999"}), PROFIL)
    s = stavka(izvjestaj, "", "zaglavlje")
    assert s.status is Status.NE_PROLAZI and "0036999999" in s.poruke[0]
    assert stavka(provjeri(NAZIV, rjesenje, PROFIL), "", "zaglavlje").status is Status.PROSLO


def test_predlozak_je_zasticen_lozinkom():
    wb = openpyxl.load_workbook(io.BytesIO(izradi(PROFIL, ANA, lozinka="tajna")))
    ws = wb["Sinusov poučak"]
    assert ws.protection.sheet and ws.protection.password
    assert all(ws[f"{s}{r}"].protection.locked for s in "ABCDEFGH" for r in range(1, 5))
    assert wb.security.lockStructure

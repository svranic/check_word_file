"""R-PROMJENE i R-OKVIRI: prihvaćeno stanje praćenih promjena i tekstni okviri (priča 2.6)."""

import io
from pathlib import Path

import docx as python_docx
import pytest
from alati.generator import VARIJANTE, ispravan, okvir_run
from docx.oxml import OxmlElement, parse_xml
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import nsdecls, qn

from provjera.jezgra import model, profil
from provjera.jezgra.model.uloge import UlogaOdlomka, uloga_okvira
from provjera.jezgra.nalaz import Dio, KontekstProvjere, Status, VrstaIzvora
from provjera.jezgra.poruke import KATALOG, tekst_nalaza
from provjera.jezgra.pravila import REGISTAR
from provjera.jezgra.provjera import provjeri, status_pravila

KORIJEN = Path(__file__).resolve().parents[1]
TOML = (KORIJEN / "profili" / "vjezba_1.toml").read_text(encoding="utf-8")
PROFIL = profil.ucitaj(TOML.encode("utf-8"))
KONTEKST = KontekstProvjere("Ana_Horvat_GiG_OG_Vjezba_1.docx", 54, None, None)
W = nsdecls("w")
PROMJENA = 'w:author="Ana" w:date="2026-10-14T10:00:00Z"'


def _izmijeni(promjena, sadrzaj: bytes | None = None) -> bytes:
    dokument = python_docx.Document(io.BytesIO(sadrzaj or ispravan(54, PROFIL)))
    promjena(dokument)
    spremnik = io.BytesIO()
    dokument.save(spremnik)
    return spremnik.getvalue()


def _model(sadrzaj: bytes):
    return model.procitaj(sadrzaj, PROFIL.uloge)


def _nalazi(sadrzaj: bytes, pravilo_id: str):
    return [n for n in provjeri(sadrzaj, PROFIL, KONTEKST).nalazi if n.pravilo_id == pravilo_id]


def _padaju(sadrzaj: bytes) -> set[str]:
    return {n.pravilo_id for n in provjeri(sadrzaj, PROFIL, KONTEKST).nalazi if n.status is not Status.PROSLO}


def _osnovni(dokument):
    """Odlomci osnovnog teksta ispravnog rada (Normal iza prvog naslova, bez literature)."""
    odlomci = dokument.paragraphs
    prvi = next(i for i, o in enumerate(odlomci) if o.style.name == "Heading 1")
    return [o for o in odlomci[prvi:] if o.style.name == "Normal" and o.text]


# --- matrica --------------------------------------------------------------------------


@pytest.mark.parametrize("ime", ["promjene_umetanje", "promjene_brisanje", "promjene_oblik"])
def test_promjene_daju_jedan_nalaz_prisutne(ime):
    sadrzaj = VARIJANTE[ime](54, PROFIL)
    [nalaz] = _nalazi(sadrzaj, "R-PROMJENE")
    assert nalaz.status is Status.NE_PROLAZI
    assert nalaz.poruka_kljuc == "R-PROMJENE.prisutne"
    assert nalaz.lokacija.dio is Dio.DOKUMENT
    tekst = tekst_nalaza(nalaz)
    assert "Pregled (Review)" in tekst and "Accept All Changes and Stop Tracking" in tekst
    assert _padaju(sadrzaj) == {"R-PROMJENE"}


def test_ispravan_rad_nema_promjena_ni_okvira():
    m = _model(ispravan(54, PROFIL))
    assert m.pracene_promjene == 0 and m.okviri == ()
    [nalaz] = _nalazi(ispravan(54, PROFIL), "R-PROMJENE")
    assert nalaz.status is Status.PROSLO and nalaz.poruka_kljuc == "R-PROMJENE.prolazi"
    [nalaz] = _nalazi(ispravan(54, PROFIL), "R-OKVIRI")
    assert nalaz.status is Status.PROSLO and nalaz.poruka_kljuc == "R-OKVIRI.prolazi"


def test_umetnuti_tekst_je_u_prihvacenom_stanju():
    m = _model(VARIJANTE["promjene_umetanje"](54, PROFIL))
    assert any("umetnuta uz praćenje" in o.tekst for o in m.odlomci)
    assert m.pracene_promjene == 1


def test_izbrisani_run_ne_postoji():
    m = _model(VARIJANTE["promjene_brisanje"](54, PROFIL))
    assert not any("Izbrisani tekst" in o.tekst for o in m.odlomci)
    fontovi = {r.font.vrijednost for o in m.odlomci for r in o.runovi}
    assert "Arial" not in fontovi


def test_rpr_change_ne_utjece_na_svojstva():
    m = _model(VARIJANTE["promjene_oblik"](54, PROFIL))
    [run] = [r for o in m.odlomci for r in o.runovi if "promijenjen je font" in r.tekst]
    assert run.font.vrijednost == "Calibri" and run.font.izvor.vrsta is not VrstaIzvora.IZRAVNO
    assert run.velicina.vrijednost == 24 and run.velicina.izvor.vrsta is not VrstaIzvora.IZRAVNO


def test_ppr_change_ne_utjece_na_svojstva():
    """Prije lijevo poravnanje i 0 pt razmaka (w:pPrChange), sada iz stila Normal."""

    def promjena(d):
        ppr = _osnovni(d)[0]._p.get_or_add_pPr()
        ppr.append(parse_xml(
            f'<w:pPrChange {W} w:id="1" {PROMJENA}><w:pPr><w:jc w:val="left"/>'
            '<w:spacing w:before="0" w:after="0"/><w:ind w:left="720"/></w:pPr></w:pPrChange>'
        ))

    sadrzaj = _izmijeni(promjena)
    m = _model(sadrzaj)
    o = next(o for o in m.odlomci if o.uloga is UlogaOdlomka.OSNOVNI_TEKST)
    for svojstvo in (o.poravnanje, o.razmak_ispred, o.razmak_iza, o.uvlaka_lijevo):
        assert svojstvo.izvor.vrsta is not VrstaIzvora.IZRAVNO
    assert o.razmak_ispred.vrijednost == 120
    assert _padaju(sadrzaj) == {"R-PROMJENE"}


def test_izbrisani_odlomak_ne_postoji_ni_u_indeksima_ni_u_numeraciji():
    """Cijeli numerirani naslov i njegova oznaka u w:del: odlomak ne postoji."""
    izvorni = _model(ispravan(54, PROFIL))

    def promjena(d):
        uvod = next(o for o in d.paragraphs if o.text == "Uvod")
        uvod._p.addprevious(parse_xml(
            f'<w:p {W}><w:pPr><w:pStyle w:val="Heading1"/><w:rPr><w:del w:id="1" {PROMJENA}/></w:rPr></w:pPr>'
            f'<w:del w:id="2" {PROMJENA}><w:r><w:delText>Stari naslov</w:delText></w:r></w:del></w:p>'
        ))

    sadrzaj = _izmijeni(promjena)
    m = _model(sadrzaj)
    assert [o.tekst for o in m.odlomci] == [o.tekst for o in izvorni.odlomci]
    assert [o.indeks for o in m.odlomci] == list(range(1, len(m.odlomci) + 1))
    oznake = [o.numeracija.vrijednost.oznaka for o in m.odlomci if o.uloga is UlogaOdlomka.NASLOV_POGLAVLJA]
    assert oznake == ["1.", "2.", "3."]
    assert m.pracene_promjene == 2
    assert _padaju(sadrzaj) == {"R-PROMJENE"}


def test_odlomak_s_izbrisanom_oznakom_i_tekstom_postoji():
    def promjena(d):
        ppr = _osnovni(d)[0]._p.get_or_add_pPr()
        ppr.append(parse_xml(f'<w:rPr {W}><w:del w:id="1" {PROMJENA}/></w:rPr>'))

    m = _model(_izmijeni(promjena))
    assert len(m.odlomci) == len(_model(ispravan(54, PROFIL)).odlomci)


def test_promjene_u_zaglavlju_i_podnozju_se_broje():
    def promjena(d):
        sekcija = d.sections[1]
        zaglavlje = sekcija.header
        zaglavlje.is_linked_to_previous = False
        zaglavlje.paragraphs[0]._p.append(parse_xml(
            f'<w:ins {W} w:id="1" {PROMJENA}><w:r><w:t>Zaglavlje</w:t></w:r></w:ins>'
        ))
        sekcija.footer.paragraphs[0]._p.append(parse_xml(
            f'<w:del {W} w:id="2" {PROMJENA}><w:r><w:delText>3</w:delText></w:r></w:del>'
        ))

    sadrzaj = _izmijeni(promjena)
    assert _model(sadrzaj).pracene_promjene == 2
    assert _padaju(sadrzaj) == {"R-PROMJENE"}


def test_premjestanje_se_broji_i_prihvaca():
    def promjena(d):
        _osnovni(d)[0]._p.append(parse_xml(
            f'<w:moveTo {W} w:id="1" {PROMJENA}><w:r><w:t xml:space="preserve"> Premješteno.</w:t></w:r></w:moveTo>'
        ))
        _osnovni(d)[1]._p.append(parse_xml(
            f'<w:moveFrom {W} w:id="2" {PROMJENA}><w:r><w:t xml:space="preserve"> Odavde.</w:t></w:r></w:moveFrom>'
        ))

    m = _model(_izmijeni(promjena))
    assert m.pracene_promjene == 2
    tekst = " ".join(o.tekst for o in m.odlomci)
    assert "Premješteno." in tekst and "Odavde." not in tekst


# --- okviri ---------------------------------------------------------------------------


def test_naslov_u_okviru():
    sadrzaj = VARIJANTE["okvir_naslov"](54, PROFIL)
    m = _model(sadrzaj)
    # Choice i Fallback: okvir se broji jednom.
    [okvir] = m.okviri
    assert okvir.uloga is UlogaOdlomka.NASLOV_POGLAVLJA and okvir.tekst == "Zaključak"
    assert okvir.kanonski_stil == "Heading1" and okvir.sekcija == 2
    [nalaz] = _nalazi(sadrzaj, "R-OKVIRI")
    assert nalaz.poruka_kljuc == "R-OKVIRI.naslov_u_okviru"
    assert nalaz.lokacija.dio is Dio.TIJELO and nalaz.lokacija.isjecak == "Zaključak"
    assert "Ctrl+X" in tekst_nalaza(nalaz)
    # Samo R-OKVIRI prijavljuje uvjet „u okviru“: R-NASLOVI ne javlja „nedostaje“.
    assert status_pravila(provjeri(sadrzaj, PROFIL, KONTEKST).nalazi, "R-NASLOVI") is Status.PROSLO


def test_tekst_u_okviru():
    sadrzaj = VARIJANTE["okvir_tekst"](54, PROFIL)
    [okvir] = _model(sadrzaj).okviri
    assert okvir.uloga is UlogaOdlomka.OSNOVNI_TEKST
    [nalaz] = _nalazi(sadrzaj, "R-OKVIRI")
    assert nalaz.poruka_kljuc == "R-OKVIRI.tekst_u_okviru"


def test_okvir_na_naslovnici_se_ne_provjerava():
    sadrzaj = VARIJANTE["okvir_naslovnica"](54, PROFIL)
    assert _model(sadrzaj).okviri == ()
    assert _padaju(sadrzaj) == set()


def test_sav_osnovni_tekst_u_okvirima_ne_daje_nema_teksta():
    def promjena(d):
        for o in _osnovni(d):
            p = o._p
            sidro = OxmlElement("w:p")
            p.addprevious(sidro)
            p.getparent().remove(p)
            p.append(parse_xml(
                f'<w:r {W}><w:t xml:space="preserve"> Dovoljno dug dodatak da odlomak bude osnovni tekst.</w:t></w:r>'
            ))
            sidro.append(okvir_run([p]))

    sadrzaj = _izmijeni(promjena)
    m = _model(sadrzaj)
    assert not any(o.uloga is UlogaOdlomka.OSNOVNI_TEKST for o in m.odlomci)
    assert {o.uloga for o in m.okviri} == {UlogaOdlomka.OSNOVNI_TEKST}
    assert status_pravila(provjeri(sadrzaj, PROFIL, KONTEKST).nalazi, "R-TEKST") is Status.PROSLO
    assert _padaju(sadrzaj) == {"R-OKVIRI"}
    # Isti nalaz u više okvira sažima se u jedan s brojem pojava.
    [nalaz] = _nalazi(sadrzaj, "R-OKVIRI")
    assert nalaz.pojave == len(m.okviri)


def test_kratka_oznaka_u_okviru_je_dopustena():
    def promjena(d):
        oznaka = parse_xml(f'<w:p {W}><w:r><w:t>Slika 1: karta</w:t></w:r></w:p>')
        sidro = OxmlElement("w:p")
        sidro.append(okvir_run([oznaka]))
        _osnovni(d)[0]._p.addnext(sidro)

    sadrzaj = _izmijeni(promjena)
    [okvir] = _model(sadrzaj).okviri
    assert okvir.uloga is UlogaOdlomka.OSTALO
    assert _padaju(sadrzaj) == set()


def test_okvir_u_izbrisanom_sadrzaju_ne_postoji():
    def promjena(d):
        naslov = parse_xml(f'<w:p {W}><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Uvod</w:t></w:r></w:p>')
        izbrisano = parse_xml(f'<w:del {W} w:id="1" {PROMJENA}/>')
        izbrisano.append(okvir_run([naslov]))
        _osnovni(d)[0]._p.append(izbrisano)

    sadrzaj = _izmijeni(promjena)
    assert _model(sadrzaj).okviri == ()
    assert _padaju(sadrzaj) == {"R-PROMJENE"}


def test_samo_vml_okvir_i_ugnijezdeni_okvir():
    def promjena(d):
        unutarnji = parse_xml(f'<w:p {W}><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Zaključak</w:t></w:r></w:p>')
        vanjski = parse_xml(f'<w:p {W}><w:r><w:t>Okvir</w:t></w:r></w:p>')
        vanjski.append(okvir_run([unutarnji]))
        run = okvir_run([vanjski])
        alternativa = run[0]
        alternativa.remove(alternativa[0])  # bez mc:Choice: čita se Fallback
        sidro = OxmlElement("w:p")
        sidro.append(run)
        _osnovni(d)[0]._p.addnext(sidro)

    m = _model(_izmijeni(promjena))
    assert [(o.tekst, o.uloga) for o in m.okviri] == [
        ("Okvir", UlogaOdlomka.OSTALO),
        ("Zaključak", UlogaOdlomka.NASLOV_POGLAVLJA),
    ]


@pytest.mark.parametrize(
    ("stil", "tekst", "uloga"),
    [
        ("Heading1", "Bilo što", UlogaOdlomka.NASLOV_POGLAVLJA),
        ("Normal", "1. UVOD:", UlogaOdlomka.NASLOV_POGLAVLJA),
        ("Normal", "4. LITERATURA:", UlogaOdlomka.NASLOV_LITERATURE),
        ("Normal", "SAŽETAK", UlogaOdlomka.NASLOV_SAZETKA),
        ("Normal", "x" * 40, UlogaOdlomka.OSNOVNI_TEKST),
        ("Normal", "x" * 39, UlogaOdlomka.OSTALO),
        ("Normal", "Slika 1. " + "x" * 60, UlogaOdlomka.OSTALO),
        ("Normal", "tablica 2: " + "x" * 60, UlogaOdlomka.OSTALO),
        ("Caption", "Karta studijskih programa sveučilišta u sjevernoj Hrvatskoj", UlogaOdlomka.OSTALO),
        ("Heading1", "   ", UlogaOdlomka.OSTALO),
    ],
)
def test_uloga_okvira(stil, tekst, uloga):
    assert uloga_okvira(stil, tekst, PROFIL.uloge) is uloga


# --- registracija, poruke, profil ----------------------------------------------------------


def test_registracija_nakon_broja_stranice_i_poruke():
    redoslijed = list(REGISTAR)
    assert redoslijed[-3:] == ["R-BROJ-STRANICE", "R-PROMJENE", "R-OKVIRI"]
    assert set(REGISTAR["R-PROMJENE"].poruke) == {"R-PROMJENE.prolazi", "R-PROMJENE.prisutne"}
    assert set(REGISTAR["R-OKVIRI"].poruke) == {
        "R-OKVIRI.prolazi", "R-OKVIRI.naslov_u_okviru", "R-OKVIRI.tekst_u_okviru",
    }
    for pravilo_id in ("R-PROMJENE", "R-OKVIRI"):
        assert f"pravilo.{pravilo_id}" in KATALOG
        assert all(k in KATALOG for k in REGISTAR[pravilo_id].poruke)


@pytest.mark.parametrize(("pravilo_id", "ime"), [("R-PROMJENE", "promjene_umetanje"), ("R-OKVIRI", "okvir_naslov")])
def test_pravilo_se_moze_iskljuciti(pravilo_id, ime):
    odjeljak = f"[pravila.{pravilo_id}]"
    pocetak = TOML.index(odjeljak)
    kraj = TOML.find("\n[", pocetak)
    toml = TOML[:pocetak] + f"{odjeljak}\nukljuceno = false\n" + (TOML[kraj:] if kraj != -1 else "")
    p = profil.ucitaj(toml.encode("utf-8"))
    izvjestaj = provjeri(VARIJANTE[ime](54, PROFIL), p, KONTEKST)
    assert status_pravila(izvjestaj.nalazi, pravilo_id) is None
    assert all(n.status is Status.PROSLO for n in izvjestaj.nalazi)


def test_okvir_naslova_u_samoj_sekciji_tijela_jedne_sekcije():
    """Jedina sekcija je tijelo: okvir u njoj se provjerava."""
    sadrzaj = VARIJANTE["jedna_sekcija"](54, PROFIL)

    def promjena(d):
        naslov = parse_xml(f'<w:p {W}><w:r><w:t>Zaključak</w:t></w:r></w:p>')
        sidro = OxmlElement("w:p")
        sidro.append(okvir_run([naslov]))
        d.paragraphs[0]._p.addnext(sidro)

    m = _model(_izmijeni(promjena, sadrzaj))
    assert [o.uloga for o in m.okviri] == [UlogaOdlomka.NASLOV_POGLAVLJA]
    assert m.okviri[0].sekcija == 1


# --- dopune iz reviewa ------------------------------------------------------------------


@pytest.mark.parametrize(
    ("ime", "korijen", "tip_sadrzaja", "tip_veze"),
    [
        ("footnotes", "footnotes", CT.WML_FOOTNOTES, RT.FOOTNOTES),
        ("endnotes", "endnotes", CT.WML_ENDNOTES, RT.ENDNOTES),
    ],
)
def test_promjene_u_fusnotama_i_endnotama_se_broje(ime, korijen, tip_sadrzaja, tip_veze):
    oznaka = korijen[:-1]
    xml = (
        f'<w:{korijen} {W}><w:{oznaka} w:id="1"><w:p><w:ins w:id="5" {PROMJENA}><w:r><w:t>Bilješka</w:t>'
        f"</w:r></w:ins></w:p></w:{oznaka}></w:{korijen}>"
    ).encode("utf-8")

    def promjena(d):
        dio = Part(PackURI(f"/word/{ime}.xml"), tip_sadrzaja, xml, d.part.package)
        d.part.relate_to(dio, tip_veze)

    sadrzaj = _izmijeni(promjena)
    assert _model(sadrzaj).pracene_promjene == 1
    assert _padaju(sadrzaj) == {"R-PROMJENE"}


@pytest.mark.parametrize("oznaka", ["w:tcPrChange", "w:trPrChange", "w:sectPrChange", "w:numberingChange"])
def test_ostale_promjene_svojstava_daju_prisutne(oznaka):
    def promjena(d):
        element = parse_xml(f'<{oznaka} {W} w:id="1" {PROMJENA}/>')
        if oznaka == "w:sectPrChange":
            d.element.body.find(qn("w:sectPr")).append(element)
        else:
            _osnovni(d)[0]._p.get_or_add_pPr().append(element)

    sadrzaj = _izmijeni(promjena)
    [nalaz] = _nalazi(sadrzaj, "R-PROMJENE")
    assert nalaz.poruka_kljuc == "R-PROMJENE.prisutne"
    assert _padaju(sadrzaj) == {"R-PROMJENE"}


def test_premjestena_oznaka_odlomka_ne_postoji_ni_u_numeraciji():
    """Numerirani naslov premješten odavde (w:pPr/w:rPr/w:moveFrom) bez preostalog teksta."""
    izvorni = _model(ispravan(54, PROFIL))

    def promjena(d):
        uvod = next(o for o in d.paragraphs if o.text == "Uvod")
        uvod._p.addprevious(parse_xml(
            f'<w:p {W}><w:pPr><w:pStyle w:val="Heading1"/><w:rPr><w:moveFrom w:id="1" {PROMJENA}/></w:rPr>'
            f'</w:pPr><w:moveFrom w:id="2" {PROMJENA}><w:r><w:t>Premješteni naslov</w:t></w:r></w:moveFrom></w:p>'
        ))

    m = _model(_izmijeni(promjena))
    assert [o.tekst for o in m.odlomci] == [o.tekst for o in izvorni.odlomci]
    oznake = [o.numeracija.vrijednost.oznaka for o in m.odlomci if o.uloga is UlogaOdlomka.NASLOV_POGLAVLJA]
    assert oznake == ["1.", "2.", "3."]


def test_odlomak_s_izbrisanom_oznakom_i_slikom_postoji():
    def promjena(d):
        _osnovni(d)[0]._p.addnext(parse_xml(
            f'<w:p {W}><w:pPr><w:rPr><w:del w:id="1" {PROMJENA}/></w:rPr></w:pPr><w:r><w:drawing/></w:r></w:p>'
        ))

    m = _model(_izmijeni(promjena))
    assert len(m.odlomci) == len(_model(ispravan(54, PROFIL)).odlomci) + 1


def _u_okvir(d, tekst: str) -> None:
    """Premješta odlomak tijela s danim tekstom u tekstni okvir na istom mjestu."""
    p = next(o for o in d.paragraphs if o.text == tekst)._p
    sidro = OxmlElement("w:p")
    p.addprevious(sidro)
    p.getparent().remove(p)
    sidro.append(okvir_run([p]))


@pytest.mark.parametrize(
    ("naslov", "uloga"),
    [("Sažetak", UlogaOdlomka.NASLOV_SAZETKA), ("Literatura", UlogaOdlomka.NASLOV_LITERATURE)],
)
def test_naslov_sazetka_ili_literature_u_okviru_prijavljuje_samo_r_okviri(naslov, uloga):
    sadrzaj = _izmijeni(lambda d: _u_okvir(d, naslov))
    [okvir] = _model(sadrzaj).okviri
    assert okvir.uloga is uloga
    assert _padaju(sadrzaj) == {"R-OKVIRI"}
    [nalaz] = _nalazi(sadrzaj, "R-OKVIRI")
    assert nalaz.poruka_kljuc == "R-OKVIRI.naslov_u_okviru"


def test_dug_opis_slike_u_okviru_je_dopusten():
    def promjena(d):
        opis = parse_xml(
            f'<w:p {W}><w:r><w:t>Slika 1. Karta studijskih programa geodezije i geomatike u Hrvatskoj</w:t></w:r></w:p>'
        )
        sidro = OxmlElement("w:p")
        sidro.append(okvir_run([opis]))
        _osnovni(d)[0]._p.addnext(sidro)

    sadrzaj = _izmijeni(promjena)
    [okvir] = _model(sadrzaj).okviri
    assert len(okvir.tekst) >= 40 and okvir.uloga is UlogaOdlomka.OSTALO
    assert _padaju(sadrzaj) == set()


def test_naslov_u_tablici_i_kontroli_sadrzaja_unutar_okvira():
    def promjena(d):
        naslov = f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>{{}}</w:t></w:r></w:p>'
        tablica = parse_xml(
            f'<w:tbl {W}><w:tblPr/><w:tblGrid><w:gridCol w:w="2000"/></w:tblGrid><w:tr><w:tc><w:tcPr/>'
            + naslov.format("Zaključak") + "</w:tc></w:tr></w:tbl>"
        )
        sdt = parse_xml(f'<w:sdt {W}><w:sdtPr/><w:sdtContent>' + naslov.format("Uvod") + "</w:sdtContent></w:sdt>")
        sidro = OxmlElement("w:p")
        sidro.append(okvir_run([tablica, sdt]))
        _osnovni(d)[0]._p.addnext(sidro)

    sadrzaj = _izmijeni(promjena)
    m = _model(sadrzaj)
    assert [(o.tekst, o.uloga) for o in m.okviri] == [
        ("Zaključak", UlogaOdlomka.NASLOV_POGLAVLJA),
        ("Uvod", UlogaOdlomka.NASLOV_POGLAVLJA),
    ]
    assert {n.poruka_kljuc for n in _nalazi(sadrzaj, "R-OKVIRI")} == {"R-OKVIRI.naslov_u_okviru"}
    assert _padaju(sadrzaj) == {"R-OKVIRI"}

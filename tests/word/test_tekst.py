"""R-TEKST, razrješavanje stilova i teme te uloge odlomaka (priča 2.1)."""

import io
import re
from pathlib import Path

import docx as python_docx
import pytest
from alati.generator import VARIJANTE, ispravan, tema_s_fontom, zamijeni_u_zipu
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt

from provjera.word.jezgra import model, profil
from provjera.word.jezgra.model import Poravnanje, PraviloProreda, Prored
from provjera.word.jezgra.model.stilovi import kanonski_id
from provjera.word.jezgra.model.uloge import (
    NaziviUloga,
    Uloga,
    UlogaOdlomka,
    ZnacajkeOdlomka,
    normaliziraj_naziv,
    uloge_odlomaka,
)
from provjera.word.jezgra.nalaz import KontekstProvjere, Status, VrstaIzvora
from provjera.word.jezgra.poruke import KATALOG, tekst_nalaza
from provjera.word.jezgra.pravila import REGISTAR
from provjera.word.jezgra.provjera import provjeri, status_pravila

KORIJEN = Path(__file__).resolve().parents[2]
PROFIL = profil.ucitaj((KORIJEN / "profili" / "word" / "vjezba_1.toml").read_bytes())
KONTEKST = KontekstProvjere("Ana_Horvat_GiG_OG_Vjezba_1.docx", 54, None, None)


def _izmijeni(promjena, sadrzaj: bytes | None = None) -> bytes:
    dokument = python_docx.Document(io.BytesIO(sadrzaj or ispravan(54, PROFIL)))
    promjena(dokument)
    spremnik = io.BytesIO()
    dokument.save(spremnik)
    return spremnik.getvalue()


def _tekst_nalazi(sadrzaj: bytes):
    return [n for n in provjeri(sadrzaj, PROFIL, KONTEKST).nalazi if n.pravilo_id == "R-TEKST"]


def _osnovni(m):
    return [o for o in m.odlomci if o.uloga is UlogaOdlomka.OSNOVNI_TEKST]


# --- matrica iz specifikacije ---------------------------------------------------


def test_ispravan_rad_daje_jedan_prolaz():
    [nalaz] = _tekst_nalazi(ispravan(54, PROFIL))
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.PROSLO, "R-TEKST.prolazi")


def test_aptos_iz_teme():
    nalazi = _tekst_nalazi(VARIJANTE["tekst_aptos"](54, PROFIL))
    [nalaz] = nalazi
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, "R-TEKST.font_tema")
    assert nalaz.izvor.vrsta is VrstaIzvora.TEMA
    assert nalaz.pronadeno.iznos == "Aptos"
    tekst = tekst_nalaza(nalaz)
    assert "Dizajn (Design) → Fontovi (Fonts)" in tekst and "Aptos" in tekst


def test_drugi_stil_s_istim_svojstvima_ne_prolazi():
    [nalaz] = _tekst_nalazi(VARIJANTE["tekst_drugi_stil"](54, PROFIL))
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, "R-TEKST.drugi_stil")
    assert nalaz.pronadeno.iznos == "Moj tekst"
    assert nalaz.lokacija.odlomak is not None and nalaz.lokacija.isjecak


def _hrvatski_word(sadrzaj: bytes) -> bytes:
    def preimenuj(xml: bytes) -> bytes:
        return xml.replace(b'"Heading1"', b'"Naslov1"')

    return zamijeni_u_zipu(sadrzaj, {"word/styles.xml": preimenuj, "word/document.xml": preimenuj})


def test_hrvatski_word_naslov1_je_heading1():
    sadrzaj = _hrvatski_word(ispravan(54, PROFIL))
    m = model.procitaj(sadrzaj, PROFIL.uloge)
    naslovi = [o for o in m.odlomci if o.uloga in (UlogaOdlomka.NASLOV_POGLAVLJA, UlogaOdlomka.NASLOV_LITERATURE)]
    assert len(naslovi) == 4
    assert {o.stil.kanonski_id for o in naslovi} == {"Heading1"}
    assert {o.stil.ime for o in naslovi} == {"heading 1"}
    assert status_pravila(provjeri(sadrzaj, PROFIL, KONTEKST).nalazi, "R-TEKST") is Status.PROSLO


def _bez_teme(sadrzaj: bytes) -> bytes:
    def bez_veze(xml: bytes) -> bytes:
        return re.sub(rb"<Relationship [^>]*/theme\"[^>]*/>", b"", xml)

    return zamijeni_u_zipu(sadrzaj, {
        "word/theme/theme1.xml": lambda _: None,
        "word/_rels/document.xml.rels": bez_veze,
    })


def test_bez_teme_font_je_neprovjerljiv():
    def bez_fonta(d):
        d.styles["Normal"].font.name = None

    sadrzaj = _bez_teme(_izmijeni(bez_fonta))
    nalazi = _tekst_nalazi(sadrzaj)
    assert {(n.status, n.poruka_kljuc) for n in nalazi} == {(Status.NEPROVJERLJIVO, "R-TEKST.font_nepoznat")}
    assert "{" not in tekst_nalaza(nalazi[0])


def test_bez_teme_s_fontom_u_stilu_prolazi():
    assert status_pravila(provjeri(_bez_teme(ispravan(54, PROFIL)), PROFIL, KONTEKST).nalazi, "R-TEKST") is Status.PROSLO


def _strict(sadrzaj: bytes) -> bytes:
    def u_strict(xml: bytes) -> bytes:
        return xml.replace(
            b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument",
            b"http://purl.oclc.org/ooxml/officeDocument/relationships/officeDocument",
        )

    return zamijeni_u_zipu(sadrzaj, {"_rels/.rels": u_strict})


def test_strict_ooxml_je_neprovjerljiv_s_razlogom():
    sadrzaj = _strict(ispravan(54, PROFIL))
    with pytest.raises(model.NepodrzaniDokument):
        model.procitaj(sadrzaj, PROFIL.uloge)
    izvjestaj = provjeri(sadrzaj, PROFIL, KONTEKST)
    assert izvjestaj.odbijeno is None
    assert status_pravila(izvjestaj.nalazi, "R-NAZIV") is Status.PROSLO
    for pravilo_id in ("R-SEKCIJE", "R-MARGINE", "R-TEKST"):
        [nalaz] = [n for n in izvjestaj.nalazi if n.pravilo_id == pravilo_id]
        assert nalaz.status is Status.NEPROVJERLJIVO
        assert nalaz.parametri["razlog"] == "Strict OOXML"
        assert "Strict OOXML" in tekst_nalaza(nalaz)


def test_bez_osnovnog_teksta():
    dokument = python_docx.Document()
    dokument.add_paragraph("Ana Horvat")
    dokument.add_section()
    for poglavlje in PROFIL.uloge.poglavlja:
        dokument.add_heading(poglavlje, level=1)
    spremnik = io.BytesIO()
    dokument.save(spremnik)
    [nalaz] = _tekst_nalazi(spremnik.getvalue())
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, "R-TEKST.nema_teksta")


def test_podebljana_rijec_ne_mijenja_status():
    def podebljaj(d):
        odlomak = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst"))
        odlomak.add_run(" važno").bold = True
        odlomak.add_run(" nagnuto").italic = True

    assert status_pravila(provjeri(_izmijeni(podebljaj), PROFIL, KONTEKST).nalazi, "R-TEKST") is Status.PROSLO


# --- sažimanje i varijante -----------------------------------------------------------


@pytest.mark.parametrize(
    ("varijanta", "kljuc"),
    [
        ("tekst_font", "R-TEKST.font"),
        ("tekst_velicina", "R-TEKST.velicina"),
        ("tekst_poravnanje", "R-TEKST.poravnanje"),
        ("tekst_razmak", "R-TEKST.razmak_iza"),
        ("tekst_prored", "R-TEKST.prored"),
    ],
)
def test_odstupanje_u_stilu_je_jedan_sazeti_nalaz(varijanta, kljuc):
    nalazi = _tekst_nalazi(VARIJANTE[varijanta](54, PROFIL))
    [nalaz] = nalazi
    assert nalaz.poruka_kljuc == kljuc
    broj_osnovnih = len(_osnovni(model.procitaj(VARIJANTE[varijanta](54, PROFIL), PROFIL.uloge)))
    assert nalaz.pojave == broj_osnovnih > 1
    assert nalaz.izvor.vrsta is VrstaIzvora.STIL and nalaz.izvor.stil == "Normal"
    tekst = tekst_nalaza(nalaz)
    assert "Početna (Home) → Stilovi (Styles) → Normal → Izmijeni (Modify)" in tekst
    assert "{" not in tekst


def test_prikaz_vrijednosti_u_porukama():
    tekstovi = {
        ime: tekst_nalaza(_tekst_nalazi(VARIJANTE[ime](54, PROFIL))[0])
        for ime in ("tekst_razmak", "tekst_prored", "tekst_poravnanje", "tekst_velicina")
    }
    assert "10 pt" in tekstovi["tekst_razmak"] and "6 pt" in tekstovi["tekst_razmak"]
    assert "1,15" in tekstovi["tekst_prored"] and "jednostruko" in tekstovi["tekst_prored"]
    assert "obostrano" in tekstovi["tekst_poravnanje"] and "lijevo" in tekstovi["tekst_poravnanje"]
    assert "11 pt" in tekstovi["tekst_velicina"]


def test_izravno_oblikovanje_ima_izvor_izravno():
    def izravno(d):
        odlomak = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst"))
        odlomak.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        odlomak.runs[0].font.size = Pt(14)

    nalazi = _tekst_nalazi(_izmijeni(izravno))
    po_kljucu = {n.poruka_kljuc: n for n in nalazi}
    assert set(po_kljucu) == {"R-TEKST.poravnanje", "R-TEKST.velicina"}
    assert all(n.izvor.vrsta is VrstaIzvora.IZRAVNO and n.pojave == 1 for n in nalazi)


def test_tocan_prored_se_prijavljuje_u_tockama():
    def tocno(d):
        d.styles["Normal"].paragraph_format.line_spacing = Pt(12)

    [nalaz] = _tekst_nalazi(_izmijeni(tocno))
    assert nalaz.poruka_kljuc == "R-TEKST.prored"
    assert "12 pt" in tekst_nalaza(nalaz)


def test_heading2_u_tekstu_pada_kao_drugi_stil():
    def podnaslov(d):
        odlomak = next(o for o in d.paragraphs if o.text.startswith("Drugi odlomak"))
        odlomak.insert_paragraph_before("Podnaslov", style="Heading 2")

    nalazi = _tekst_nalazi(_izmijeni(podnaslov))
    assert "R-TEKST.drugi_stil" in {n.poruka_kljuc for n in nalazi}


def test_kljucevi_r_tekst_u_katalogu():
    for kljuc in REGISTAR["R-TEKST"].poruke:
        assert kljuc in KATALOG
    assert list(REGISTAR).index("R-TEKST") == list(REGISTAR).index("R-MARGINE") + 1
    for poravnanje in Poravnanje:
        assert f"poravnanje.{poravnanje.value}" in KATALOG


# --- razrješavanje stilova -----------------------------------------------------------


def test_lanac_based_on_i_izvori():
    def lanac(d):
        osnova = d.styles.add_style("Osnova", WD_STYLE_TYPE.PARAGRAPH)
        osnova.base_style = d.styles["Normal"]
        osnova.font.size = Pt(13)
        izvedeni = d.styles.add_style("Izvedeni", WD_STYLE_TYPE.PARAGRAPH)
        izvedeni.base_style = osnova
        izvedeni.paragraph_format.space_before = Pt(3)
        odlomak = d.add_paragraph("Odlomak u izvedenom stilu.", style=izvedeni)
        odlomak.paragraph_format.space_after = Pt(9)

    m = model.procitaj(_izmijeni(lanac), PROFIL.uloge)
    o = m.odlomci[-1]
    assert (o.stil.kanonski_id, o.stil.ime) == ("Izvedeni", "Izvedeni")
    assert o.runovi[0].velicina.vrijednost == 26
    assert o.runovi[0].velicina.izvor.stil == "Osnova"
    assert (o.razmak_ispred.vrijednost, o.razmak_ispred.izvor.stil) == (60, "Izvedeni")
    assert (o.razmak_iza.vrijednost, o.razmak_iza.izvor.vrsta) == (180, VrstaIzvora.IZRAVNO)
    assert o.poravnanje.vrijednost is Poravnanje.OBOSTRANO and o.poravnanje.izvor.stil == "Normal"
    assert o.runovi[0].font.vrijednost == "Calibri" and o.runovi[0].font.izvor.stil == "Normal"
    assert o.prored.vrijednost == Prored(PraviloProreda.AUTO, 240)


def test_zadane_vrijednosti_iz_docdefaults_i_worda():
    dokument = python_docx.Document()
    dokument.add_paragraph("Ana Horvat")
    dokument.add_section()
    dokument.add_heading("Uvod", level=1)
    dokument.add_paragraph("Tekst.")
    spremnik = io.BytesIO()
    dokument.save(spremnik)
    [o] = _osnovni(model.procitaj(spremnik.getvalue(), PROFIL.uloge))
    assert (o.razmak_iza.vrijednost, o.razmak_iza.izvor.vrsta) == (200, VrstaIzvora.ZADANO)
    assert (o.razmak_ispred.vrijednost, o.razmak_ispred.izvor.vrsta) == (0, VrstaIzvora.ZADANO)
    assert o.prored.vrijednost == Prored(PraviloProreda.AUTO, 276)
    # Predložak python-docx: font tijela iz teme je Cambria.
    assert (o.runovi[0].font.vrijednost, o.runovi[0].font.izvor.vrsta) == ("Cambria", VrstaIzvora.TEMA)
    assert (o.runovi[0].velicina.vrijednost, o.runovi[0].velicina.izvor.vrsta) == (22, VrstaIzvora.ZADANO)


def test_font_teme_prati_temu():
    sadrzaj = zamijeni_u_zipu(_izmijeni(lambda d: setattr(d.styles["Normal"].font, "name", None)),
                              {"word/theme/theme1.xml": tema_s_fontom("Calibri")})
    assert status_pravila(provjeri(sadrzaj, PROFIL, KONTEKST).nalazi, "R-TEKST") is Status.PROSLO


@pytest.mark.parametrize(
    ("style_id", "ime", "prilagoden", "kanonski"),
    [
        ("Naslov1", "heading 1", False, "Heading1"),
        ("Odlomakpopisa", "List Paragraph", False, "ListParagraph"),
        ("Sadraj1", "toc 1", False, "TOC1"),
        ("TOCNaslov", "TOC Heading", False, "TOCHeading"),
        ("Opisslike", "caption", False, "Caption"),
        ("Normal", "Normal", False, "Normal"),
        ("Mojstil", "Moj stil", True, "Mojstil"),
        ("Bezimena", None, False, "Bezimena"),
    ],
)
def test_kanonski_id(style_id, ime, prilagoden, kanonski):
    assert kanonski_id(style_id, ime, prilagoden) == kanonski


def test_runovi_bez_izbrisanog_teksta_i_tekstnih_okvira():
    def dodaj(d):
        odlomak = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst"))
        odlomak._p.append(parse_xml(
            f'<w:del {nsdecls("w")} w:id="1" w:author="A"><w:r><w:delText>IZBRISANO</w:delText></w:r></w:del>'
        ))
        odlomak._p.append(parse_xml(
            f'<w:ins {nsdecls("w")} w:id="2" w:author="A"><w:r><w:t xml:space="preserve"> umetnuto</w:t></w:r></w:ins>'
        ))
        odlomak._p.append(parse_xml(
            f'<w:r {nsdecls("w")}><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox>'
            f'<w:txbxContent><w:p><w:r><w:t>OKVIR</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'
        ))

    m = model.procitaj(_izmijeni(dodaj), PROFIL.uloge)
    [o] = [o for o in m.odlomci if o.tekst.startswith("Ovo je osnovni tekst poglavlja Uvod")]
    assert o.tekst == "Ovo je osnovni tekst poglavlja Uvod. umetnuto"
    assert all("OKVIR" not in x.tekst and "IZBRISANO" not in x.tekst for x in m.odlomci)


def test_odlomci_u_tablici_i_sadrzaju():
    def dodaj(d):
        tablica = d.add_table(rows=1, cols=1)
        tablica.cell(0, 0).paragraphs[0].add_run("U tablici")
        d.element.body.append(parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
            f'</w:docPartObj></w:sdtPr><w:sdtContent><w:p><w:r><w:t>Stavka sadržaja</w:t></w:r></w:p>'
            f'</w:sdtContent></w:sdt>'
        ))
        # Tijelo mora završiti s w:sectPr.
        sect = d.element.body.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr")
        d.element.body.remove(sect)
        d.element.body.append(sect)

    m = model.procitaj(_izmijeni(dodaj), PROFIL.uloge)
    po_tekstu = {o.tekst: o for o in m.odlomci}
    assert po_tekstu["U tablici"].u_tablici and po_tekstu["U tablici"].uloga is UlogaOdlomka.TABLICA
    assert po_tekstu["Stavka sadržaja"].uloga is UlogaOdlomka.OSTALO
    assert [o.indeks for o in m.odlomci] == list(range(1, len(m.odlomci) + 1))
    assert {o.sekcija for o in m.odlomci} == {1, 2}


def test_model_ispravnog_rada_ima_uloge_i_isjecke():
    m = model.procitaj(ispravan(54, PROFIL), PROFIL.uloge)
    uloge = [o.uloga for o in m.odlomci]
    naslovnica = [o for o in m.odlomci if o.sekcija == 1]
    assert {o.uloga for o in naslovnica} == {UlogaOdlomka.NASLOVNICA}
    assert uloge[len(naslovnica):len(naslovnica) + 2] == [UlogaOdlomka.NASLOV_SAZETKA, UlogaOdlomka.SAZETAK]
    assert uloge.count(UlogaOdlomka.NASLOV_POGLAVLJA) == 3
    assert uloge.count(UlogaOdlomka.NASLOV_LITERATURE) == 1
    assert uloge[-1] is UlogaOdlomka.LITERATURA
    assert all(o.isjecak is None or len(o.isjecak) <= 60 for o in m.odlomci)


# --- uloge odlomaka ------------------------------------------------------------------

NAZIVI = PROFIL.uloge


def _z(tekst="Tekst", stil="Normal", sekcija=Uloga.TIJELO, **kwargs):
    return ZnacajkeOdlomka(uloga_sekcije=sekcija, kanonski_stil=stil, tekst=tekst, **kwargs)


def _uloga(*znacajke):
    return uloge_odlomaka(znacajke, NAZIVI)


def test_normaliziraj_naziv():
    assert normaliziraj_naziv("UVOD:") == normaliziraj_naziv("Uvod")
    assert normaliziraj_naziv("  ZAKLJUČAK. ") == "zakljucak"
    assert normaliziraj_naziv("Đurđevac") == "durdevac"
    assert normaliziraj_naziv("Studij  geodezije i geomatike u Varaždinu!") == (
        "studij geodezije i geomatike u varazdinu"
    )


@pytest.mark.parametrize(
    ("znacajke", "uloga"),
    [
        (_z(sekcija=Uloga.NASLOVNICA, stil="Heading1", tekst="Uvod"), UlogaOdlomka.NASLOVNICA),
        (_z(u_tablici=True, stil="Heading1"), UlogaOdlomka.TABLICA),
        (_z(tekst="  "), UlogaOdlomka.PRAZNO),
        (_z(stil="TOC1", tekst="1. Uvod 1"), UlogaOdlomka.OSTALO),
        (_z(stil="TOCHeading", tekst="Sadržaj"), UlogaOdlomka.OSTALO),
        (_z(u_sadrzaju=True, tekst="Uvod"), UlogaOdlomka.OSTALO),
        (_z(stil="Caption", tekst="Slika 1."), UlogaOdlomka.OSTALO),
        (_z(tekst="", ima_sliku=True), UlogaOdlomka.OSTALO),
        (_z(stil="Heading1", tekst="Bilo što"), UlogaOdlomka.NASLOV_POGLAVLJA),
        (_z(tekst="1. UVOD:"), UlogaOdlomka.NASLOV_POGLAVLJA),
        (_z(tekst="4. Literatura"), UlogaOdlomka.NASLOV_LITERATURE),
        (_z(tekst="SAŽETAK:"), UlogaOdlomka.NASLOV_SAZETKA),
    ],
)
def test_uloga_po_prednosti(znacajke, uloga):
    assert _uloga(znacajke) == (uloga,)


def test_uloge_po_kontekstu():
    uloge = _uloga(
        _z("Ana Horvat", sekcija=Uloga.NASLOVNICA),
        _z("Sažetak"),
        _z("Sažetak rada."),
        _z("Uvod", stil="Heading1"),
        _z("Osnovni tekst."),
        _z("Stavka", numeriran=True),
        _z("Stavka", stil="ListParagraph"),
        _z("Literatura", stil="Heading1"),
        _z("Izvor 1", stil="ListParagraph", numeriran=True),
        _z("Izvor 2"),
        _z("Zaključak", stil="Heading1"),
        _z("Opet tekst."),
    )
    assert uloge == (
        UlogaOdlomka.NASLOVNICA,
        UlogaOdlomka.NASLOV_SAZETKA,
        UlogaOdlomka.SAZETAK,
        UlogaOdlomka.NASLOV_POGLAVLJA,
        UlogaOdlomka.OSNOVNI_TEKST,
        UlogaOdlomka.OSTALO,
        UlogaOdlomka.OSTALO,
        UlogaOdlomka.NASLOV_LITERATURE,
        UlogaOdlomka.LITERATURA,
        UlogaOdlomka.LITERATURA,
        UlogaOdlomka.NASLOV_POGLAVLJA,
        UlogaOdlomka.OSNOVNI_TEKST,
    )


def test_nazivi_uloga_su_strogi():
    with pytest.raises(ValueError):
        NaziviUloga(sazetak="Sažetak", literatura="Literatura", poglavlja=("Uvod",), nepoznato=1)
    with pytest.raises(ValueError):
        NaziviUloga(sazetak="", literatura="Literatura", poglavlja=("Uvod",))


def test_sazetak_u_heading1_je_naslov_sazetka():
    uloge = _uloga(
        _z("Sažetak", stil="Heading1"),
        _z("Sažetak rada."),
        _z("Uvod", stil="Heading1"),
        _z("Osnovni tekst."),
    )
    assert uloge == (
        UlogaOdlomka.NASLOV_SAZETKA,
        UlogaOdlomka.SAZETAK,
        UlogaOdlomka.NASLOV_POGLAVLJA,
        UlogaOdlomka.OSNOVNI_TEKST,
    )


def _numeriraj(odlomak, num_id: str) -> None:
    odlomak._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
    ))


def test_numeracija_odlomka_kroz_model():
    def numeracija(d):
        numeriran = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst poglavlja Uvod"))
        _numeriraj(numeriran, "1")
        ponisten = next(o for o in d.paragraphs if o.text.startswith("Drugi odlomak"))
        ponisten.style = d.styles["List Number"]
        _numeriraj(ponisten, "0")

    m = model.procitaj(_izmijeni(numeracija), PROFIL.uloge)
    numeriran = next(o for o in m.odlomci if o.tekst.startswith("Ovo je osnovni tekst poglavlja Uvod"))
    ponisten = next(o for o in m.odlomci if o.tekst.startswith("Drugi odlomak"))
    assert numeriran.uloga is UlogaOdlomka.OSTALO
    assert (ponisten.stil.kanonski_id, ponisten.uloga) == ("ListNumber", UlogaOdlomka.OSNOVNI_TEKST)


def test_stil_znaka_s_drugim_fontom():
    def stil_znaka(d):
        stil = d.styles.add_style("Moj znak", WD_STYLE_TYPE.CHARACTER)
        stil.font.name = "Arial"
        odlomak = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst"))
        odlomak.runs[0].style = stil

    [nalaz] = _tekst_nalazi(_izmijeni(stil_znaka))
    assert (nalaz.poruka_kljuc, nalaz.pronadeno.iznos) == ("R-TEKST.font", "Arial")
    assert (nalaz.izvor.vrsta, nalaz.izvor.stil) == (VrstaIzvora.STIL, "Moj znak")


def test_stil_iz_profila_bez_obzira_na_velika_slova():
    malim = PROFIL.model_copy(update={"pravila": {
        **PROFIL.pravila, "R-TEKST": PROFIL.pravila["R-TEKST"].model_copy(update={"stil": "normal"}),
    }})
    nalazi = [n for n in provjeri(ispravan(54, PROFIL), malim, KONTEKST).nalazi if n.pravilo_id == "R-TEKST"]
    assert [n.status for n in nalazi] == [Status.PROSLO]
    assert "u stilu normal" in tekst_nalaza(nalazi[0])


def test_kriv_razmak_ispred():
    def razmak(d):
        d.styles["Normal"].paragraph_format.space_before = Pt(12)

    [nalaz] = _tekst_nalazi(_izmijeni(razmak))
    assert nalaz.poruka_kljuc == "R-TEKST.razmak_ispred"
    assert (nalaz.ocekivano.iznos, nalaz.pronadeno.iznos) == (120, 240)
    tekst = tekst_nalaza(nalaz)
    assert "12 pt" in tekst and "6 pt" in tekst


def test_tema_ima_prednost_pred_imenom_fonta_u_istom_rfonts():
    def tema_i_ime(d):
        rfonts = d.styles["Normal"].element.rPr.rFonts
        rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}asciiTheme", "minorHAnsi")
        assert rfonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii") == "Calibri"

    sadrzaj = zamijeni_u_zipu(_izmijeni(tema_i_ime), {"word/theme/theme1.xml": tema_s_fontom("Aptos")})
    [nalaz] = _tekst_nalazi(sadrzaj)
    assert (nalaz.poruka_kljuc, nalaz.pronadeno.iznos) == ("R-TEKST.font_tema", "Aptos")


def test_stil_iz_profila_s_razmakom_odgovara_kanonskom_id():
    def body_text(d):
        for o in d.paragraphs:
            if o.style.name == "Normal" and o.text.startswith(("Ovo je", "Drugi")):
                o.style = d.styles["Body Text"]

    tekst = PROFIL.pravila["R-TEKST"].model_copy(update={"stil": "Body Text"})
    body = PROFIL.model_copy(update={"pravila": {**PROFIL.pravila, "R-TEKST": tekst}})
    nalazi = [n for n in provjeri(_izmijeni(body_text), body, KONTEKST).nalazi if n.pravilo_id == "R-TEKST"]
    assert "R-TEKST.drugi_stil" not in {n.poruka_kljuc for n in nalazi}


def _odlomak_uvoda(m):
    return next(o for o in m.odlomci if o.tekst.startswith("Ovo je osnovni tekst poglavlja Uvod"))


def _razmak_u_ppr(atributi: str):
    def promjena(d):
        odlomak = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst poglavlja Uvod"))
        odlomak._p.get_or_add_pPr().append(parse_xml(f"<w:spacing {nsdecls('w')} {atributi}/>"))

    return promjena


def test_automatski_razmak_ima_prednost_i_iznosi_14_pt():
    m = model.procitaj(_izmijeni(_razmak_u_ppr('w:beforeAutospacing="1" w:before="100"')), PROFIL.uloge)
    assert _odlomak_uvoda(m).razmak_ispred.vrijednost == 280
    nalazi = provjeri(_izmijeni(_razmak_u_ppr('w:beforeAutospacing="1" w:before="100"')), PROFIL, KONTEKST)
    tekstovi = [tekst_nalaza(n) for n in nalazi.nalazi if n.poruka_kljuc == "R-TEKST.razmak_ispred"]
    assert tekstovi and "14 pt" in tekstovi[0]


def test_decimalni_atribut_ne_rusi_model():
    m = model.procitaj(_izmijeni(_razmak_u_ppr('w:before="120.0"')), PROFIL.uloge)
    assert _odlomak_uvoda(m).razmak_ispred.vrijednost == 120


def test_mjera_s_jedinicom_ne_rusi_model_ni_sekcije():
    sadrzaj = _izmijeni(_razmak_u_ppr('w:before="12pt"'))
    izvjestaj = provjeri(sadrzaj, PROFIL, KONTEKST)
    assert status_pravila(izvjestaj.nalazi, "R-SEKCIJE") is Status.PROSLO
    assert status_pravila(izvjestaj.nalazi, "R-MARGINE") is Status.PROSLO


def test_dokument_bez_dijela_stilova_koristi_wordove_zadane_vrijednosti():
    def bez_veze(xml: bytes) -> bytes:
        return re.sub(rb"<Relationship [^>]*/styles\"[^>]*/>", b"", xml)

    sadrzaj = zamijeni_u_zipu(ispravan(54, PROFIL), {
        "word/styles.xml": lambda _: None,
        "word/_rels/document.xml.rels": bez_veze,
    })
    m = model.procitaj(sadrzaj, PROFIL.uloge)
    o = _odlomak_uvoda(m)
    assert o.stil.kanonski_id == "Normal"
    assert o.poravnanje.vrijednost is Poravnanje.LIJEVO
    assert o.poravnanje.izvor.vrsta is VrstaIzvora.ZADANO
    assert o.razmak_iza.vrijednost == 0


@pytest.mark.parametrize("num_id", ["77"])
def test_viseci_num_id_nije_popis(num_id):
    def numeracija(d):
        odlomak = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst poglavlja Uvod"))
        _numeriraj(odlomak, num_id)

    o = _odlomak_uvoda(model.procitaj(_izmijeni(numeracija), PROFIL.uloge))
    assert o.numeracija.vrijednost is None
    assert o.uloga is UlogaOdlomka.OSNOVNI_TEKST


def test_isjecak_bez_tabulatora_i_prijeloma():
    def promjena(d):
        odlomak = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst poglavlja Uvod"))
        odlomak.runs[0].text = "Prvi\tdrugi\nTreći"

    o = next(o for o in model.procitaj(_izmijeni(promjena), PROFIL.uloge).odlomci if o.tekst.startswith("Prvi"))
    assert o.isjecak.startswith("Prvi drugi Treći")

"""Zapis zbirne tablice u Excel (priča 4.3, AD-12): tekst nikad nije formula, nema prepisivanja."""

import io
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_zbirna import PROFIL, VRIJEME, _rezultat

from provjera.word.adapteri import excel_izlaz
from provjera.word.adapteri.popis import ZONA
from provjera.word.jezgra.zbirna import List, ZbirnaTablica, izradi


def _knjiga():
    return load_workbook(io.BytesIO(excel_izlaz.izradi(izradi(_rezultat(), PROFIL, VRIJEME))))


def test_excel_bez_formula_i_tekst_ostaje_tekst():
    knjiga = _knjiga()
    assert knjiga.sheetnames == ["Rezultati", "Nalazi", "Nepovezano"]
    for list_ in knjiga.worksheets:
        for redak in list_.iter_rows():
            for celija in redak:
                assert celija.data_type != "f", (list_.title, celija.coordinate)
    rezultati = knjiga["Rezultati"]
    assert rezultati["E4"].value == "=1+1.docx" and rezultati["E4"].data_type == "s"
    assert rezultati["A4"].value == 1
    assert knjiga["Nepovezano"]["A4"].value == "-rad.docx"
    assert any("@SUM(A1)" in str(c.value) for c in knjiga["Nalazi"]["D"])
    assert rezultati.freeze_panes == "A4"
    assert [c.value for c in rezultati[1]][:6] == [
        "Vježba", "Vježba 1", "Verzija profila", "abcd1234", "Vrijeme provjere", VRIJEME,
    ]


def test_formula_ne_izvrsava_ni_u_sirovom_xml():
    import zipfile

    with zipfile.ZipFile(io.BytesIO(excel_izlaz.izradi(izradi(_rezultat(), PROFIL, VRIJEME)))) as z:
        for ime in z.namelist():
            if ime.startswith("xl/worksheets/"):
                assert b"<f>" not in z.read(ime)


T = datetime(2026, 10, 15, 12, 30, 5, tzinfo=ZONA)


def test_spremi_prvi_put_i_postojeca(tmp_path):
    prva = excel_izlaz.spremi(b"prvi", tmp_path, "vjezba_1", T)
    assert prva == tmp_path / "rezultati_vjezba_1.xlsx"
    druga = excel_izlaz.spremi(b"drugi", tmp_path, "vjezba_1", T)
    assert druga == tmp_path / "rezultati_vjezba_1_20261015-123005.xlsx"
    treca = excel_izlaz.spremi(b"treci", tmp_path, "vjezba_1", T)
    assert treca == tmp_path / "rezultati_vjezba_1_20261015-123005_2.xlsx"
    assert prva.read_bytes() == b"prvi" and druga.read_bytes() == b"drugi"


def test_spremi_zakljucana(tmp_path, monkeypatch):
    pravi = open

    def otvori(putanja, nacin="r", *a, **k):
        if Path(putanja).name == "rezultati_vjezba_1.xlsx":
            raise PermissionError("otvoreno u Excelu")
        return pravi(putanja, nacin, *a, **k)

    monkeypatch.setattr("builtins.open", otvori)
    putanja = excel_izlaz.spremi(b"x", tmp_path, "vjezba_1", T)
    assert putanja.name == "rezultati_vjezba_1_20261015-123005.xlsx"


def test_spremi_greska_pisanja_brise_nedovrsenu(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_izlaz.os, "fsync", lambda _f: (_ for _ in ()).throw(OSError("disk")))
    with pytest.raises(OSError):
        excel_izlaz.spremi(b"x", tmp_path, "vjezba_1", T)
    assert list(tmp_path.iterdir()) == []


def _jedan_list(*vrijednosti):
    return ZbirnaTablica((List("Nalazi", ("Vježba", "x"), ("Opis",), tuple((v,) for v in vrijednosti)),))


@pytest.mark.parametrize("vrijednost", ["=1+1", "+1", "-1", "@SUM(A1)", "\t=1", "\r=1", "=cmd|' /C calc'!A0"])
def test_svaki_prefiks_formule_ostaje_tekst(vrijednost):
    sadrzaj = excel_izlaz.izradi(_jedan_list(vrijednost))
    celija = load_workbook(io.BytesIO(sadrzaj))["Nalazi"]["A4"]
    assert celija.data_type == "s"
    assert celija.value == vrijednost


def test_kontrolni_znakovi_se_uklanjaju():
    sadrzaj = excel_izlaz.izradi(_jedan_list("ra\x00d\x07.docx"))
    assert load_workbook(io.BytesIO(sadrzaj))["Nalazi"]["A4"].value == "rad.docx"


def test_spremi_nezapisiva_mapa_javlja_pravi_razlog(tmp_path, monkeypatch):
    def otvori(putanja, nacin="r", *a, **k):
        raise PermissionError(13, "Pristup odbijen")

    monkeypatch.setattr("builtins.open", otvori)
    with pytest.raises(PermissionError) as greska:
        excel_izlaz.spremi(b"x", tmp_path, "vjezba_1", T)
    assert greska.value.strerror == "Pristup odbijen"


def test_spremi_bez_slobodnog_imena_ima_razlog(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_izlaz, "NAJVISE_POKUSAJA", 2)
    for _ in range(2):
        excel_izlaz.spremi(b"x", tmp_path, "vjezba_1", T)
    with pytest.raises(FileExistsError) as greska:
        excel_izlaz.spremi(b"x", tmp_path, "vjezba_1", T)
    assert greska.value.strerror


def test_tekst_s_prefiksom_formule_ostaje_tekst_i_nakon_uredivanja():
    knjiga = load_workbook(io.BytesIO(excel_izlaz.izradi(_jedan_list("=1+1", "obican"))))
    list_ = knjiga["Nalazi"]
    assert list_["A4"].quotePrefix is True
    assert not list_["A5"].quotePrefix


def test_filtar_na_retku_naslova():
    knjiga = _knjiga()
    rezultati = knjiga["Rezultati"]
    assert rezultati.auto_filter.ref.startswith("A3:")
    assert rezultati.auto_filter.ref.endswith(str(3 + rezultati.max_row - 3))


def test_predug_tekst_se_skracuje_na_granicu_excela():
    knjiga = load_workbook(io.BytesIO(excel_izlaz.izradi(_jedan_list("x" * 40_000))))
    assert len(knjiga["Nalazi"]["A4"].value) == excel_izlaz.NAJVECA_DULJINA_CELIJE

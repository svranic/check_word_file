"""R-SAZETAK i R-LITERATURA te zajednička provjera runova (priča 2.3)."""

import io
import re
from pathlib import Path

import docx as python_docx
import pytest
from alati.generator import VARIJANTE, ispravan, tema_s_fontom, zamijeni_u_zipu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Twips

from provjera.word.jezgra import profil
from provjera.word.jezgra.nalaz import KontekstProvjere, Status, VrstaIzvora
from provjera.word.jezgra.poruke import KATALOG, tekst_nalaza
from provjera.word.jezgra.pravila import REGISTAR
from provjera.word.jezgra.provjera import provjeri

KORIJEN = Path(__file__).resolve().parents[2]
TOML = (KORIJEN / "profili" / "word" / "vjezba_1.toml").read_text(encoding="utf-8")
PROFIL = profil.ucitaj(TOML.encode("utf-8"))
KONTEKST = KontekstProvjere("Ana_Horvat_GiG_OG_Vjezba_1.docx", 54, None, None)
SAZETAK = "R-SAZETAK"
LITERATURA = "R-LITERATURA"


def _izmijeni(promjena, sadrzaj: bytes | None = None) -> bytes:
    dokument = python_docx.Document(io.BytesIO(sadrzaj or ispravan(54, PROFIL)))
    promjena(dokument)
    spremnik = io.BytesIO()
    dokument.save(spremnik)
    return spremnik.getvalue()


def _nalazi(sadrzaj: bytes, pravilo_id: str, p=PROFIL):
    return [n for n in provjeri(sadrzaj, p, KONTEKST).nalazi if n.pravilo_id == pravilo_id]


def _kljucevi(sadrzaj: bytes, pravilo_id: str) -> set[str]:
    return {n.poruka_kljuc.removeprefix(f"{pravilo_id}.") for n in _nalazi(sadrzaj, pravilo_id)}


def _tekst_sazetka(d):
    odlomci = d.paragraphs
    naslov = next(i for i, o in enumerate(odlomci) if o.text == PROFIL.uloge.sazetak)
    return odlomci[naslov + 1]


def _stavka(d):
    return next(o for o in d.paragraphs if o.style.name == "List Paragraph")


# --- matrica iz specifikacije ---------------------------------------------------


def test_ispravan_rad_prolazi_oba_pravila():
    sadrzaj = ispravan(54, PROFIL)
    for pravilo_id in (SAZETAK, LITERATURA):
        [nalaz] = _nalazi(sadrzaj, pravilo_id)
        assert (nalaz.status, nalaz.poruka_kljuc) == (Status.PROSLO, f"{pravilo_id}.prolazi")


def test_bez_uvlake_obje_strane():
    nalazi = _nalazi(VARIJANTE["sazetak_bez_uvlake"](54, PROFIL), SAZETAK)
    assert {n.poruka_kljuc for n in nalazi} == {"R-SAZETAK.uvlaka_lijevo", "R-SAZETAK.uvlaka_desno"}
    assert all(n.status is Status.NE_PROLAZI for n in nalazi)
    lijevo = next(n for n in nalazi if n.poruka_kljuc == "R-SAZETAK.uvlaka_lijevo")
    tekst = tekst_nalaza(lijevo)
    assert "3,5 cm" in tekst and "0 cm" in tekst
    assert "Početna (Home) → Odlomak (Paragraph) → Uvlaka (Indentation) → Lijevo (Left)" in tekst
    assert "vlastiti stil" in tekst and "{" not in tekst


def test_uvlaka_samo_lijevo():
    sadrzaj = VARIJANTE["sazetak_jedna_strana"](54, PROFIL)
    assert _kljucevi(sadrzaj, SAZETAK) == {"uvlaka_desno"}
    [nalaz] = _nalazi(sadrzaj, SAZETAK)
    assert "Desno (Right)" in tekst_nalaza(nalaz)


def test_vlastiti_stil_prolazi():
    sadrzaj = VARIJANTE["sazetak_vlastiti_stil"](54, PROFIL)
    assert _kljucevi(sadrzaj, SAZETAK) == {"prolazi"}


@pytest.mark.parametrize(("twips", "prolazi"), [(1984, True), (1985, True), (2001, True), (2002, False), (1967, True), (1966, False)])
def test_uvlaka_unutar_tolerancije(twips, prolazi):
    def uvuci(d):
        oblik = _tekst_sazetka(d).paragraph_format
        oblik.left_indent = oblik.right_indent = Twips(twips)

    kljucevi = _kljucevi(_izmijeni(uvuci), SAZETAK)
    assert kljucevi == ({"prolazi"} if prolazi else {"uvlaka_lijevo", "uvlaka_desno"})


def test_nema_sazetka():
    [nalaz] = _nalazi(VARIJANTE["sazetak_nedostaje"](54, PROFIL), SAZETAK)
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, "R-SAZETAK.nema_sazetka")
    assert "„Sažetak“" in tekst_nalaza(nalaz)


def test_literatura_u_drugom_fontu():
    sadrzaj = VARIJANTE["literatura_font"](54, PROFIL)
    [nalaz] = _nalazi(sadrzaj, LITERATURA)
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, "R-LITERATURA.font")
    assert nalaz.pronadeno.iznos == "Times New Roman" and nalaz.ocekivano.iznos == "Calibri"
    assert nalaz.izvor.vrsta is VrstaIzvora.STIL
    assert "Početna (Home)" in tekst_nalaza(nalaz)


def test_literatura_u_drugoj_velicini():
    def manja(d):
        d.styles["List Paragraph"].font.size = Pt(11)

    [nalaz] = _nalazi(_izmijeni(manja), LITERATURA)
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, "R-LITERATURA.velicina")
    assert (nalaz.pronadeno.iznos, nalaz.ocekivano.iznos) == (22, 24)
    tekst = tekst_nalaza(nalaz)
    assert "11 pt" in tekst and "12 pt" in tekst and "{" not in tekst


def test_literatura_bez_stavki():
    def obrisi(d):
        stavka = _stavka(d)
        stavka._p.getparent().remove(stavka._p)

    [nalaz] = _nalazi(_izmijeni(obrisi), LITERATURA)
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, "R-LITERATURA.nema_stavki")
    assert "„Literatura“" in tekst_nalaza(nalaz)


# --- ostali slučajevi -------------------------------------------------------------


def test_tekst_bez_naslova_sazetka():
    nalazi = _nalazi(_izmijeni(_bez_naslova_sazetka), SAZETAK)
    assert [(n.status, n.poruka_kljuc) for n in nalazi] == [(Status.NE_PROLAZI, "R-SAZETAK.nema_naslova")]
    assert "„Sažetak“" in tekst_nalaza(nalazi[0])


def _bez_naslova_sazetka(d):
    naslov = next(o for o in d.paragraphs if o.text == PROFIL.uloge.sazetak)
    naslov._p.getparent().remove(naslov._p)


def test_jedna_sekcija_bez_naslova_ne_ocjenjuje_naslovnicu():
    """Bez naslova su i redci naslovnice ispred poglavlja SAZETAK: provjerava se samo naslov."""
    sadrzaj = _izmijeni(_bez_naslova_sazetka, VARIJANTE["jedna_sekcija"](54, PROFIL))
    nalazi = _nalazi(sadrzaj, SAZETAK)
    assert [(n.status, n.poruka_kljuc) for n in nalazi] == [(Status.NE_PROLAZI, "R-SAZETAK.nema_naslova")]


def test_naslov_bez_teksta_sazetka():
    def obrisi_tekst(d):
        tekst = _tekst_sazetka(d)
        tekst._p.getparent().remove(tekst._p)

    [nalaz] = _nalazi(_izmijeni(obrisi_tekst), SAZETAK)
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, "R-SAZETAK.prazan_sazetak")
    tekst = tekst_nalaza(nalaz)
    assert "Ispod naslova „Sažetak“ nema teksta sažetka" in tekst and "{" not in tekst


def test_sazetak_iza_prvog_poglavlja():
    def premjesti(d):
        naslov = next(o for o in d.paragraphs if o.text == PROFIL.uloge.sazetak)
        uvod = next(o for o in d.paragraphs if o.style.name == "Heading 1")
        uvod._p.addnext(naslov._p)

    [nalaz] = _nalazi(_izmijeni(premjesti), SAZETAK)
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, "R-SAZETAK.krivo_mjesto")
    assert "Sažetak mora biti ispred prvog poglavlja" in tekst_nalaza(nalaz)


def test_naslovnica_bez_sekcije_nije_sazetak():
    """U jednoj sekciji tekst naslovnice ispred naslova „Sažetak“ ne ocjenjuje se kao sažetak."""
    assert _kljucevi(VARIJANTE["jedna_sekcija"](54, PROFIL), SAZETAK) == {"prolazi"}


def test_font_velicina_i_poravnanje_sazetka():
    def pokvari(d):
        odlomak = _tekst_sazetka(d)
        odlomak.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in odlomak.runs:
            run.font.name = "Arial"
            run.font.size = Pt(11)

    nalazi = _nalazi(_izmijeni(pokvari), SAZETAK)
    assert {n.poruka_kljuc for n in nalazi} == {"R-SAZETAK.font", "R-SAZETAK.velicina", "R-SAZETAK.poravnanje"}
    poravnanje = next(n for n in nalazi if n.poruka_kljuc == "R-SAZETAK.poravnanje")
    assert tekst_nalaza(poravnanje).endswith("Poravnanje (Alignment) → obostrano (Justify).")
    assert all(n.izvor.vrsta is VrstaIzvora.IZRAVNO for n in nalazi)
    for nalaz in nalazi:
        assert "{" not in tekst_nalaza(nalaz)


def test_ponavljanje_po_odlomcima_se_sazima():
    def jos_odlomaka(d):
        prvi = _tekst_sazetka(d)
        for _ in range(2):
            prvi.insert_paragraph_before("Još jedan odlomak sažetka bez uvlake.")

    nalazi = _nalazi(_izmijeni(jos_odlomaka), SAZETAK)
    assert sorted((n.poruka_kljuc, n.pojave) for n in nalazi) == [
        ("R-SAZETAK.uvlaka_desno", 2), ("R-SAZETAK.uvlaka_lijevo", 2)]


def test_font_iz_teme():
    def bez_fonta(d):
        d.styles["Normal"].font.name = None
        d.styles["List Paragraph"].font.name = None

    sadrzaj = zamijeni_u_zipu(_izmijeni(bez_fonta), {"word/theme/theme1.xml": tema_s_fontom("Aptos")})
    for pravilo_id in (SAZETAK, LITERATURA):
        [nalaz] = _nalazi(sadrzaj, pravilo_id)
        assert nalaz.poruka_kljuc == f"{pravilo_id}.font_tema"
        assert "Dizajn (Design) → Fontovi (Fonts)" in tekst_nalaza(nalaz)


def test_bez_teme_font_je_neprovjerljiv():
    def bez_fonta(d):
        d.styles["Normal"].font.name = None

    sadrzaj = zamijeni_u_zipu(_izmijeni(bez_fonta), {
        "word/theme/theme1.xml": lambda _: None,
        "word/_rels/document.xml.rels": lambda xml: re.sub(rb"<Relationship [^>]*/theme\"[^>]*/>", b"", xml),
    })
    for pravilo_id in (SAZETAK, LITERATURA):
        nalazi = _nalazi(sadrzaj, pravilo_id)
        assert {(n.status, n.poruka_kljuc) for n in nalazi} == {(Status.NEPROVJERLJIVO, f"{pravilo_id}.font_nepoznat")}


# --- profil, registar, poruke ------------------------------------------------------


def test_ciljevi_iz_profila():
    s, lit = PROFIL.pravila[SAZETAK], PROFIL.pravila[LITERATURA]
    assert (s.font, s.velicina, s.poravnanje.value) == ("Calibri", 24, "obostrano")
    assert (s.uvlaka_lijevo, s.uvlaka_desno, s.tolerancija) == (1984, 1984, 17)
    assert (lit.font, lit.velicina) == ("Calibri", 24)


def _odjeljak(pravilo_id: str, sadrzaj: str) -> str:
    pocetak = TOML.index(f"[pravila.{pravilo_id}]")
    kraj = TOML.find("\n[", pocetak + 1)
    return TOML[:pocetak] + f"[pravila.{pravilo_id}]\n{sadrzaj}\n" + (TOML[kraj + 1:] if kraj != -1 else "")


@pytest.mark.parametrize("pravilo_id", [SAZETAK, LITERATURA])
def test_iskljuceno_pravilo_ne_treba_ciljeve_i_nema_nalaza(pravilo_id):
    iskljucen = profil.ucitaj(_odjeljak(pravilo_id, "ukljuceno = false").encode("utf-8"))
    pada = {SAZETAK: "sazetak_nedostaje", LITERATURA: "literatura_font"}[pravilo_id]
    assert _nalazi(VARIJANTE[pada](54, PROFIL), pravilo_id) != []
    assert _nalazi(VARIJANTE[pada](54, PROFIL), pravilo_id, iskljucen) == []
    with pytest.raises(profil.GreskaProfila):
        profil.ucitaj(_odjeljak(pravilo_id, "ukljuceno = true").encode("utf-8"))


def test_registracija_nakon_naslova():
    redoslijed = list(REGISTAR)
    i = redoslijed.index("R-NASLOVI")
    assert redoslijed[i + 1:i + 3] == [SAZETAK, LITERATURA]
    assert "runovi" not in " ".join(redoslijed).lower()


def test_svi_kljucevi_imaju_poruku():
    for pravilo_id in (SAZETAK, LITERATURA):
        assert f"pravilo.{pravilo_id}" in KATALOG
        assert all(k in KATALOG for k in REGISTAR[pravilo_id].poruke)


def test_poravnanje_ostalo_nije_dopusteno_u_sazetku():
    pocetak = TOML.index("[pravila.R-SAZETAK]")
    kraj = TOML.find("\n[", pocetak + 1)
    odjeljak = TOML[pocetak:kraj].replace('poravnanje = "obostrano"', 'poravnanje = "ostalo"')
    assert 'poravnanje = "ostalo"' in odjeljak
    with pytest.raises(profil.GreskaProfila):
        profil.ucitaj((TOML[:pocetak] + odjeljak + TOML[kraj:]).encode("utf-8"))

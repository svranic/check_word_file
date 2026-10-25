"""ISVU popis studenata (priča 3.1): adapter ``provjera.word.adapteri.popis``."""

import csv
import errno
import io
import json
import logging
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
import pytest
from alati.isvu import (
    ISPIS,
    ISPISAN,
    ISPRAVNI,
    KASNI_UPIS,
    PONOVNI_UPIS,
    PRENUMERIRANI,
    S_GRESKAMA,
    UPISAN,
    ZAGLAVLJE,
    fixturei,
    izvoz,
)

from provjera.word.adapteri.popis import (
    GreskaIzdanog,
    GreskaRetka,
    GreskaZamrznutog,
    Izdano,
    PopisVecZamrznut,
    Razlike,
    Student,
    UcitaniPopis,
    primijeni,
    procitaj_izdano,
    procitaj_zamrznuti,
    putanja_izdanog,
    putanja_pdf,
    spremi,
    ucitaj_isvu,
    usporedi,
    zamrzni,
    zapisi_izdano,
)
from provjera.word.jezgra import formula as formula_modul
from provjera.word.jezgra import profil as profil_modul
from provjera.word.jezgra.poruke import KATALOG, poruka

FIXTURES = Path(__file__).resolve().parent / "fixtures"
PRVI_REDAK = 17  # 14 metapodataka, prazan redak, zaglavlje u 16. retku


def _kljucevi(popis: UcitaniPopis) -> list[tuple[int | None, str]]:
    return [(g.redak, g.poruka_kljuc) for g in popis.greske]


def test_fixturei_su_azurni():
    for naziv, sadrzaj in fixturei().items():
        # Sadržaj .xlsx-a nije bajtno stabilan (vrijeme u metapodacima), pa se uspoređuje rezultat.
        assert ucitaj_isvu((FIXTURES / naziv).read_bytes()) == ucitaj_isvu(sadrzaj)


def test_stvarni_oblik_bez_gresaka():
    popis = ucitaj_isvu((FIXTURES / "isvu_ispravan.xlsx").read_bytes())
    assert popis.greske == ()
    assert popis.moze_zamrznuti
    assert [s.rbr for s in popis.studenti] == [1, 2, 3, 4, 5]
    assert popis.studenti[0] == Student(1, "0036123456", "Ana Marija", "Horvat-Kovač", True)
    assert popis.broj_neaktivnih == 1


def test_zaglavlje_bez_metapodataka_i_drugaciji_redoslijed():
    zaglavlje = ("Status upisa", " prezime ", "IME", "jmbag", "RBR.")
    sadrzaj = izvoz(ISPRAVNI[:1], metapodaci=False, zaglavlje=zaglavlje)
    popis = ucitaj_isvu(sadrzaj)
    assert popis.greske == ()
    assert popis.studenti[0].jmbag == "0036123456"


@pytest.mark.parametrize(
    "jmbag", ["0036123456", 36123456, 36123456.0, " 0036123456 ", "36123456", "0036 123 456"]
)
def test_jmbag_zadrzava_vodece_nule(jmbag):
    popis = ucitaj_isvu(izvoz([("1.", jmbag, "Ana", "Horvat", UPISAN)]))
    assert popis.greske == ()
    assert popis.studenti[0].jmbag == "0036123456"


@pytest.mark.parametrize("rbr", ["78.", "78", " 078 ", 78, 78.0])
def test_rbr_tekst_i_broj(rbr):
    popis = ucitaj_isvu(izvoz([(rbr, "0036000001", "Ana", "Horvat", UPISAN)]))
    assert popis.greske == ()
    assert popis.studenti[0].rbr == 78


def test_razmaci_i_nfc_u_imenu():
    rastavljeno = "S\u030cimek"  # Š kao S + kvačica (NFD)
    popis = ucitaj_isvu(izvoz([("1.", "0036000001", "  Ana   Marija ", rastavljeno, UPISAN)]))
    student = popis.studenti[0]
    assert student.ime == "Ana Marija"
    assert student.prezime == "Šimek"


@pytest.mark.parametrize(
    ("status", "aktivan"),
    [
        (ISPISAN, False),
        ("ISPISAN", False),
        ("  Ispisan s predmeta", False),
        (UPISAN, True),
        (None, True),
    ],
)
def test_ispisan_je_neaktivan(status, aktivan):
    popis = ucitaj_isvu(izvoz([("1.", "0036000001", "Ana", "Horvat", status)]))
    assert popis.studenti[0].aktivan is aktivan


def test_greske_navode_brojeve_redaka():
    popis = ucitaj_isvu(izvoz(S_GRESKAMA))
    r = PRVI_REDAK
    assert _kljucevi(popis) == [
        (r + 1, "popis.rbr_neispravan"),  # „a“
        (r + 2, "popis.rbr_neispravan"),  # „0“
        (r + 4, "popis.dupli_rbr"),
        (r + 5, "popis.dupli_jmbag"),
        (r + 6, "popis.prazan_jmbag"),
        (r + 7, "popis.prazno_ime"),
        (r + 8, "popis.prazno_prezime"),
    ]
    dupli_rbr, dupli_jmbag = popis.greske[2], popis.greske[3]
    assert dict(dupli_rbr.parametri) == {"rbr": 5, "prvi": r + 3}
    assert dict(dupli_jmbag.parametri) == {"prvi": r}
    assert not popis.moze_zamrznuti
    # Retci bez grešaka ostaju u popisu.
    assert [s.rbr for s in popis.studenti] == [1, 5]


@pytest.mark.parametrize("rbr", ["a", "0", "-3", "5,5", "", 0, 5.5, True])
def test_nevaljan_rbr(rbr):
    popis = ucitaj_isvu(izvoz([(rbr, "0036000001", "Ana", "Horvat", UPISAN)]))
    assert _kljucevi(popis) == [(PRVI_REDAK, "popis.rbr_neispravan")]


def test_prazni_retci_se_preskacu_a_brojevi_redaka_ostaju_excelovi():
    prazan = (None, None, None, None, None)
    retci = [ISPRAVNI[0], prazan, ("  ", None, "", None, None), ("2.", "", "Ivo", "Babić", UPISAN)]
    popis = ucitaj_isvu(izvoz(retci))
    assert _kljucevi(popis) == [(PRVI_REDAK + 3, "popis.prazan_jmbag")]


def test_bez_zaglavlja_navodi_stupce_koji_nedostaju():
    zaglavlje = tuple(s for s in ZAGLAVLJE if s != "JMBAG")
    popis = ucitaj_isvu(izvoz(ISPRAVNI, zaglavlje=zaglavlje))
    assert popis.studenti == ()
    assert popis.greske == (
        GreskaRetka(None, "popis.bez_zaglavlja", {"stupci": "JMBAG", "najvise": 50}),
    )
    assert "JMBAG" in poruka("popis.bez_zaglavlja", **popis.greske[0].parametri)
    assert not popis.moze_zamrznuti


def test_zaglavlje_ispod_50_redaka_se_ne_trazi():
    sadrzaj = izvoz([], metapodaci=False)
    knjiga = openpyxl.load_workbook(io.BytesIO(sadrzaj))
    knjiga.active.insert_rows(1, 50)
    izlaz = io.BytesIO()
    knjiga.save(izlaz)
    popis = ucitaj_isvu(izlaz.getvalue())
    assert _kljucevi(popis) == [(None, "popis.bez_zaglavlja")]


@pytest.mark.parametrize("sadrzaj", [b"", b"nije xlsx", b"PK\x03\x04pokvareno"])
def test_necitljiva_datoteka(sadrzaj):
    popis = ucitaj_isvu(sadrzaj)
    assert popis == UcitaniPopis((), (GreskaRetka(None, "popis.necitljiv"),))
    assert not popis.moze_zamrznuti


def test_popis_bez_studenata_nije_zamrzljiv():
    popis = ucitaj_isvu(izvoz([]))
    assert _kljucevi(popis) == [(None, "popis.bez_studenata")]


def test_svaka_greska_ima_poruku():
    popis = ucitaj_isvu(izvoz(S_GRESKAMA))
    kljucevi = {g.poruka_kljuc for g in popis.greske} | {
        "popis.necitljiv",
        "popis.bez_zaglavlja",
        "popis.bez_studenata",
    }
    for g in popis.greske:
        assert poruka(g.poruka_kljuc, **g.parametri)
    assert kljucevi <= KATALOG.keys()


def test_log_ne_sadrzi_osobne_podatke(caplog):
    with caplog.at_level(logging.DEBUG):
        ucitaj_isvu(izvoz(ISPRAVNI))
        ucitaj_isvu(izvoz(S_GRESKAMA))
    for _rbr, jmbag, ime, prezime, _status in ISPRAVNI + S_GRESKAMA:
        for podatak in (jmbag, ime, prezime):
            if str(podatak).strip():
                assert str(podatak) not in caplog.text


@pytest.mark.parametrize("jmbag", ["abc", 36123456.5, "00361234567", 12345678901, "0036-12345", -5])
def test_nevaljan_jmbag(jmbag):
    popis = ucitaj_isvu(izvoz([("1.", jmbag, "Ana", "Horvat", UPISAN)]))
    assert _kljucevi(popis) == [(PRVI_REDAK, "popis.jmbag_neispravan")]


def test_isti_jmbag_kao_tekst_i_broj_je_duplikat():
    retci = [("1.", "0036123456", "Ana", "Horvat", UPISAN), ("2.", 36123456, "Ivo", "Babić", UPISAN)]
    popis = ucitaj_isvu(izvoz(retci))
    assert _kljucevi(popis) == [(PRVI_REDAK + 1, "popis.dupli_jmbag")]


def _knjiga(*listovi: list[tuple[object, ...]], aktivan: int = 0) -> bytes:
    knjiga = openpyxl.Workbook()
    knjiga.remove(knjiga.active)
    for i, retci in enumerate(listovi):
        list_ = knjiga.create_sheet(f"List{i + 1}")
        for redak in retci:
            list_.append(redak)
    knjiga.active = aktivan
    izlaz = io.BytesIO()
    knjiga.save(izlaz)
    return izlaz.getvalue()


def test_redak_zbroja_samo_u_drugim_stupcima_se_preskace():
    zaglavlje = ("Rbr.", "JMBAG", "Ime", "Prezime", "Status upisa", "Napomena")
    sadrzaj = _knjiga(
        [zaglavlje, ("1.", "0036000001", "Ana", "Horvat", UPISAN, None), (None, " ", None, "", None, "Ukupno: 1")]
    )
    popis = ucitaj_isvu(sadrzaj)
    assert popis.greske == ()
    assert len(popis.studenti) == 1


@pytest.mark.parametrize("status", ["Status  upisa", "Status\u00a0upisa", " status upisa "])
def test_zaglavlje_s_visestrukim_razmacima_i_nbsp(status):
    sadrzaj = _knjiga([("Rbr.", "JMBAG", "Ime", "Prezime", status), ISPRAVNI[0]])
    popis = ucitaj_isvu(sadrzaj)
    assert popis.greske == ()
    assert popis.studenti[0].jmbag == "0036123456"


def test_zaglavlje_na_listu_koji_nije_aktivan():
    sadrzaj = _knjiga(
        [("Napomena",), ("nema popisa",)],
        [("Rbr.", "JMBAG", "Ime", "Prezime", "Status upisa"), ISPRAVNI[0]],
        [("Rbr.", "JMBAG", "Ime", "Prezime", "Status upisa"), ISPRAVNI[1]],
        aktivan=0,
    )
    popis = ucitaj_isvu(sadrzaj)
    assert popis.greske == ()
    assert [s.rbr for s in popis.studenti] == [1]  # prvi list sa zaglavljem


def test_parametri_greske_su_nepromjenjivi():
    izvorni = {"rbr": 5}
    greska = GreskaRetka(3, "popis.dupli_rbr", izvorni)
    izvorni["rbr"] = 6
    assert greska.parametri["rbr"] == 5
    with pytest.raises(TypeError):
        greska.parametri["rbr"] = 7  # type: ignore[index]


# --- Zamrznuti popis (priča 3.2) ----------------------------------------------


def _studenti(retci) -> tuple[Student, ...]:
    popis = ucitaj_isvu(izvoz(retci))
    assert popis.moze_zamrznuti
    return popis.studenti


def _zamrznuto(tmp_path: Path, retci=ISPRAVNI) -> tuple[Path, tuple[Student, ...]]:
    putanja = zamrzni(_studenti(retci), tmp_path / "vjezba_1")
    return putanja, procitaj_zamrznuti(putanja)


def test_prvo_zamrzavanje_zapisuje_bom_stupce_i_redoslijed(tmp_path):
    mapa_vjezbe = tmp_path / "nova" / "vjezba_1"  # mapa se stvara po potrebi
    putanja = zamrzni(_studenti(ISPRAVNI), mapa_vjezbe)
    assert putanja == mapa_vjezbe / "popis.csv"
    sadrzaj = putanja.read_bytes()
    assert sadrzaj.startswith(b"\xef\xbb\xbf")
    retci = sadrzaj.decode("utf-8-sig").splitlines()
    assert retci[0] == "rbr,jmbag,ime,prezime,aktivan"
    assert retci[1] == "1,0036123456,Ana Marija,Horvat-Kovač,true"
    assert retci[3] == "3,0036000003,Ivana,Žužić,false"
    assert len(retci) == 1 + len(ISPRAVNI)
    assert [p.name for p in mapa_vjezbe.iterdir()] == ["popis.csv"]  # bez privremenih datoteka


def test_kruzni_put(tmp_path):
    studenti = _studenti(ISPRAVNI)
    putanja = zamrzni(studenti, tmp_path)
    assert procitaj_zamrznuti(putanja) == studenti
    assert usporedi(procitaj_zamrznuti(putanja), studenti) == Razlike()


def test_postojeci_popis_se_ne_prepisuje(tmp_path):
    putanja, _ = _zamrznuto(tmp_path)
    prije = putanja.read_bytes()
    with pytest.raises(PopisVecZamrznut):
        zamrzni(_studenti(KASNI_UPIS), putanja.parent)
    assert putanja.read_bytes() == prije


def test_zamrzavanje_ne_prihvaca_neispravan_popis(tmp_path):
    with pytest.raises(ValueError):
        zamrzni((), tmp_path)
    dupli = (Student(1, "0036000001", "Ana", "Horvat", True),) * 2
    with pytest.raises(ValueError):
        zamrzni(dupli, tmp_path)
    assert not (tmp_path / "popis.csv").exists()


def test_novi_izvoz_bez_promjena_nema_razlika_i_ne_mijenja_datoteku(tmp_path):
    putanja, zamrznuti = _zamrznuto(tmp_path)
    prije = putanja.read_bytes()
    razlike = usporedi(zamrznuti, _studenti(ISPRAVNI))
    assert razlike == Razlike()
    assert not razlike.ima_promjena
    assert putanja.read_bytes() == prije


def test_prenumeriran_izvoz_nema_razlika(tmp_path):
    _, zamrznuti = _zamrznuto(tmp_path)
    izvoz_ = _studenti(PRENUMERIRANI)
    po_jmbagu = {s.jmbag: s.rbr for s in zamrznuti}
    assert any(po_jmbagu[s.jmbag] != s.rbr for s in izvoz_)  # ISVU je prenumerirao
    assert usporedi(zamrznuti, izvoz_) == Razlike()


def test_kasni_upis_dobiva_max_rbr_plus_1():
    zamrznuti = tuple(
        Student(rbr, f"00360000{rbr:02d}", "Ime", f"Prezime{rbr}", True) for rbr in (1, 5, 78)
    )
    izvoz_ = (
        Student(1, "0036000099", "Marta", "Babić", True),  # ISVU Rbr. se zanemaruje
        *(Student(i + 2, s.jmbag, s.ime, s.prezime, True) for i, s in enumerate(zamrznuti)),
        Student(9, "0036000098", "Nera", "Zec", True),
        Student(10, "0036000097", "Iva", "Ispisana", False),  # nije aktivna, pa nije nova
    )
    razlike = usporedi(zamrznuti, izvoz_)
    assert [s.jmbag for s in razlike.novi] == ["0036000099", "0036000098"]
    assert razlike.ispisani == razlike.ponovno_upisani == ()
    novi = primijeni(zamrznuti, razlike)
    assert novi[:3] == zamrznuti
    assert novi[3:] == (
        Student(79, "0036000099", "Marta", "Babić", True),
        Student(80, "0036000098", "Nera", "Zec", True),
    )


def test_kasni_upis_iz_isvu_izvoza(tmp_path):
    putanja, zamrznuti = _zamrznuto(tmp_path)
    razlike = usporedi(zamrznuti, _studenti(KASNI_UPIS))
    assert razlike.ima_promjena
    spremi(primijeni(zamrznuti, razlike), putanja)
    novi = procitaj_zamrznuti(putanja)
    assert novi[:5] == zamrznuti
    assert novi[5] == Student(6, "0036000006", "Marta", "Babić", True)
    assert usporedi(novi, _studenti(KASNI_UPIS)) == Razlike()


def test_ispis_zadrzava_rbr_i_postaje_neaktivan(tmp_path):
    putanja, zamrznuti = _zamrznuto(tmp_path)
    razlike = usporedi(zamrznuti, _studenti(ISPIS))
    # Šimić je „ispisan …“, Čačića nema u izvozu; Žužić je već neaktivna.
    assert [s.rbr for s in razlike.ispisani] == [2, 4]
    assert razlike.novi == razlike.ponovno_upisani == ()
    spremi(primijeni(zamrznuti, razlike), putanja)
    novi = procitaj_zamrznuti(putanja)
    assert [(s.rbr, s.jmbag, s.aktivan) for s in novi] == [
        (s.rbr, s.jmbag, s.rbr not in (2, 3, 4)) for s in zamrznuti
    ]
    assert [(s.ime, s.prezime) for s in novi] == [(s.ime, s.prezime) for s in zamrznuti]


def test_ponovni_upis_postaje_aktivan(tmp_path):
    _, zamrznuti = _zamrznuto(tmp_path)
    razlike = usporedi(zamrznuti, _studenti(PONOVNI_UPIS))
    assert [s.rbr for s in razlike.ponovno_upisani] == [3]
    assert razlike.novi == razlike.ispisani == ()
    novi = primijeni(zamrznuti, razlike)
    assert novi[2] == Student(3, "0036000003", "Ivana", "Žužić", True)
    assert all(n == z for n, z in zip(novi, zamrznuti, strict=True) if n.rbr != 3)


def test_promjena_imena_je_samo_upozorenje(tmp_path):
    _, zamrznuti = _zamrznuto(tmp_path)
    retci = tuple(
        (*redak[:2], "Đurđa", *redak[3:]) if redak[1] == "0036000002" else redak
        for redak in ISPRAVNI
    )
    razlike = usporedi(zamrznuti, _studenti(retci))
    assert not razlike.ima_promjena
    assert len(razlike.promijenjeno_ime) == 1
    promjena = razlike.promijenjeno_ime[0]
    assert (promjena.zamrznuti.ime, promjena.izvoz.ime) == ("Đuro", "Đurđa")
    assert primijeni(zamrznuti, razlike) == zamrznuti


@pytest.mark.parametrize(
    ("ime", "zapisano"),
    [
        ("=SUM(A1)", "'=SUM(A1)"),
        ("+1", "'+1"),
        ("-x", "'-x"),
        ("@cmd", "'@cmd"),
        ("\tx", "'\tx"),
        ("\rx", "'\rx"),
        ("'=a", "''=a"),
        ("'Ana", "''Ana"),
        ("Ana", "Ana"),
    ],
)
def test_zastita_od_formula(tmp_path, ime, zapisano):
    studenti = (Student(1, "0036000001", ime, "=Horvat", True),)
    putanja = zamrzni(studenti, tmp_path)
    tekst = putanja.read_bytes().decode("utf-8-sig")  # bez prevođenja \r
    _, redak = list(csv.reader(io.StringIO(tekst, newline="")))
    assert redak == ["1", "0036000001", zapisano, "'=Horvat", "true"]
    assert procitaj_zamrznuti(putanja) == studenti


def test_istodobno_zamrzavanje_ne_prepisuje_popis(tmp_path, monkeypatch):
    import provjera.word.adapteri.popis as modul

    putanja = tmp_path / "popis.csv"
    pravi_link = modul.os.link

    def drugi_stigao_prvi(izvor, cilj):
        Path(cilj).write_bytes(b"drugi")  # zamrznut nakon provjere postojanja
        pravi_link(izvor, cilj)

    monkeypatch.setattr(modul.os, "link", drugi_stigao_prvi)
    with pytest.raises(PopisVecZamrznut):
        zamrzni(_studenti(ISPRAVNI), tmp_path)
    assert putanja.read_bytes() == b"drugi"
    assert [p.name for p in tmp_path.iterdir()] == ["popis.csv"]


def _bez_tvrdih_poveznica(*_args):
    raise OSError(errno.EPERM, "tvrde poveznice nisu podržane")


def test_zamrzavanje_bez_tvrdih_poveznica(tmp_path, monkeypatch):
    import provjera.word.adapteri.popis as modul

    ocekivano = zamrzni(_studenti(ISPRAVNI), tmp_path / "s_poveznicom").read_bytes()
    monkeypatch.setattr(modul.os, "link", _bez_tvrdih_poveznica)
    putanja = zamrzni(_studenti(ISPRAVNI), tmp_path / "bez_poveznice")
    assert putanja.read_bytes() == ocekivano
    assert [p.name for p in putanja.parent.iterdir()] == ["popis.csv"]


def test_zamrzavanje_bez_tvrdih_poveznica_ne_prepisuje(tmp_path, monkeypatch):
    import provjera.word.adapteri.popis as modul

    putanja = tmp_path / "popis.csv"

    def drugi_stigao_prvi(_izvor, cilj):
        Path(cilj).write_bytes(b"drugi")  # pojavio se prije zamjenskog stvaranja
        _bez_tvrdih_poveznica()

    monkeypatch.setattr(modul.os, "link", drugi_stigao_prvi)
    with pytest.raises(PopisVecZamrznut):
        zamrzni(_studenti(ISPRAVNI), tmp_path)
    assert putanja.read_bytes() == b"drugi"
    assert [p.name for p in tmp_path.iterdir()] == ["popis.csv"]


def test_ime_sa_zarezom_i_navodnicima(tmp_path):
    studenti = (Student(1, "0036000001", 'Ana, "Mia"', "Horvat", False),)
    assert procitaj_zamrznuti(zamrzni(studenti, tmp_path)) == studenti


def _csv(tmp_path: Path, tekst: str, bom: bool = True) -> Path:
    putanja = tmp_path / "popis.csv"
    putanja.write_bytes((b"\xef\xbb\xbf" if bom else b"") + tekst.encode("utf-8"))
    return putanja


ZAGLAVLJE_CSV = "rbr,jmbag,ime,prezime,aktivan\r\n"
ANA = "1,0036000001,Ana,Horvat,true\r\n"


@pytest.mark.parametrize(
    ("sadrzaj", "redak", "kljuc"),
    [
        ("", None, "popis.zamrznuti_prazan"),
        (ZAGLAVLJE_CSV, None, "popis.zamrznuti_prazan"),
        ("rbr,jmbag,ime,prezime\r\n1,0036000001,Ana,Horvat\r\n", 1, "popis.zamrznuti_zaglavlje"),
        ("jmbag,rbr,ime,prezime,aktivan\r\n", 1, "popis.zamrznuti_zaglavlje"),
        (ZAGLAVLJE_CSV + "1,0036000001,Ana,Horvat\r\n", 2, "popis.zamrznuti_broj_stupaca"),
        (ZAGLAVLJE_CSV + ANA + "a,0036000002,Ivo,Babić,true\r\n", 3, "popis.rbr_neispravan"),
        (ZAGLAVLJE_CSV + "0,0036000001,Ana,Horvat,true\r\n", 2, "popis.rbr_neispravan"),
        (ZAGLAVLJE_CSV + "1,36000001,Ana,Horvat,true\r\n", 2, "popis.jmbag_neispravan"),
        (ZAGLAVLJE_CSV + "1,0036000001,Ana,Horvat,da\r\n", 2, "popis.zamrznuti_aktivan"),
        (ZAGLAVLJE_CSV + "1,0036000001,,Horvat,true\r\n", 2, "popis.prazno_ime"),
        (ZAGLAVLJE_CSV + "1,0036000001,Ana, ,true\r\n", 2, "popis.prazno_prezime"),
        (ZAGLAVLJE_CSV + ANA + "1,0036000002,Ivo,Babić,true\r\n", 3, "popis.dupli_rbr"),
        (ZAGLAVLJE_CSV + ANA + "2,0036000001,Ivo,Babić,true\r\n", 3, "popis.dupli_jmbag"),
    ],
)
def test_ostecen_csv_daje_gresku_s_retkom(tmp_path, sadrzaj, redak, kljuc):
    with pytest.raises(GreskaZamrznutog) as greska:
        procitaj_zamrznuti(_csv(tmp_path, sadrzaj))
    assert (greska.value.redak, greska.value.poruka_kljuc) == (redak, kljuc)
    assert poruka(greska.value.poruka_kljuc, **greska.value.parametri)


@pytest.mark.parametrize(
    "sadrzaj",
    [b"\xff\xfe\x00nije utf-8", (ZAGLAVLJE_CSV + '1,0036000001,"Ana,Horvat,true\r\n').encode()],
)
def test_necitljiv_csv(tmp_path, sadrzaj):
    putanja = tmp_path / "popis.csv"
    putanja.write_bytes(sadrzaj)
    with pytest.raises(GreskaZamrznutog) as greska:
        procitaj_zamrznuti(putanja)
    assert greska.value.poruka_kljuc == "popis.zamrznuti_necitljiv"


def test_nepostojeci_csv_daje_gresku(tmp_path):
    with pytest.raises(GreskaZamrznutog) as greska:
        procitaj_zamrznuti(tmp_path / "ne_postoji.csv")
    assert (greska.value.redak, greska.value.poruka_kljuc) == (None, "popis.zamrznuti_necitljiv")


def test_csv_bez_bom_a_i_s_lf_je_citljiv(tmp_path):
    tekst = "rbr,jmbag,ime,prezime,aktivan\n1,0036000001,Ana,Horvat,true\n"
    putanja = _csv(tmp_path, tekst, bom=False)
    assert procitaj_zamrznuti(putanja) == (Student(1, "0036000001", "Ana", "Horvat", True),)


def test_neuspjelo_pisanje_ne_mijenja_popis_ni_ostavlja_privremenu(tmp_path, monkeypatch):
    import provjera.word.adapteri.popis as modul

    putanja, zamrznuti = _zamrznuto(tmp_path)
    prije = putanja.read_bytes()

    def pukni(*_args):
        raise OSError("disk pun")

    monkeypatch.setattr(modul.os, "replace", pukni)
    with pytest.raises(OSError):
        spremi(zamrznuti[:1], putanja)
    assert putanja.read_bytes() == prije
    assert [p.name for p in putanja.parent.iterdir()] == ["popis.csv"]


def test_log_zamrzavanja_ne_sadrzi_osobne_podatke(tmp_path, caplog):
    with caplog.at_level(logging.DEBUG):
        putanja, zamrznuti = _zamrznuto(tmp_path)
        spremi(primijeni(zamrznuti, usporedi(zamrznuti, _studenti(KASNI_UPIS))), putanja)
    for s in procitaj_zamrznuti(putanja):
        for podatak in (s.jmbag, s.ime, s.prezime):
            assert podatak not in caplog.text


# --- izdano.json (priča 3.5) --------------------------------------------------

PROFIL_VJEZBE_1 = Path(__file__).resolve().parents[2] / "profili" / "word" / "vjezba_1.toml"
ZAGREB = ZoneInfo("Europe/Zagreb")


def _profil(zamjena: tuple[str, str] | None = None) -> profil_modul.Profil:
    tekst = PROFIL_VJEZBE_1.read_text(encoding="utf-8")
    if zamjena:
        tekst = tekst.replace(*zamjena)
    return profil_modul.ucitaj(tekst.encode("utf-8"))


def test_putanje_izdanog_i_pdf(tmp_path):
    mapa = tmp_path / "vjezba_1"
    assert putanja_izdanog(mapa) == mapa / "izdano.json"
    assert putanja_pdf(mapa, "vjezba_1") == mapa / "svojstva_vjezba_1.pdf"


def test_zapisi_izdano_otisak_verzija_i_vrijeme(tmp_path):
    mapa = tmp_path / "vjezba_1"
    profil = _profil()
    # Ljetno vrijeme: 10:30 UTC je 12:30 u Zagrebu (+02:00).
    putanja = zapisi_izdano(mapa, profil, datetime(2026, 7, 1, 10, 30, tzinfo=timezone.utc))
    assert putanja == mapa / "izdano.json"
    podaci = json.loads(putanja.read_text(encoding="utf-8"))
    assert podaci == {
        "formula_otisak": formula_modul.otisak(profil.formula),
        "profil_verzija": profil_modul.verzija(profil),
        "vrijeme": "2026-07-01T12:30:00+02:00",
    }
    izdano = procitaj_izdano(mapa)
    assert izdano == Izdano(
        formula_modul.otisak(profil.formula),
        profil_modul.verzija(profil),
        datetime(2026, 7, 1, 12, 30, tzinfo=ZAGREB),
    )
    assert izdano.vrijeme.utcoffset() == timedelta(hours=2)


def test_zapisi_izdano_zimsko_vrijeme(tmp_path):
    zapisi_izdano(tmp_path, _profil(), datetime(2026, 12, 1, 9, 0, tzinfo=ZAGREB))
    podaci = json.loads(putanja_izdanog(tmp_path).read_text(encoding="utf-8"))
    assert podaci["vrijeme"] == "2026-12-01T09:00:00+01:00"


def test_zapisi_izdano_prepisuje_postojeci(tmp_path):
    vrijeme = datetime(2026, 10, 14, 12, 0, tzinfo=ZAGREB)
    zapisi_izdano(tmp_path, _profil(), vrijeme)
    drugi = _profil(("sjeme = 20261012", "sjeme = 1"))
    zapisi_izdano(tmp_path, drugi, vrijeme + timedelta(hours=1))
    izdano = procitaj_izdano(tmp_path)
    assert izdano.formula_otisak == formula_modul.otisak(drugi.formula)
    assert izdano.profil_verzija == profil_modul.verzija(drugi)
    assert izdano.vrijeme == vrijeme + timedelta(hours=1)
    assert [p.name for p in tmp_path.iterdir()] == ["izdano.json"]


def test_zapisi_izdano_bez_vremenske_zone_baca(tmp_path):
    with pytest.raises(ValueError):
        zapisi_izdano(tmp_path, _profil(), datetime(2026, 10, 14, 12, 0))
    assert not putanja_izdanog(tmp_path).exists()


def test_neuspjelo_pisanje_izdanog_ne_mijenja_postojeci(tmp_path, monkeypatch):
    import provjera.word.adapteri.popis as modul

    zapisi_izdano(tmp_path, _profil(), datetime(2026, 10, 14, 12, 0, tzinfo=ZAGREB))
    prije = putanja_izdanog(tmp_path).read_bytes()

    def pukni(*_args):
        raise OSError("disk pun")

    monkeypatch.setattr(modul.os, "replace", pukni)
    with pytest.raises(OSError):
        zapisi_izdano(tmp_path, _profil(), datetime(2026, 10, 15, 12, 0, tzinfo=ZAGREB))
    assert putanja_izdanog(tmp_path).read_bytes() == prije
    assert [p.name for p in tmp_path.iterdir()] == ["izdano.json"]


def test_procitaj_izdano_bez_datoteke(tmp_path):
    assert procitaj_izdano(tmp_path) is None


@pytest.mark.parametrize(
    "sadrzaj",
    [
        b"{",
        b"[]",
        b"\xff\xfe",
        b'{"formula_otisak": "abc", "profil_verzija": "def"}',
        b'{"formula_otisak": "abc", "profil_verzija": "def", "vrijeme": "2026-10-14T12:00:00"}',
        b'{"formula_otisak": "", "profil_verzija": "def", "vrijeme": "2026-10-14T12:00:00+02:00"}',
        b'{"formula_otisak": 1, "profil_verzija": "def", "vrijeme": "2026-10-14T12:00:00+02:00"}',
        b'{"formula_otisak": "abc", "profil_verzija": "def", "vrijeme": 5}',
        b'{"formula_otisak": "abc", "profil_verzija": "def", "vrijeme": "jucer"}',
    ],
)
def test_procitaj_izdano_neispravno(tmp_path, sadrzaj):
    putanja_izdanog(tmp_path).write_bytes(sadrzaj)
    with pytest.raises(GreskaIzdanog) as greska:
        procitaj_izdano(tmp_path)
    assert greska.value.poruka_kljuc in KATALOG


def test_bez_zaglavlja_vise_listova_navodi_najblizi_list():
    bez_jmbaga = tuple(s for s in ZAGLAVLJE if s != "JMBAG")
    popis = ucitaj_isvu(_knjiga([("Naslovnica",)], [bez_jmbaga, (1, "Ana", "Horvat", "upisan")]))
    assert popis.greske == (
        GreskaRetka(None, "popis.bez_zaglavlja", {"stupci": "JMBAG", "najvise": 50}),
    )


def test_zaglavlje_u_50_retku_se_pronalazi():
    sadrzaj = izvoz(ISPRAVNI, metapodaci=False)
    knjiga = openpyxl.load_workbook(io.BytesIO(sadrzaj))
    knjiga.active.insert_rows(1, 49)
    izlaz = io.BytesIO()
    knjiga.save(izlaz)
    popis = ucitaj_isvu(izlaz.getvalue())
    assert popis.greske == ()
    assert len(popis.studenti) == len(ISPRAVNI)


def test_zakljucan_zamrznuti_popis_nije_ostecen(tmp_path, monkeypatch):
    import provjera.word.adapteri.popis as modul

    putanja = zamrzni(_studenti(ISPRAVNI), tmp_path)

    def zakljucano(*_a, **_k):
        raise PermissionError(13, "Pristup odbijen")

    monkeypatch.setattr(modul.Path, "read_bytes", zakljucano)
    with pytest.raises(GreskaZamrznutog) as greska:
        procitaj_zamrznuti(putanja)
    assert greska.value.poruka_kljuc == "popis.zamrznuti_zakljucan"


def test_zamrzavanje_bez_tvrdih_poveznica_neuspjelo_pisanje_ne_ostavlja_datoteku(tmp_path, monkeypatch):
    import provjera.word.adapteri.popis as modul

    monkeypatch.setattr(modul.os, "link", _bez_tvrdih_poveznica)

    def puca(_opisnik):
        raise OSError("disk je pun")

    monkeypatch.setattr(modul.os, "fsync", puca)
    with pytest.raises(OSError):
        zamrzni(_studenti(ISPRAVNI), tmp_path)
    assert list(tmp_path.iterdir()) == []

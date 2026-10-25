"""R-BROJ-STRANICE: polje broja stranice u podnožjima (priča 2.5)."""

import io
import re
from pathlib import Path

import docx as python_docx
import pytest
from alati.generator import (
    VARIJANTE,
    ispravan,
    podnozje_s_brojem,
    polje_page,
    tema_s_fontom,
    zamijeni_u_zipu,
)
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from provjera.word.jezgra import model, profil
from provjera.word.jezgra.model import VrstaPodnozja
from provjera.word.jezgra.nalaz import Dio, KontekstProvjere, Status, VrstaIzvora
from provjera.word.jezgra.poruke import KATALOG, tekst_nalaza
from provjera.word.jezgra.pravila import REGISTAR
from provjera.word.jezgra.provjera import provjeri, status_pravila

KORIJEN = Path(__file__).resolve().parents[2]
TOML = (KORIJEN / "profili" / "word" / "vjezba_1.toml").read_text(encoding="utf-8")
PROFIL = profil.ucitaj(TOML.encode("utf-8"))
KONTEKST = KontekstProvjere("Ana_Horvat_GiG_OG_Vjezba_1.docx", 54, None, None)
UZORAK = KORIJEN / "samples" / "Lucija_Šimek_GiG_OG_Vjezba_1.docx"
ID = "R-BROJ-STRANICE"


def _izmijeni(promjena) -> bytes:
    dokument = python_docx.Document(io.BytesIO(ispravan(54, PROFIL)))
    promjena(dokument)
    spremnik = io.BytesIO()
    dokument.save(spremnik)
    return spremnik.getvalue()


def _nalazi(sadrzaj: bytes, p=PROFIL):
    return [n for n in provjeri(sadrzaj, p, KONTEKST).nalazi if n.pravilo_id == ID]


def _kljucevi(sadrzaj: bytes) -> set[str]:
    return {n.poruka_kljuc.removeprefix(f"{ID}.") for n in _nalazi(sadrzaj)}


def _model(sadrzaj: bytes):
    return model.procitaj(sadrzaj, PROFIL.uloge)


def _isprazni(podnozje):
    korijen = podnozje._element
    for dijete in list(korijen):
        korijen.remove(dijete)
    return podnozje.add_paragraph(style="Footer")


def _bez_reference(sekcija) -> None:
    for ref in sekcija._sectPr.findall(qn("w:footerReference")):
        sekcija._sectPr.remove(ref)


def _parne_neparne(dokument) -> None:
    dokument.settings.element.append(OxmlElement("w:evenAndOddHeaders"))


# --- matrica iz specifikacije ---------------------------------------------------


def test_ispravan_rad_daje_jedan_prolaz():
    [nalaz] = _nalazi(ispravan(54, PROFIL))
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.PROSLO, f"{ID}.prolazi")


@pytest.mark.parametrize("ime", ["broj_galerija", "broj_okvir"])
def test_galerija_i_tekstni_okvir_prolaze(ime):
    sadrzaj = VARIJANTE[ime](54, PROFIL)
    assert [n.status for n in _nalazi(sadrzaj)] == [Status.PROSLO]
    [polje] = _model(sadrzaj).sekcije[1].podnozja.zadano.polja_stranice
    assert [r.tekst for r in polje.runovi] == ["2"]


@pytest.mark.parametrize(
    ("ime", "kljuc"),
    [
        ("broj_rucni", "rucni_broj"),
        ("broj_nema", "bez_broja"),
        ("broj_naslovnica", "naslovnica_ima_broj"),
        ("broj_prva_stranica", "bez_broja_prva"),
        ("broj_font", "velicina"),
    ],
)
def test_pokvarene_varijante(ime, kljuc):
    nalazi = _nalazi(VARIJANTE[ime](54, PROFIL))
    assert [(n.status, n.poruka_kljuc) for n in nalazi] == [(Status.NE_PROLAZI, f"{ID}.{kljuc}")]
    assert nalazi[0].lokacija.dio is Dio.PODNOZJE
    assert "{" not in tekst_nalaza(nalazi[0])


def test_povezano_s_naslovnicom_tijelo_prolazi():
    [nalaz] = _nalazi(VARIJANTE["broj_naslovnica"](54, PROFIL))
    assert nalaz.lokacija.sekcija == 1
    tekst = tekst_nalaza(nalaz)
    assert "Poveži s prethodnim (Link to Previous)" in tekst
    assert "Drukčija prva stranica (Different First Page)" in tekst
    tijelo = _model(VARIJANTE["broj_naslovnica"](54, PROFIL)).sekcije[1]
    assert len(tijelo.podnozja.zadano.polja_stranice) == 1


def test_krivi_font_navodi_velicinu_i_izvor():
    [nalaz] = _nalazi(VARIJANTE["broj_font"](54, PROFIL))
    assert (nalaz.pronadeno.iznos, nalaz.ocekivano.iznos) == (22, 20)
    assert nalaz.izvor.vrsta is VrstaIzvora.IZRAVNO
    assert "11 pt" in tekst_nalaza(nalaz) and "10 pt" in tekst_nalaza(nalaz)


def test_rucni_broj_daje_uputu_za_umetanje():
    [nalaz] = _nalazi(VARIJANTE["broj_rucni"](54, PROFIL))
    assert "Umetanje (Insert) → Broj stranice (Page Number) → Dno stranice (Bottom of Page)" in tekst_nalaza(nalaz)


# --- model ------------------------------------------------------------------------


def test_nasljedivanje_kroz_tri_sekcije():
    def tri(d):
        d.add_section(WD_SECTION_START.NEW_PAGE)
        _bez_reference(d.sections[2])

    sadrzaj = _izmijeni(tri)
    m = _model(sadrzaj)
    assert [len(s.podnozja.zadano.polja_stranice) for s in m.sekcije] == [0, 1, 1]
    assert [n.status for n in _nalazi(sadrzaj)] == [Status.PROSLO]


def test_treca_sekcija_s_rucnim_brojem():
    def tri(d):
        sekcija = d.add_section(WD_SECTION_START.NEW_PAGE)
        sekcija.footer.is_linked_to_previous = False
        _isprazni(sekcija.footer).add_run("3")

    [nalaz] = _nalazi(_izmijeni(tri))
    assert (nalaz.poruka_kljuc, nalaz.lokacija.sekcija) == (f"{ID}.rucni_broj", 3)


def test_naslovnica_s_drukcijom_prvom_stranicom_prolazi():
    """Naslovnica ima broj samo u zadanom podnožju, a prvo podnožje joj je prazno."""

    def prva(d):
        naslovnica = d.sections[0]
        podnozje_s_brojem(naslovnica)
        naslovnica.different_first_page_header_footer = True
        naslovnica.first_page_footer.is_linked_to_previous = False
        d.sections[1].footer.is_linked_to_previous = False

    assert [n.status for n in _nalazi(_izmijeni(prva))] == [Status.PROSLO]


def test_drukcija_prva_bez_prvog_podnozja_je_prazna():
    """Prvo podnožje bez reference ni u jednoj sekciji je prazno (ne nasljeđuje zadano)."""

    def prva(d):
        d.sections[1].different_first_page_header_footer = True

    sadrzaj = _izmijeni(prva)
    assert _model(sadrzaj).sekcije[1].prva_stranica
    assert _kljucevi(sadrzaj) == {"bez_broja_prva"}


def test_parna_podnozja():
    def bez_parnog(d):
        _parne_neparne(d)

    def s_parnim(d):
        _parne_neparne(d)
        podnozje = d.sections[1].even_page_footer
        podnozje.is_linked_to_previous = False
        polje_page(podnozje.paragraphs[0])

    bez = _izmijeni(bez_parnog)
    assert _model(bez).parna_neparna
    assert _kljucevi(bez) == {"bez_broja_parna"}
    s = _izmijeni(s_parnim)
    assert len(_model(s).sekcije[1].podnozja.po_vrsti(VrstaPodnozja.PARNA).polja_stranice) == 1
    assert [n.status for n in _nalazi(s)] == [Status.PROSLO]


def test_parna_podnozja_se_ne_provjeravaju_bez_postavke():
    def parno_bez_postavke(d):
        podnozje = d.sections[1].even_page_footer
        podnozje.is_linked_to_previous = False

    assert [n.status for n in _nalazi(_izmijeni(parno_bez_postavke))] == [Status.PROSLO]


def test_numpages_nije_broj_stranice():
    def numpages(d):
        polje_page(_isprazni(d.sections[1].footer), naredba=" NUMPAGES ", rezultat="9")

    sadrzaj = _izmijeni(numpages)
    podnozje = _model(sadrzaj).sekcije[1].podnozja.zadano
    assert podnozje.polja_stranice == () and podnozje.tekst_izvan_polja == ""
    assert _kljucevi(sadrzaj) == {"bez_broja"}


def test_page_i_numpages_zajedno():
    def oba(d):
        odlomak = d.sections[1].footer.paragraphs[0]
        odlomak.add_run(" od ")
        polje_page(odlomak, naredba=" NUMPAGES ", rezultat="9")

    sadrzaj = _izmijeni(oba)
    podnozje = _model(sadrzaj).sekcije[1].podnozja.zadano
    assert len(podnozje.polja_stranice) == 1 and podnozje.tekst_izvan_polja == " od "
    assert [n.status for n in _nalazi(sadrzaj)] == [Status.PROSLO]


@pytest.mark.parametrize("naredba", ["page", r" Page \* roman ", "PAGE\\* Arabic"])
def test_naredba_bez_obzira_na_velika_slova_i_prekidace(naredba):
    def polje(d):
        polje_page(_isprazni(d.sections[1].footer), naredba=naredba)

    assert len(_model(_izmijeni(polje)).sekcije[1].podnozja.zadano.polja_stranice) == 1


def test_slozeno_polje_preko_vise_runova_i_odlomaka():
    """Naredba je razlomljena na runove, a rezultat je u drugom odlomku."""
    w = nsdecls("w")

    def razlomljeno(d):
        prvi = _isprazni(d.sections[1].footer)
        for xml in (
            f'<w:r {w}><w:fldChar w:fldCharType="begin"/></w:r>',
            f'<w:r {w}><w:instrText xml:space="preserve"> PA</w:instrText></w:r>',
            f'<w:r {w}><w:instrText xml:space="preserve">GE </w:instrText></w:r>',
            f'<w:r {w}><w:fldChar w:fldCharType="separate"/></w:r>',
        ):
            prvi._p.append(parse_xml(xml))
        drugi = d.sections[1].footer.add_paragraph(style="Footer")
        for xml in (
            f'<w:r {w}><w:t>1</w:t></w:r>',
            f'<w:r {w}><w:rPr><w:sz w:val="22"/></w:rPr><w:t>2</w:t></w:r>',
            f'<w:r {w}><w:fldChar w:fldCharType="end"/></w:r>',
        ):
            drugi._p.append(parse_xml(xml))

    sadrzaj = _izmijeni(razlomljeno)
    [polje] = _model(sadrzaj).sekcije[1].podnozja.zadano.polja_stranice
    assert [(r.tekst, r.velicina.vrijednost) for r in polje.runovi] == [("1", 20), ("2", 22)]
    assert _kljucevi(sadrzaj) == {"velicina"}


def test_jednostavno_polje():
    def fld_simple(d):
        _isprazni(d.sections[1].footer)._p.append(parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>4</w:t></w:r></w:fldSimple>'
        ))

    sadrzaj = _izmijeni(fld_simple)
    [polje] = _model(sadrzaj).sekcije[1].podnozja.zadano.polja_stranice
    assert [r.tekst for r in polje.runovi] == ["4"]
    assert [n.status for n in _nalazi(sadrzaj)] == [Status.PROSLO]


def test_polje_bez_rezultata_provjerava_run_s_naredbom():
    w = nsdecls("w")

    def bez_rezultata(d):
        odlomak = _isprazni(d.sections[1].footer)
        for xml in (
            f'<w:r {w}><w:fldChar w:fldCharType="begin"/></w:r>',
            f'<w:r {w}><w:rPr><w:sz w:val="28"/></w:rPr><w:instrText> PAGE </w:instrText></w:r>',
            f'<w:r {w}><w:fldChar w:fldCharType="end"/></w:r>',
        ):
            odlomak._p.append(parse_xml(xml))

    sadrzaj = _izmijeni(bez_rezultata)
    [polje] = _model(sadrzaj).sekcije[1].podnozja.zadano.polja_stranice
    assert [r.velicina.vrijednost for r in polje.runovi] == [28]
    assert _kljucevi(sadrzaj) == {"velicina"}


def test_alternativni_sadrzaj_okvira_broji_se_jednom():
    """Word okvir zapisuje dvaput (mc:Choice i mc:Fallback); čita se samo prva grana."""
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"

    def okvir(d):
        podnozje = d.sections[1].footer
        unutarnji = podnozje.paragraphs[0]._p
        vanjski = _isprazni(podnozje)
        run = parse_xml(
            f'<w:r {nsdecls("w")} xmlns:mc="{mc}"><mc:AlternateContent>'
            "<mc:Choice Requires=\"wps\"><w:txbxContent/></mc:Choice>"
            "<mc:Fallback><w:txbxContent/></mc:Fallback>"
            "</mc:AlternateContent></w:r>"
        )
        for spremnik in run.iter(qn("w:txbxContent")):
            spremnik.append(parse_xml(unutarnji.xml))
        vanjski._p.append(run)

    assert len(_model(_izmijeni(okvir)).sekcije[1].podnozja.zadano.polja_stranice) == 1


def test_izbrisani_broj_se_ne_cita():
    def izbrisano(d):
        _isprazni(d.sections[1].footer)._p.append(parse_xml(
            f'<w:del {nsdecls("w")} w:id="1" w:author="A"><w:r><w:delText>2</w:delText></w:r></w:del>'
        ))

    assert _kljucevi(_izmijeni(izbrisano)) == {"bez_broja"}


def _footer_iz_teme(d):
    stil = d.styles["Footer"]
    rfonts = stil.element.get_or_add_rPr().get_or_add_rFonts()
    for atribut in list(rfonts.attrib):
        del rfonts.attrib[atribut]
    rfonts.set(qn("w:asciiTheme"), "minorHAnsi")


def test_font_teme_u_podnozju():
    sadrzaj = zamijeni_u_zipu(_izmijeni(_footer_iz_teme), {"word/theme/theme1.xml": tema_s_fontom("Aptos")})
    [polje] = _model(sadrzaj).sekcije[1].podnozja.zadano.polja_stranice
    assert polje.runovi[0].font == model.Efektivno("Aptos", polje.runovi[0].font.izvor)
    [nalaz] = _nalazi(sadrzaj)
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, f"{ID}.font_tema")
    assert nalaz.izvor.vrsta is VrstaIzvora.TEMA
    assert "{" not in tekst_nalaza(nalaz)


def test_bez_teme_font_podnozja_je_neprovjerljiv():
    def bez_veze(xml: bytes) -> bytes:
        return re.sub(rb"<Relationship [^>]*/theme\"[^>]*/>", b"", xml)

    sadrzaj = zamijeni_u_zipu(_izmijeni(_footer_iz_teme), {
        "word/theme/theme1.xml": lambda _: None,
        "word/_rels/document.xml.rels": bez_veze,
    })
    [nalaz] = _nalazi(sadrzaj)
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NEPROVJERLJIVO, f"{ID}.font_nepoznat")
    assert "{" not in tekst_nalaza(nalaz)


# --- samostalni ručni broj ----------------------------------------------------------


def test_godina_u_podnozju_naslovnice_prolazi():
    def godina(d):
        podnozje = d.sections[0].footer
        podnozje.is_linked_to_previous = False
        podnozje.paragraphs[0].add_run("Varaždin, 2025.")
        podnozje.add_paragraph("0000000000", style="Footer")
        d.sections[1].footer.is_linked_to_previous = False

    assert [n.status for n in _nalazi(_izmijeni(godina))] == [Status.PROSLO]


@pytest.mark.parametrize("tekst", ["- 2 -", "Stranica 2", "str. 12", "2"])
def test_samostalni_broj_u_tijelu_je_rucni(tekst):
    def rucni(d):
        _isprazni(d.sections[1].footer).add_run(tekst)

    assert _kljucevi(_izmijeni(rucni)) == {"rucni_broj"}


def test_godina_u_tijelu_nije_rucni_broj():
    def godina(d):
        _isprazni(d.sections[1].footer).add_run("Varaždin, 2025.")

    assert _kljucevi(_izmijeni(godina)) == {"bez_broja"}


def test_upisani_broj_na_naslovnici():
    def jedan(d):
        podnozje = d.sections[0].footer
        podnozje.is_linked_to_previous = False
        podnozje.paragraphs[0].add_run("1")
        d.sections[1].footer.is_linked_to_previous = False

    [nalaz] = _nalazi(_izmijeni(jedan))
    assert (nalaz.poruka_kljuc, nalaz.lokacija.sekcija) == (f"{ID}.naslovnica_ima_broj", 1)


# --- povezana podnožja ----------------------------------------------------------


def test_povezano_podnozje_prijavljuje_se_jednom():
    """Tri sekcije tijela dijele isto podnožje s krivom veličinom: nalaz samo na prvoj."""

    def tri(d):
        _isprazni(d.sections[1].footer)
        polje_page(d.sections[1].footer.paragraphs[0], rpr=f'<w:rPr {nsdecls("w")}><w:sz w:val="22"/></w:rPr>')
        for _ in range(2):
            d.add_section(WD_SECTION_START.NEW_PAGE)
            _bez_reference(d.sections[-1])

    sadrzaj = _izmijeni(tri)
    m = _model(sadrzaj)
    assert len({s.podnozja.zadano.dio for s in m.sekcije[1:]}) == 1
    nalazi = _nalazi(sadrzaj)
    assert [(n.poruka_kljuc, n.lokacija.sekcija, n.pojave) for n in nalazi] == [(f"{ID}.velicina", 2, 1)]


def test_povezani_rucni_broj_prijavljuje_se_jednom():
    def tri(d):
        _isprazni(d.sections[1].footer).add_run("2")
        for _ in range(2):
            d.add_section(WD_SECTION_START.NEW_PAGE)
            _bez_reference(d.sections[-1])

    nalazi = _nalazi(_izmijeni(tri))
    assert [(n.poruka_kljuc, n.lokacija.sekcija, n.pojave) for n in nalazi] == [(f"{ID}.rucni_broj", 2, 1)]


def test_poruke_bez_broja_spominju_zaglavlje():
    for kljuc in ("bez_broja", "bez_broja_prva", "bez_broja_parna"):
        assert "zaglavlje (Header" in KATALOG[f"{ID}.{kljuc}"]
    assert "naslovnic" not in KATALOG[f"{ID}.prolazi"]


def test_model_ne_mijenja_paket(monkeypatch):
    """titlePg i evenAndOddHeaders bez podnožja: model ne smije dodati definicije podnožja."""

    def bez_podnozja(d):
        d.sections[1].footer.is_linked_to_previous = True
        d.sections[1].different_first_page_header_footer = True
        _parne_neparne(d)

    sadrzaj = _izmijeni(bez_podnozja)
    otvoreni = []
    izvorni = python_docx.Document

    def prati(*args, **kwargs):
        dokument = izvorni(*args, **kwargs)
        otvoreni.append(dokument)
        return dokument

    def stanje(d):
        dijelovi = sorted(str(p.partname) for p in d.part.package.iter_parts())
        reference = [len(s._sectPr.findall(qn("w:footerReference"))) for s in d.sections]
        return dijelovi, reference, sorted(d.part.rels)

    monkeypatch.setattr(model.docx, "Document", prati)
    m = model.procitaj(sadrzaj, PROFIL.uloge)
    assert stanje(otvoreni[0]) == stanje(izvorni(io.BytesIO(sadrzaj)))
    assert m.parna_neparna and m.sekcije[1].prva_stranica
    assert m.sekcije[1].podnozja == model.PodnozjaSekcije()


# --- pravilo i profil -------------------------------------------------------------


def test_registracija_nakon_stilova_i_poruke():
    redoslijed = list(REGISTAR)
    assert redoslijed.index(ID) == redoslijed.index("R-STILOVI") + 1
    assert f"pravilo.{ID}" in KATALOG
    for kljuc in REGISTAR[ID].poruke:
        assert kljuc in KATALOG, kljuc
    kljucevi = {k.removeprefix(f"{ID}.") for k in REGISTAR[ID].poruke}
    assert {"prolazi", "bez_broja", "bez_broja_prva", "bez_broja_parna", "rucni_broj", "naslovnica_ima_broj",
            "font", "font_tema", "font_nepoznat", "velicina"} <= kljucevi


def test_dokument_bez_sekcija_je_neprovjerljiv():
    pravilo = REGISTAR[ID]
    [nalaz] = pravilo.funkcija(model.DocModel(), KONTEKST, PROFIL.pravila[ID], PROFIL)
    assert nalaz.status is Status.NEPROVJERLJIVO


def _odjeljak(sadrzaj: str) -> str:
    pocetak = TOML.index(f"[pravila.{ID}]")
    kraj = TOML.find("\n[", pocetak)
    return TOML[:pocetak] + f"[pravila.{ID}]\n{sadrzaj}\n" + (TOML[kraj:] if kraj != -1 else "")


def test_pravilo_se_moze_iskljuciti_bez_ciljeva():
    iskljuceno = profil.ucitaj(_odjeljak("ukljuceno = false").encode("utf-8"))
    assert _nalazi(VARIJANTE["broj_nema"](54, PROFIL), iskljuceno) == []
    with pytest.raises(profil.GreskaProfila):
        profil.ucitaj(_odjeljak("ukljuceno = true").encode("utf-8"))


def test_ostala_pravila_ne_vide_podnozje():
    """Broj stranice u podnožju ne mijenja nalaze pravila tijela."""
    for ime in ("broj_rucni", "broj_nema", "broj_naslovnica", "broj_prva_stranica", "broj_font"):
        nalazi = provjeri(VARIJANTE[ime](54, PROFIL), PROFIL, KONTEKST).nalazi
        assert {n.pravilo_id for n in nalazi if n.status is not Status.PROSLO} == {ID}
        assert status_pravila(nalazi, "R-TEKST") is Status.PROSLO


@pytest.mark.skipif(not UZORAK.exists(), reason="stvarni uzorak postoji samo lokalno")
def test_stvarni_uzorak_prepoznaje_polje_iz_galerije():
    sadrzaj = UZORAK.read_bytes()
    m = model.procitaj(sadrzaj, PROFIL.uloge)
    assert any(s.podnozja.zadano.polja_stranice for s in m.sekcije)
    statusi = {n.status for n in provjeri(sadrzaj, PROFIL, KontekstProvjere(UZORAK.name, 54, None, None)).nalazi
               if n.pravilo_id == ID}
    assert Status.NEPROVJERLJIVO not in statusi


def _podnozje_tijela_s_tekstom(tekst: str):
    def promjena(d):
        _isprazni(d.sections[1].footer).add_run(tekst)

    return promjena


def test_rucni_broj_uz_tekst_odvojen_tabulatorom():
    assert _kljucevi(_izmijeni(_podnozje_tijela_s_tekstom("Ana Horvat\t\t2"))) == {"rucni_broj"}


def test_rucni_broj_iza_pozicijskog_tabulatora():
    def promjena(d):
        odlomak = _isprazni(d.sections[1].footer)
        odlomak._p.append(parse_xml(
            f'<w:r {nsdecls("w")}><w:t>Ana Horvat</w:t>'
            '<w:ptab w:relativeTo="margin" w:alignment="right" w:leader="none"/><w:t>2</w:t></w:r>'
        ))

    assert _kljucevi(_izmijeni(promjena)) == {"rucni_broj"}


@pytest.mark.parametrize(
    ("vrijednost", "ukljuceno"),
    [(None, True), ("1", True), ("true", True), ("on", True), ("0", False), ("false", False), ("off", False)],
)
def test_prva_stranica_i_parne_prema_vrijednosti(vrijednost, ukljuceno):
    atribut = "" if vrijednost is None else f' w:val="{vrijednost}"'

    def promjena(d):
        d.sections[1]._sectPr.append(parse_xml(f'<w:titlePg {nsdecls("w")}{atribut}/>'))
        d.settings.element.append(parse_xml(f'<w:evenAndOddHeaders {nsdecls("w")}{atribut}/>'))

    m = _model(_izmijeni(promjena))
    assert m.sekcije[1].prva_stranica is ukljuceno
    assert m.parna_neparna is ukljuceno


def test_necitljivo_podnozje_je_ostecena_datoteka():
    import zipfile

    sadrzaj = ispravan(54, PROFIL)
    with zipfile.ZipFile(io.BytesIO(sadrzaj)) as arhiva:
        podnozja = [i for i in arhiva.namelist() if re.fullmatch(r"word/footer\d+\.xml", i)]
    assert podnozja
    pokvaren = zamijeni_u_zipu(sadrzaj, {ime: lambda _x: b"<smece" for ime in podnozja})
    # python-docx parsira XML dijelove pri otvaranju: pokvareno podnožje je oštećena datoteka.
    from provjera.zajednicko.ulaz import RazlogOdbijanja

    assert provjeri(pokvaren, PROFIL, KONTEKST).odbijeno is RazlogOdbijanja.OSTECENA

"""R-STILOVI: ručno oblikovanje osnovnog teksta i naslova (priča 2.4)."""

import io
import re
from pathlib import Path

import docx as python_docx
import pytest
from alati.generator import VARIJANTE, ispravan, tema_s_fontom, zamijeni_u_zipu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, Twips

from provjera.jezgra import model, profil
from provjera.jezgra.model.uloge import UlogaOdlomka
from provjera.jezgra.nalaz import KontekstProvjere, Status, VrstaIzvora
from provjera.jezgra.poruke import KATALOG, tekst_nalaza
from provjera.jezgra.pravila import REGISTAR
from provjera.jezgra.provjera import provjeri, status_pravila

KORIJEN = Path(__file__).resolve().parents[1]
TOML = (KORIJEN / "profili" / "vjezba_1.toml").read_text(encoding="utf-8")
PROFIL = profil.ucitaj(TOML.encode("utf-8"))
KONTEKST = KontekstProvjere("Ana_Horvat_GiG_OG_Vjezba_1.docx", 54, None, None)
ID = "R-STILOVI"
KLJUCEVI = {"izravni_font", "izravna_velicina", "izravno_poravnanje", "izravni_razmak", "izravni_prored",
            "izravna_uvlaka"}


def _izmijeni(promjena) -> bytes:
    dokument = python_docx.Document(io.BytesIO(ispravan(54, PROFIL)))
    promjena(dokument)
    spremnik = io.BytesIO()
    dokument.save(spremnik)
    return spremnik.getvalue()


def _nalazi(sadrzaj: bytes, pravilo_id: str = ID):
    return [n for n in provjeri(sadrzaj, PROFIL, KONTEKST).nalazi if n.pravilo_id == pravilo_id]


def _kljucevi(sadrzaj: bytes) -> set[str]:
    return {n.poruka_kljuc.removeprefix(f"{ID}.") for n in _nalazi(sadrzaj)}


def _osnovni(d):
    return [o for o in d.paragraphs if o.text.startswith(("Ovo je osnovni", "Drugi odlomak"))]


def _naslov(d):
    return next(o for o in d.paragraphs if o.style.name == "Heading 1")


def _sazetak(d):
    odlomci = d.paragraphs
    naslov = next(i for i, o in enumerate(odlomci) if o.text == PROFIL.uloge.sazetak)
    return odlomci[naslov + 1]


# --- matrica iz specifikacije ---------------------------------------------------


def test_ispravan_rad_daje_jedan_prolaz():
    [nalaz] = _nalazi(ispravan(54, PROFIL))
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.PROSLO, f"{ID}.prolazi")


def test_ispravan_font_rucno():
    sadrzaj = VARIJANTE["stilovi_font"](54, PROFIL)
    nalazi = _nalazi(sadrzaj)
    assert {n.poruka_kljuc for n in nalazi} == {f"{ID}.izravni_font", f"{ID}.izravna_velicina"}
    assert all(n.status is Status.NE_PROLAZI and n.izvor.vrsta is VrstaIzvora.IZRAVNO for n in nalazi)
    po_kljucu = {n.poruka_kljuc: n for n in nalazi}
    assert po_kljucu[f"{ID}.izravni_font"].pronadeno.iznos == "Calibri"
    assert tekst_nalaza(po_kljucu[f"{ID}.izravni_font"]).startswith("Font je zadan ručno")
    assert "(12 pt)" in tekst_nalaza(po_kljucu[f"{ID}.izravna_velicina"])
    assert status_pravila(provjeri(sadrzaj, PROFIL, KONTEKST).nalazi, "R-TEKST") is Status.PROSLO


def test_ispravno_poravnanje_rucno():
    [nalaz] = _nalazi(VARIJANTE["stilovi_poravnanje"](54, PROFIL))
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, f"{ID}.izravno_poravnanje")
    assert nalaz.pronadeno.iznos == "obostrano"
    assert "obostrano (Justify)" in tekst_nalaza(nalaz)


def test_naslov_rucno():
    [nalaz] = _nalazi(VARIJANTE["stilovi_naslov"](54, PROFIL))
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.NE_PROLAZI, f"{ID}.izravni_razmak")
    assert "Razmak ispred odlomka (6 pt)" in tekst_nalaza(nalaz)
    assert nalaz.lokacija.isjecak == PROFIL.uloge.poglavlja[0]


def test_podebljana_rijec_prolazi():
    [nalaz] = _nalazi(VARIJANTE["stilovi_podebljano"](54, PROFIL))
    assert nalaz.status is Status.PROSLO


def test_isti_razmak_u_40_odlomaka_daje_jedan_nalaz_s_primjerima():
    def promjena(d):
        prvi = _osnovni(d)[0]
        for i in range(38):
            prvi.insert_paragraph_before(f"Dodatni odlomak {i}.")
        for o in _osnovni(d)[:1] + [o for o in d.paragraphs if o.text.startswith("Dodatni")] + _osnovni(d)[1:2]:
            o.paragraph_format.space_after = Pt(6)

    [nalaz] = _nalazi(_izmijeni(promjena))
    assert (nalaz.poruka_kljuc, nalaz.pojave) == (f"{ID}.izravni_razmak", 40)
    assert len(nalaz.primjeri) == 3
    assert nalaz.primjeri[0] == nalaz.lokacija
    assert len({p.odlomak for p in nalaz.primjeri}) == 3


def test_sazetak_rucno_se_ne_provjerava():
    def promjena(d):
        _sazetak(d).paragraph_format.first_line_indent = Cm(1)

    [nalaz] = _nalazi(_izmijeni(promjena))
    assert nalaz.status is Status.PROSLO


# --- svojstva i rubni slučajevi -------------------------------------------------


@pytest.mark.parametrize(
    ("promjena", "kljuc"),
    [
        (lambda o: setattr(o.paragraph_format, "space_before", Pt(6)), "izravni_razmak"),
        (lambda o: setattr(o.paragraph_format, "line_spacing", 1.0), "izravni_prored"),
        (lambda o: setattr(o.paragraph_format, "line_spacing", Pt(14)), "izravni_prored"),
        (lambda o: setattr(o.paragraph_format, "left_indent", Cm(1)), "izravna_uvlaka"),
        (lambda o: setattr(o.paragraph_format, "right_indent", Cm(1)), "izravna_uvlaka"),
        (lambda o: setattr(o.paragraph_format, "first_line_indent", Cm(1)), "izravna_uvlaka"),
        (lambda o: setattr(o.paragraph_format, "alignment", WD_ALIGN_PARAGRAPH.LEFT), "izravno_poravnanje"),
    ],
)
def test_svako_svojstvo_odlomka(promjena, kljuc):
    assert _kljucevi(_izmijeni(lambda d: promjena(_osnovni(d)[0]))) == {kljuc}


def test_razmak_ispred_i_iza_su_zasebni_nalazi_sa_stranom():
    def promjena(d):
        for o in _osnovni(d):
            oblik = o.paragraph_format
            oblik.space_before = oblik.space_after = Pt(6)

    nalazi = _nalazi(_izmijeni(promjena))
    broj_osnovnih = len(_osnovni(python_docx.Document(io.BytesIO(ispravan(54, PROFIL)))))
    assert sorted((n.parametri["strana"], n.pojave) for n in nalazi) == [
        ("ispred", broj_osnovnih), ("iza", broj_osnovnih)]
    tekstovi = sorted(tekst_nalaza(n) for n in nalazi)
    assert tekstovi[0].startswith("Razmak ispred odlomka (6 pt)")
    assert tekstovi[1].startswith("Razmak iza odlomka (6 pt)")


def test_uvlake_navode_stranu():
    def promjena(d):
        oblik = _osnovni(d)[0].paragraph_format
        oblik.left_indent, oblik.right_indent, oblik.first_line_indent = Cm(1), Cm(1), Cm(1)

    nalazi = _nalazi(_izmijeni(promjena))
    assert sorted(tekst_nalaza(n).split(")")[0] for n in nalazi) == [
        "Uvlaka odlomka (desno: 1 cm", "Uvlaka odlomka (lijevo: 1 cm", "Uvlaka odlomka (prvi redak: 1 cm"]


@pytest.mark.parametrize(("prored", "prikaz"), [(1.0, KATALOG["prored.jednostruko"]), (Pt(14), "14 pt")])
def test_prikaz_proreda(prored, prikaz):
    def promjena(d):
        _osnovni(d)[0].paragraph_format.line_spacing = prored

    [nalaz] = _nalazi(_izmijeni(promjena))
    assert f"Prored ({prikaz})" in tekst_nalaza(nalaz)


# --- font teme zadan izravno na runu ---------------------------------------------

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _font_teme_na_runu(d):
    """Kao odabir „Calibri (Body)“ s popisa fontova: referenca na temu u w:rPr runa."""
    rfonts = _osnovni(d)[0].runs[0]._r.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(f"{W}asciiTheme", "minorHAnsi")
    rfonts.set(f"{W}hAnsiTheme", "minorHAnsi")


def _tema_bez_fonta_tijela(xml: bytes) -> bytes:
    return re.sub(rb'(<a:minorFont>\s*<a:latin typeface=")[^"]*"', rb'\g<1>"', xml)


def test_font_teme_zadan_izravno_je_rucno_oblikovanje():
    sadrzaj = zamijeni_u_zipu(_izmijeni(_font_teme_na_runu), {"word/theme/theme1.xml": tema_s_fontom("Calibri")})
    o = next(x for x in model.procitaj(sadrzaj, PROFIL.uloge).odlomci if x.uloga is UlogaOdlomka.OSNOVNI_TEKST)
    assert (o.runovi[0].font.vrijednost, o.runovi[0].font.izvor.vrsta) == ("Calibri", VrstaIzvora.IZRAVNO)
    [nalaz] = _nalazi(sadrzaj)
    assert (nalaz.poruka_kljuc, nalaz.pronadeno.iznos) == (f"{ID}.izravni_font", "Calibri")
    assert nalaz.izvor.vrsta is VrstaIzvora.IZRAVNO
    assert status_pravila(provjeri(sadrzaj, PROFIL, KONTEKST).nalazi, "R-TEKST") is Status.PROSLO


def test_font_teme_zadan_izravno_bez_imena_u_temi():
    sadrzaj = zamijeni_u_zipu(_izmijeni(_font_teme_na_runu), {"word/theme/theme1.xml": _tema_bez_fonta_tijela})
    [nalaz] = _nalazi(sadrzaj)
    assert (nalaz.status, nalaz.poruka_kljuc, nalaz.pronadeno) == (Status.NE_PROLAZI, f"{ID}.izravni_font", None)
    assert "{" not in tekst_nalaza(nalaz)


def test_font_teme_iz_stila_ostaje_tema():
    def bez_fonta(d):
        d.styles["Normal"].font.name = None

    m = model.procitaj(_izmijeni(bez_fonta), PROFIL.uloge)
    o = next(x for x in m.odlomci if x.uloga is UlogaOdlomka.OSNOVNI_TEKST)
    assert o.runovi[0].font.izvor.vrsta is VrstaIzvora.TEMA


def test_vise_runova_s_istim_fontom_u_odlomku_daje_jedan_nalaz():
    def promjena(d):
        odlomak = _osnovni(d)[0]
        for tekst in (" jedan", " dva"):
            odlomak.add_run(tekst).font.name = "Calibri"
        odlomak.runs[0].font.name = "Calibri"

    [nalaz] = _nalazi(_izmijeni(promjena))
    assert (nalaz.poruka_kljuc, nalaz.pojave) == (f"{ID}.izravni_font", 1)


def test_run_bez_vidljivog_teksta_se_ne_provjerava():
    def promjena(d):
        d_run = _osnovni(d)[0].add_run("   ")
        d_run.font.name = "Arial"
        d_run.font.size = Pt(20)

    assert _kljucevi(_izmijeni(promjena)) == {"prolazi"}


def test_naslov_literature_se_provjerava():
    def promjena(d):
        naslov = next(o for o in d.paragraphs if o.text == PROFIL.uloge.literatura)
        naslov.runs[0].font.size = Pt(12)

    assert _kljucevi(_izmijeni(promjena)) == {"izravna_velicina"}


def test_literatura_se_ne_provjerava():
    def promjena(d):
        stavka = next(o for o in d.paragraphs if o.style.name == "List Paragraph")
        stavka.paragraph_format.left_indent = Twips(0)

    assert _kljucevi(_izmijeni(promjena)) == {"prolazi"}


def test_bez_odlomaka_zadanih_uloga_prolazi():
    dokument = python_docx.Document()
    dokument.add_paragraph("Ana Horvat")
    spremnik = io.BytesIO()
    dokument.save(spremnik)
    [nalaz] = _nalazi(spremnik.getvalue())
    assert nalaz.status is Status.PROSLO


def test_ispravan_rad_ima_sve_provjeravane_uloge():
    """Pravilo filtrira samo po ulozi; ispravan rad mora imati sve tri uloge."""
    m = model.procitaj(ispravan(54, PROFIL), PROFIL.uloge)
    uloge = {o.uloga for o in m.odlomci}
    assert {UlogaOdlomka.OSNOVNI_TEKST, UlogaOdlomka.NASLOV_POGLAVLJA, UlogaOdlomka.NASLOV_LITERATURE} <= uloge


# --- poruke i profil ------------------------------------------------------------


def test_poruke_imaju_uputu_i_kljuceve():
    pravilo = REGISTAR[ID]
    assert {k.removeprefix(f"{ID}.") for k in pravilo.poruke} == KLJUCEVI | {"prolazi"}
    for kljuc in KLJUCEVI:
        tekst = KATALOG[f"{ID}.{kljuc}"]
        assert "Ctrl+Space" in tekst and "Ctrl+Q" in tekst and "Stilovi (Styles)" in tekst
    assert f"pravilo.{ID}" in KATALOG


def test_profil_ima_samo_ukljuceno():
    postavke = PROFIL.pravila[ID]
    assert postavke.ukljuceno is True
    assert set(type(postavke).model_fields) == {"ukljuceno"}


def test_pravilo_se_moze_iskljuciti():
    tekst = TOML.split("[pravila.R-STILOVI]")[0] + "[pravila.R-STILOVI]\nukljuceno = false\n"
    iskljuceno = profil.ucitaj(tekst.encode("utf-8"))
    izvjestaj = provjeri(VARIJANTE["stilovi_font"](54, PROFIL), iskljuceno, KONTEKST)
    assert not [n for n in izvjestaj.nalazi if n.pravilo_id == ID]

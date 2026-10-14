"""ISVU popis studenata (priča 3.1): adapter ``provjera.adapteri.popis``."""

import io
import logging
from pathlib import Path

import openpyxl
import pytest
from alati.isvu import ISPISAN, ISPRAVNI, S_GRESKAMA, UPISAN, ZAGLAVLJE, fixturei, izvoz

from provjera.adapteri.popis import GreskaRetka, Student, UcitaniPopis, ucitaj_isvu
from provjera.jezgra.poruke import KATALOG, poruka

FIXTURES = Path(__file__).resolve().parent / "fixtures"
PRVI_REDAK = 17  # 14 metapodataka, prazan redak, zaglavlje u 16. retku


def _kljucevi(popis: UcitaniPopis) -> list[tuple[int | None, str]]:
    return [(g.redak, g.poruka_kljuc) for g in popis.greske]


def test_fixturei_su_azurni():
    for naziv, sadrzaj in fixturei().items():
        # Sadržaj .xlsx-a nije bajtno stabilan (vrijeme u metapodacima), pa se uspoređuje rezultat.
        assert ucitaj_isvu((FIXTURES / naziv).read_bytes()) == ucitaj_isvu(sadrzaj)


def test_stvarni_oblik_bez_gresaka():
    popis = ucitaj_isvu((FIXTURES / "isvu_ispravan.xlsx").read_bytes())
    assert popis.greske == ()
    assert popis.moze_zamrznuti
    assert [s.rbr for s in popis.studenti] == [1, 2, 3, 4, 5]
    assert popis.studenti[0] == Student(1, "0036123456", "Ana Marija", "Horvat-Kovač", True)
    assert popis.broj_neaktivnih == 1


def test_zaglavlje_bez_metapodataka_i_drugaciji_redoslijed():
    zaglavlje = ("Status upisa", " prezime ", "IME", "jmbag", "RBR.")
    sadrzaj = izvoz(ISPRAVNI[:1], metapodaci=False, zaglavlje=zaglavlje)
    popis = ucitaj_isvu(sadrzaj)
    assert popis.greske == ()
    assert popis.studenti[0].jmbag == "0036123456"


@pytest.mark.parametrize(
    "jmbag", ["0036123456", 36123456, 36123456.0, " 0036123456 ", "36123456", "0036 123 456"]
)
def test_jmbag_zadrzava_vodece_nule(jmbag):
    popis = ucitaj_isvu(izvoz([("1.", jmbag, "Ana", "Horvat", UPISAN)]))
    assert popis.greske == ()
    assert popis.studenti[0].jmbag == "0036123456"


@pytest.mark.parametrize("rbr", ["78.", "78", " 078 ", 78, 78.0])
def test_rbr_tekst_i_broj(rbr):
    popis = ucitaj_isvu(izvoz([(rbr, "0036000001", "Ana", "Horvat", UPISAN)]))
    assert popis.greske == ()
    assert popis.studenti[0].rbr == 78


def test_razmaci_i_nfc_u_imenu():
    rastavljeno = "S\u030cimek"  # Š kao S + kvačica (NFD)
    popis = ucitaj_isvu(izvoz([("1.", "0036000001", "  Ana   Marija ", rastavljeno, UPISAN)]))
    student = popis.studenti[0]
    assert student.ime == "Ana Marija"
    assert student.prezime == "Šimek"


@pytest.mark.parametrize(
    ("status", "aktivan"),
    [
        (ISPISAN, False),
        ("ISPISAN", False),
        ("  Ispisan s predmeta", False),
        (UPISAN, True),
        (None, True),
    ],
)
def test_ispisan_je_neaktivan(status, aktivan):
    popis = ucitaj_isvu(izvoz([("1.", "0036000001", "Ana", "Horvat", status)]))
    assert popis.studenti[0].aktivan is aktivan


def test_greske_navode_brojeve_redaka():
    popis = ucitaj_isvu(izvoz(S_GRESKAMA))
    r = PRVI_REDAK
    assert _kljucevi(popis) == [
        (r + 1, "popis.rbr_neispravan"),  # „a“
        (r + 2, "popis.rbr_neispravan"),  # „0“
        (r + 4, "popis.dupli_rbr"),
        (r + 5, "popis.dupli_jmbag"),
        (r + 6, "popis.prazan_jmbag"),
        (r + 7, "popis.prazno_ime"),
        (r + 8, "popis.prazno_prezime"),
    ]
    dupli_rbr, dupli_jmbag = popis.greske[2], popis.greske[3]
    assert dict(dupli_rbr.parametri) == {"rbr": 5, "prvi": r + 3}
    assert dict(dupli_jmbag.parametri) == {"prvi": r}
    assert not popis.moze_zamrznuti
    # Retci bez grešaka ostaju u popisu.
    assert [s.rbr for s in popis.studenti] == [1, 5]


@pytest.mark.parametrize("rbr", ["a", "0", "-3", "5,5", "", 0, 5.5, True])
def test_nevaljan_rbr(rbr):
    popis = ucitaj_isvu(izvoz([(rbr, "0036000001", "Ana", "Horvat", UPISAN)]))
    assert _kljucevi(popis) == [(PRVI_REDAK, "popis.rbr_neispravan")]


def test_prazni_retci_se_preskacu_a_brojevi_redaka_ostaju_excelovi():
    prazan = (None, None, None, None, None)
    retci = [ISPRAVNI[0], prazan, ("  ", None, "", None, None), ("2.", "", "Ivo", "Babić", UPISAN)]
    popis = ucitaj_isvu(izvoz(retci))
    assert _kljucevi(popis) == [(PRVI_REDAK + 3, "popis.prazan_jmbag")]


def test_bez_zaglavlja_navodi_stupce_koji_nedostaju():
    zaglavlje = tuple(s for s in ZAGLAVLJE if s != "JMBAG")
    popis = ucitaj_isvu(izvoz(ISPRAVNI, zaglavlje=zaglavlje))
    assert popis.studenti == ()
    assert popis.greske == (
        GreskaRetka(None, "popis.bez_zaglavlja", {"stupci": "JMBAG", "najvise": 50}),
    )
    assert "JMBAG" in poruka("popis.bez_zaglavlja", **popis.greske[0].parametri)
    assert not popis.moze_zamrznuti


def test_zaglavlje_ispod_50_redaka_se_ne_trazi():
    sadrzaj = izvoz([], metapodaci=False)
    knjiga = openpyxl.load_workbook(io.BytesIO(sadrzaj))
    knjiga.active.insert_rows(1, 50)
    izlaz = io.BytesIO()
    knjiga.save(izlaz)
    popis = ucitaj_isvu(izlaz.getvalue())
    assert _kljucevi(popis) == [(None, "popis.bez_zaglavlja")]


@pytest.mark.parametrize("sadrzaj", [b"", b"nije xlsx", b"PK\x03\x04pokvareno"])
def test_necitljiva_datoteka(sadrzaj):
    popis = ucitaj_isvu(sadrzaj)
    assert popis == UcitaniPopis((), (GreskaRetka(None, "popis.necitljiv"),))
    assert not popis.moze_zamrznuti


def test_popis_bez_studenata_nije_zamrzljiv():
    popis = ucitaj_isvu(izvoz([]))
    assert _kljucevi(popis) == [(None, "popis.bez_studenata")]


def test_svaka_greska_ima_poruku():
    popis = ucitaj_isvu(izvoz(S_GRESKAMA))
    kljucevi = {g.poruka_kljuc for g in popis.greske} | {
        "popis.necitljiv",
        "popis.bez_zaglavlja",
        "popis.bez_studenata",
    }
    for g in popis.greske:
        assert poruka(g.poruka_kljuc, **g.parametri)
    assert kljucevi <= KATALOG.keys()


def test_log_ne_sadrzi_osobne_podatke(caplog):
    with caplog.at_level(logging.DEBUG):
        ucitaj_isvu(izvoz(ISPRAVNI))
        ucitaj_isvu(izvoz(S_GRESKAMA))
    for _rbr, jmbag, ime, prezime, _status in ISPRAVNI + S_GRESKAMA:
        for podatak in (jmbag, ime, prezime):
            if str(podatak).strip():
                assert str(podatak) not in caplog.text


@pytest.mark.parametrize("jmbag", ["abc", 36123456.5, "00361234567", 12345678901, "0036-12345", -5])
def test_nevaljan_jmbag(jmbag):
    popis = ucitaj_isvu(izvoz([("1.", jmbag, "Ana", "Horvat", UPISAN)]))
    assert _kljucevi(popis) == [(PRVI_REDAK, "popis.jmbag_neispravan")]


def test_isti_jmbag_kao_tekst_i_broj_je_duplikat():
    retci = [("1.", "0036123456", "Ana", "Horvat", UPISAN), ("2.", 36123456, "Ivo", "Babić", UPISAN)]
    popis = ucitaj_isvu(izvoz(retci))
    assert _kljucevi(popis) == [(PRVI_REDAK + 1, "popis.dupli_jmbag")]


def _knjiga(*listovi: list[tuple[object, ...]], aktivan: int = 0) -> bytes:
    knjiga = openpyxl.Workbook()
    knjiga.remove(knjiga.active)
    for i, retci in enumerate(listovi):
        list_ = knjiga.create_sheet(f"List{i + 1}")
        for redak in retci:
            list_.append(redak)
    knjiga.active = aktivan
    izlaz = io.BytesIO()
    knjiga.save(izlaz)
    return izlaz.getvalue()


def test_redak_zbroja_samo_u_drugim_stupcima_se_preskace():
    zaglavlje = ("Rbr.", "JMBAG", "Ime", "Prezime", "Status upisa", "Napomena")
    sadrzaj = _knjiga(
        [zaglavlje, ("1.", "0036000001", "Ana", "Horvat", UPISAN, None), (None, " ", None, "", None, "Ukupno: 1")]
    )
    popis = ucitaj_isvu(sadrzaj)
    assert popis.greske == ()
    assert len(popis.studenti) == 1


@pytest.mark.parametrize("status", ["Status  upisa", "Status\u00a0upisa", " status upisa "])
def test_zaglavlje_s_visestrukim_razmacima_i_nbsp(status):
    sadrzaj = _knjiga([("Rbr.", "JMBAG", "Ime", "Prezime", status), ISPRAVNI[0]])
    popis = ucitaj_isvu(sadrzaj)
    assert popis.greske == ()
    assert popis.studenti[0].jmbag == "0036123456"


def test_zaglavlje_na_listu_koji_nije_aktivan():
    sadrzaj = _knjiga(
        [("Napomena",), ("nema popisa",)],
        [("Rbr.", "JMBAG", "Ime", "Prezime", "Status upisa"), ISPRAVNI[0]],
        [("Rbr.", "JMBAG", "Ime", "Prezime", "Status upisa"), ISPRAVNI[1]],
        aktivan=0,
    )
    popis = ucitaj_isvu(sadrzaj)
    assert popis.greske == ()
    assert [s.rbr for s in popis.studenti] == [1]  # prvi list sa zaglavljem


def test_parametri_greske_su_nepromjenjivi():
    izvorni = {"rbr": 5}
    greska = GreskaRetka(3, "popis.dupli_rbr", izvorni)
    izvorni["rbr"] = 6
    assert greska.parametri["rbr"] == 5
    with pytest.raises(TypeError):
        greska.parametri["rbr"] = 7  # type: ignore[index]

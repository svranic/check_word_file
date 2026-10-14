"""Generator sintetskog referentnog skupa Vježbe 1: ispravan rad i pokvarene varijante.

Pokretanje: ``uv run python tests/alati/generator.py`` zapisuje .docx primjere u
``tests/referentni/vjezba_1/``. Imena u dokumentima su izmišljena.
"""

import io
import re
import zipfile
from collections.abc import Callable
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENTATION, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, Twips

from provjera.jezgra import formula
from provjera.jezgra.jedinice import mm_u_twips
from provjera.jezgra.profil import Profil, ucitaj

KORIJEN = Path(__file__).resolve().parents[2]
MAPA = KORIJEN / "tests" / "referentni" / "vjezba_1"


def _a4(sekcija) -> None:
    sekcija.orientation = WD_ORIENTATION.PORTRAIT
    sekcija.page_width = Mm(210)
    sekcija.page_height = Mm(297)
    sekcija.gutter = 0


def _margine(sekcija, gore: int, dolje: int, lijevo: int, desno: int) -> None:
    """Margine u twipsima, točno kako ih Word zapisuje."""
    sekcija.top_margin = Twips(gore)
    sekcija.bottom_margin = Twips(dolje)
    sekcija.left_margin = Twips(lijevo)
    sekcija.right_margin = Twips(desno)


def _spremi(dokument) -> bytes:
    spremnik = io.BytesIO()
    dokument.save(spremnik)
    return spremnik.getvalue()


def _tijelo(rbr: int, profil: Profil) -> dict[str, int]:
    return formula.margine(rbr, profil.formula).u_twips()


def zamijeni_u_zipu(sadrzaj: bytes, zamjene: dict[str, Callable[[bytes], bytes | None]]) -> bytes:
    """Prepisuje dijelove paketa; funkcija koja vrati ``None`` uklanja dio."""
    ulaz, izlaz = zipfile.ZipFile(io.BytesIO(sadrzaj)), io.BytesIO()
    with zipfile.ZipFile(izlaz, "w", zipfile.ZIP_DEFLATED) as novi:
        for stavka in ulaz.infolist():
            podaci = ulaz.read(stavka)
            if stavka.filename in zamjene:
                podaci = zamjene[stavka.filename](podaci)
                if podaci is None:
                    continue
            novi.writestr(stavka, podaci)
    return izlaz.getvalue()


def _normal(dokument, profil: Profil) -> None:
    """Stil Normal točno po [pravila.R-TEKST]; font je zapisan u stilu, ne preuzet iz teme."""
    tekst = profil.pravila["R-TEKST"]
    normal = dokument.styles["Normal"]
    normal.font.name = tekst.font
    normal.font.size = Pt(tekst.velicina / 2)
    oblik = normal.paragraph_format
    oblik.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    oblik.space_before = Twips(tekst.razmak_ispred)
    oblik.space_after = Twips(tekst.razmak_iza)
    oblik.line_spacing = tekst.prored / 240


# Numeracija naslova poglavlja („1.“–„4.“) kroz stil Heading 1, kao Višerazinski popis u Wordu.
NUM_ID_NASLOVA = "90"


def _numeracija_naslova(dokument) -> None:
    numbering = dokument.part.numbering_part.element
    apstraktni = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{NUM_ID_NASLOVA}">'
        '<w:multiLevelType w:val="multilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading1"/>'
        '<w:lvlText w:val="%1."/><w:lvlJc w:val="left"/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading2"/>'
        '<w:lvlText w:val="%1.%2."/><w:lvlJc w:val="left"/></w:lvl>'
        "</w:abstractNum>"
    )
    # Svi w:abstractNum moraju biti ispred prvog w:num.
    prvi_num = numbering.find(qn("w:num"))
    if prvi_num is not None:
        prvi_num.addprevious(apstraktni)
    else:
        numbering.append(apstraktni)
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{NUM_ID_NASLOVA}"><w:abstractNumId w:val="{NUM_ID_NASLOVA}"/></w:num>'
    ))


def _heading1(dokument, profil: Profil) -> None:
    """Stil Heading 1 točno po [pravila.R-NASLOVI], numeriran kroz stil (bez ilvl, preko pStyle)."""
    naslovi = profil.pravila["R-NASLOVI"]
    stil = dokument.styles["Heading 1"]
    rpr = stil.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    # Font iz teme (majorHAnsi) imao bi prednost pred imenom fonta.
    for atribut in ("w:asciiTheme", "w:hAnsiTheme"):
        rfonts.attrib.pop(qn(atribut), None)
    boja = rpr.find(qn("w:color"))
    if boja is not None:
        rpr.remove(boja)
    stil.font.name = naslovi.font
    stil.font.bold = naslovi.podebljano
    stil.font.size = Pt(naslovi.velicina / 2)
    oblik = stil.paragraph_format
    oblik.alignment = WD_ALIGN_PARAGRAPH.LEFT
    oblik.space_before = Twips(naslovi.razmak_ispred)
    oblik.space_after = Twips(naslovi.razmak_iza)
    _numeracija_naslova(dokument)
    stil.element.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:numId w:val="{NUM_ID_NASLOVA}"/></w:numPr>'
    ))
    # Shema traži numPr ispred spacing/jc: python-docx ga ne slaže sam.
    ppr = stil.element.pPr
    numpr = ppr.find(qn("w:numPr"))
    ppr.remove(numpr)
    sljedeci = next((e for e in ppr if e.tag in {qn("w:spacing"), qn("w:jc"), qn("w:outlineLvl")}), None)
    if sljedeci is not None:
        sljedeci.addprevious(numpr)
    else:
        ppr.append(numpr)


def _sadrzaj(dokument, profil: Profil) -> None:
    """Sažetak, poglavlja u stilu Heading 1 s osnovnim tekstom i stavka literature.

    Sažetak je u stilu Normal s izravnom uvlakom po [pravila.R-SAZETAK].
    """
    sazetak = profil.pravila["R-SAZETAK"]
    dokument.add_paragraph(profil.uloge.sazetak)
    tekst = dokument.add_paragraph("Rad ukratko opisuje izbor studija i ciljeve studiranja.")
    tekst.paragraph_format.left_indent = Twips(sazetak.uvlaka_lijevo)
    tekst.paragraph_format.right_indent = Twips(sazetak.uvlaka_desno)
    for poglavlje in profil.uloge.poglavlja:
        dokument.add_heading(poglavlje, level=1)
        if poglavlje == profil.uloge.literatura:
            dokument.add_paragraph("Sveučilište Sjever, https://www.unin.hr/", style="List Paragraph")
        else:
            dokument.add_paragraph(f"Ovo je osnovni tekst poglavlja {poglavlje}.")
            dokument.add_paragraph("Drugi odlomak osnovnog teksta istog poglavlja.")


def _footer(dokument, profil: Profil) -> None:
    """Stil Footer točno po [pravila.R-BROJ-STRANICE]."""
    broj = profil.pravila["R-BROJ-STRANICE"]
    stil = dokument.styles["Footer"]
    stil.font.name = broj.font
    stil.font.size = Pt(broj.velicina / 2)


def polje_page(odlomak, naredba: str = r" PAGE   \* MERGEFORMAT ", rezultat: str = "2", rpr: str = "") -> None:
    """Složeno polje (fldChar/instrText) kako ga Word zapisuje; ``rpr`` je XML za run rezultata."""
    w = nsdecls("w")
    for xml in (
        f'<w:r {w}><w:fldChar w:fldCharType="begin"/></w:r>',
        f'<w:r {w}><w:instrText xml:space="preserve">{naredba}</w:instrText></w:r>',
        f'<w:r {w}><w:fldChar w:fldCharType="separate"/></w:r>',
        f'<w:r {w}>{rpr}<w:t>{rezultat}</w:t></w:r>',
        f'<w:r {w}><w:fldChar w:fldCharType="end"/></w:r>',
    ):
        odlomak._p.append(parse_xml(xml))


def podnozje_s_brojem(sekcija):
    """Vlastito podnožje sekcije (nije povezano s prethodnim) s poljem PAGE u stilu Footer."""
    podnozje = sekcija.footer
    podnozje.is_linked_to_previous = False
    odlomak = podnozje.paragraphs[0]
    odlomak.alignment = WD_ALIGN_PARAGRAPH.CENTER
    polje_page(odlomak)
    return podnozje


def _dokument(rbr: int, profil: Profil, jedna_sekcija: bool = False):
    dokument = docx.Document()
    _normal(dokument, profil)
    _heading1(dokument, profil)
    _footer(dokument, profil)
    naslovnica = dokument.sections[0]
    _a4(naslovnica)
    tijelo = _tijelo(rbr, profil)
    if jedna_sekcija:
        _margine(naslovnica, **tijelo)
        dokument.add_paragraph("Ana Horvat, 0000000000")
        _sadrzaj(dokument, profil)
        podnozje_s_brojem(naslovnica)
        return dokument
    _margine(naslovnica, *(mm_u_twips(profil.formula.naslovnica),) * 4)
    dokument.add_paragraph("Ana Horvat, 0000000000")
    sekcija = dokument.add_section(WD_SECTION_START.NEW_PAGE)
    _a4(sekcija)
    _margine(sekcija, **tijelo)
    _sadrzaj(dokument, profil)
    # Naslovnica nema podnožje; tijelo ima vlastito s brojem stranice.
    podnozje_s_brojem(sekcija)
    return dokument


def ispravan(rbr: int, profil: Profil) -> bytes:
    return _spremi(_dokument(rbr, profil))


def _izmijeni(promjena: Callable) -> Callable[[int, Profil], bytes]:
    def varijanta(rbr: int, profil: Profil) -> bytes:
        dokument = _dokument(rbr, profil)
        promjena(dokument, rbr, profil)
        return _spremi(dokument)

    return varijanta


def _neprekinuto(dokument, rbr, profil):
    dokument.sections[1].start_type = WD_SECTION_START.CONTINUOUS


def _krive_margine_tijela(dokument, rbr, profil):
    dokument.sections[1].left_margin = Twips(_tijelo(rbr, profil)["lijevo"] + mm_u_twips(1))


def _krive_margine_naslovnice(dokument, rbr, profil):
    dokument.sections[0].top_margin = Mm(30)


def _unutar_tolerancije(dokument, rbr, profil):
    tolerancija = profil.tolerancije.margine
    dokument.sections[1].right_margin = Twips(_tijelo(rbr, profil)["desno"] + tolerancija)


def _letter(dokument, rbr, profil):
    for sekcija in dokument.sections:
        sekcija.page_width, sekcija.page_height = Mm(215.9), Mm(279.4)


def _pejzazno(dokument, rbr, profil):
    sekcija = dokument.sections[1]
    sekcija.orientation = WD_ORIENTATION.LANDSCAPE
    sekcija.page_width, sekcija.page_height = sekcija.page_height, sekcija.page_width


def _gutter(dokument, rbr, profil):
    dokument.sections[1].gutter = Mm(5)


def _zrcalne(dokument, rbr, profil):
    dokument.settings.element.append(OxmlElement("w:mirrorMargins"))


def _tekst_sazetka(dokument, profil):
    """Odlomak teksta sažetka (iza naslova „Sažetak“)."""
    odlomci = dokument.paragraphs
    naslov = next(i for i, o in enumerate(odlomci) if o.text == profil.uloge.sazetak)
    return odlomci[naslov + 1]


def _zastiti_sazetak_i_literaturu(dokument, profil) -> None:
    """Varijanta koja kvari stil Normal smije pasti samo na R-TEKST.

    Sažetak (u Normal) i literatura (List Paragraph, nasljeđuje Normal) dobivaju ispravne
    vrijednosti izravno: smiju imati vlastito oblikovanje.
    """
    sazetak, literatura = profil.pravila["R-SAZETAK"], profil.pravila["R-LITERATURA"]
    odlomak = _tekst_sazetka(dokument, profil)
    odlomak.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in odlomak.runs:
        run.font.name = sazetak.font
        run.font.size = Pt(sazetak.velicina / 2)
    stil = dokument.styles["List Paragraph"]
    stil.font.name = literatura.font
    stil.font.size = Pt(literatura.velicina / 2)


def _tekst_font(dokument, rbr, profil):
    dokument.styles["Normal"].font.name = "Times New Roman"
    _zastiti_sazetak_i_literaturu(dokument, profil)


def _tekst_velicina(dokument, rbr, profil):
    dokument.styles["Normal"].font.size = Pt(11)
    _zastiti_sazetak_i_literaturu(dokument, profil)


def _tekst_poravnanje(dokument, rbr, profil):
    dokument.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _zastiti_sazetak_i_literaturu(dokument, profil)


def _tekst_razmak(dokument, rbr, profil):
    dokument.styles["Normal"].paragraph_format.space_after = Pt(10)


def _tekst_prored(dokument, rbr, profil):
    dokument.styles["Normal"].paragraph_format.line_spacing = 1.15


def _prvi_osnovni(dokument):
    """Prvi odlomak osnovnog teksta (iza prvog naslova poglavlja)."""
    odlomci = dokument.paragraphs
    prvi_naslov = next(i for i, o in enumerate(odlomci) if o.style.name == "Heading 1")
    return odlomci[prvi_naslov + 1]


def _tekst_drugi_stil(dokument, rbr, profil):
    """Prilagođeni stil s istim svojstvima kao Normal ipak nije Normal."""
    stil = dokument.styles.add_style("Moj tekst", WD_STYLE_TYPE.PARAGRAPH)
    stil.base_style = dokument.styles["Normal"]
    _prvi_osnovni(dokument).style = stil


def tema_s_fontom(font: str) -> Callable[[bytes], bytes]:
    def zamijeni(xml: bytes) -> bytes:
        return re.sub(rb'(<a:minorFont>\s*<a:latin typeface=")[^"]*', rb"\g<1>" + font.encode(), xml)

    return zamijeni


def _tekst_aptos(rbr: int, profil: Profil) -> bytes:
    """Normal bez fonta: Word uzima font teme (Aptos u novijim verzijama)."""
    dokument = _dokument(rbr, profil)
    dokument.styles["Normal"].font.name = None
    _zastiti_sazetak_i_literaturu(dokument, profil)
    return zamijeni_u_zipu(_spremi(dokument), {"word/theme/theme1.xml": tema_s_fontom("Aptos")})


def _naslovi(dokument):
    return [o for o in dokument.paragraphs if o.style.name == "Heading 1"]


def _bez_numeracije_stila(dokument) -> None:
    ppr = dokument.styles["Heading 1"].element.pPr
    ppr.remove(ppr.find(qn("w:numPr")))


def _naslovi_rucni_broj(dokument, rbr, profil):
    """Brojevi upisani u tekst naslova umjesto Wordove numeracije."""
    _bez_numeracije_stila(dokument)
    for broj, naslov in enumerate(_naslovi(dokument), start=1):
        naslov.runs[0].text = f"{broj}. {naslov.runs[0].text}"


def _naslovi_bez_numeracije(dokument, rbr, profil):
    _bez_numeracije_stila(dokument)


def _naslov(dokument, profil, indeks: int):
    naziv = profil.uloge.poglavlja[indeks]
    return next(o for o in _naslovi(dokument) if o.text == naziv)


def _naslovi_drugi_stil(dokument, rbr, profil):
    """„Zaključak“ izgleda kao naslov (podebljano), ali je u stilu Normal."""
    naslov = _naslov(dokument, profil, 2)
    naslov.style = dokument.styles["Normal"]
    naslov.runs[0].bold = True


def _naslovi_nedostaje(dokument, rbr, profil):
    naslov = _naslov(dokument, profil, 2)
    naslov._p.getparent().remove(naslov._p)


def _naslovi_font(dokument, rbr, profil):
    dokument.styles["Heading 1"].font.name = "Times New Roman"


def _naslovi_sazetak_numeriran(dokument, rbr, profil):
    sazetak = next(o for o in dokument.paragraphs if o.text == profil.uloge.sazetak)
    sazetak.style = dokument.styles["Heading 1"]


def _sazetak_bez_uvlake(dokument, rbr, profil):
    oblik = _tekst_sazetka(dokument, profil).paragraph_format
    oblik.left_indent = oblik.right_indent = None


def _sazetak_jedna_strana(dokument, rbr, profil):
    _tekst_sazetka(dokument, profil).paragraph_format.right_indent = None


def _sazetak_vlastiti_stil(dokument, rbr, profil):
    """Stil „Sažetak“ sa svim svojstvima; odlomak nema izravnu uvlaku."""
    sazetak = profil.pravila["R-SAZETAK"]
    stil = dokument.styles.add_style("Sažetak", WD_STYLE_TYPE.PARAGRAPH)
    stil.base_style = dokument.styles["Normal"]
    stil.font.name = sazetak.font
    stil.font.size = Pt(sazetak.velicina / 2)
    oblik = stil.paragraph_format
    oblik.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    oblik.left_indent = Twips(sazetak.uvlaka_lijevo)
    oblik.right_indent = Twips(sazetak.uvlaka_desno)
    odlomak = _tekst_sazetka(dokument, profil)
    odlomak.style = stil
    odlomak.paragraph_format.left_indent = odlomak.paragraph_format.right_indent = None


def _sazetak_nedostaje(dokument, rbr, profil):
    tekst = _tekst_sazetka(dokument, profil)
    naslov = next(o for o in dokument.paragraphs if o.text == profil.uloge.sazetak)
    for odlomak in (naslov, tekst):
        odlomak._p.getparent().remove(odlomak._p)


def _literatura_font(dokument, rbr, profil):
    dokument.styles["List Paragraph"].font.name = "Times New Roman"


def _stilovi_font(dokument, rbr, profil):
    """Ispravan font i veličina, ali zadani ručno na runu osnovnog teksta."""
    tekst = profil.pravila["R-TEKST"]
    for run in _prvi_osnovni(dokument).runs:
        run.font.name = tekst.font
        run.font.size = Pt(tekst.velicina / 2)


def _stilovi_poravnanje(dokument, rbr, profil):
    """Ispravno obostrano poravnanje (jc=both), ali zadano ručno na odlomku."""
    _prvi_osnovni(dokument).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _stilovi_naslov(dokument, rbr, profil):
    """Ispravan razmak 6 pt, ali zadan ručno na naslovu poglavlja."""
    naslov = _naslov(dokument, profil, 0)
    naslov.paragraph_format.space_before = Twips(profil.pravila["R-NASLOVI"].razmak_ispred)


def _stilovi_podebljano(dokument, rbr, profil):
    """Podebljana, nakošena i podcrtana riječ u osnovnom tekstu je dopuštena."""
    run = _prvi_osnovni(dokument).add_run(" važno")
    run.bold = run.italic = run.underline = True


def _isprazni(podnozje):
    """Briše sav sadržaj podnožja i vraća jedan prazan odlomak u stilu Footer."""
    korijen = podnozje._element
    for dijete in list(korijen):
        korijen.remove(dijete)
    return podnozje.add_paragraph(style="Footer")


def _broj_rucni(dokument, rbr, profil):
    """Broj „2“ upisan kao tekst umjesto polja."""
    _isprazni(dokument.sections[1].footer).add_run("2")


def _broj_nema(dokument, rbr, profil):
    dokument.sections[1].footer.is_linked_to_previous = True


def _broj_naslovnica(dokument, rbr, profil):
    """Broj samo u podnožju naslovnice; tijelo ga nasljeđuje (Poveži s prethodnim)."""
    podnozje_s_brojem(dokument.sections[0])
    dokument.sections[1].footer.is_linked_to_previous = True


def _broj_prva_stranica(dokument, rbr, profil):
    """Drukčija prva stranica u tijelu, a podnožje prve stranice je prazno."""
    sekcija = dokument.sections[1]
    sekcija.different_first_page_header_footer = True
    sekcija.first_page_footer.is_linked_to_previous = False


def _broj_font(dokument, rbr, profil):
    """Rezultat polja 11 pt (izravno na runu)."""
    odlomak = _isprazni(dokument.sections[1].footer)
    polje_page(odlomak, rpr=f'<w:rPr {nsdecls("w")}><w:sz w:val="22"/></w:rPr>')


def _broj_galerija(dokument, rbr, profil):
    """Broj iz Wordove galerije: odlomak s poljem unutar sdt „Page Numbers (Bottom of Page)“."""
    odlomak = dokument.sections[1].footer.paragraphs[0]._p
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:id w:val="1824306713"/><w:docPartObj>'
        '<w:docPartGallery w:val="Page Numbers (Bottom of Page)"/><w:docPartUnique/>'
        "</w:docPartObj></w:sdtPr><w:sdtContent/></w:sdt>"
    )
    odlomak.addprevious(sdt)
    sdt.find(qn("w:sdtContent")).append(odlomak)


_VML = "urn:schemas-microsoft-com:vml"


def _broj_okvir(dokument, rbr, profil):
    """Polje PAGE u tekstnom okviru (VML v:textbox) u podnožju."""
    podnozje = dokument.sections[1].footer
    unutarnji = podnozje.paragraphs[0]._p
    vanjski = _isprazni(podnozje)
    run = parse_xml(
        f'<w:r {nsdecls("w")} xmlns:v="{_VML}"><w:pict><v:shape id="Okvir1" type="#_x0000_t202" '
        'style="position:absolute;margin-left:0;margin-top:0;width:60pt;height:20pt;'
        'mso-position-horizontal:center">'
        "<v:textbox><w:txbxContent/></v:textbox></v:shape></w:pict></w:r>"
    )
    run.find(f".//{qn('w:txbxContent')}").append(unutarnji)
    vanjski._p.append(run)


VARIJANTE: dict[str, Callable[[int, Profil], bytes]] = {
    "ispravan": ispravan,
    "jedna_sekcija": lambda rbr, profil: _spremi(_dokument(rbr, profil, jedna_sekcija=True)),
    "neprekinuti_prijelom": _izmijeni(_neprekinuto),
    "krive_margine_tijela": _izmijeni(_krive_margine_tijela),
    "krive_margine_naslovnice": _izmijeni(_krive_margine_naslovnice),
    "unutar_tolerancije": _izmijeni(_unutar_tolerancije),
    "letter": _izmijeni(_letter),
    "pejzazno": _izmijeni(_pejzazno),
    "gutter": _izmijeni(_gutter),
    "zrcalne": _izmijeni(_zrcalne),
    "tekst_font": _izmijeni(_tekst_font),
    "tekst_velicina": _izmijeni(_tekst_velicina),
    "tekst_poravnanje": _izmijeni(_tekst_poravnanje),
    "tekst_razmak": _izmijeni(_tekst_razmak),
    "tekst_prored": _izmijeni(_tekst_prored),
    "tekst_drugi_stil": _izmijeni(_tekst_drugi_stil),
    "tekst_aptos": _tekst_aptos,
    "naslovi_rucni_broj": _izmijeni(_naslovi_rucni_broj),
    "naslovi_bez_numeracije": _izmijeni(_naslovi_bez_numeracije),
    "naslovi_drugi_stil": _izmijeni(_naslovi_drugi_stil),
    "naslovi_nedostaje": _izmijeni(_naslovi_nedostaje),
    "naslovi_font": _izmijeni(_naslovi_font),
    "naslovi_sazetak_numeriran": _izmijeni(_naslovi_sazetak_numeriran),
    "sazetak_bez_uvlake": _izmijeni(_sazetak_bez_uvlake),
    "sazetak_jedna_strana": _izmijeni(_sazetak_jedna_strana),
    "sazetak_vlastiti_stil": _izmijeni(_sazetak_vlastiti_stil),
    "sazetak_nedostaje": _izmijeni(_sazetak_nedostaje),
    "literatura_font": _izmijeni(_literatura_font),
    "stilovi_font": _izmijeni(_stilovi_font),
    "stilovi_poravnanje": _izmijeni(_stilovi_poravnanje),
    "stilovi_naslov": _izmijeni(_stilovi_naslov),
    "stilovi_podebljano": _izmijeni(_stilovi_podebljano),
    "broj_rucni": _izmijeni(_broj_rucni),
    "broj_nema": _izmijeni(_broj_nema),
    "broj_naslovnica": _izmijeni(_broj_naslovnica),
    "broj_prva_stranica": _izmijeni(_broj_prva_stranica),
    "broj_font": _izmijeni(_broj_font),
    "broj_galerija": _izmijeni(_broj_galerija),
    "broj_okvir": _izmijeni(_broj_okvir),
}


def zapisi(rbr: int = 54) -> list[Path]:
    profil = ucitaj((KORIJEN / "profili" / "vjezba_1.toml").read_bytes())
    MAPA.mkdir(parents=True, exist_ok=True)
    zapisane = []
    for ime, varijanta in VARIJANTE.items():
        putanja = MAPA / f"{ime}.docx"
        putanja.write_bytes(varijanta(rbr, profil))
        zapisane.append(putanja)
    return zapisane


if __name__ == "__main__":
    for putanja in zapisi():
        print(putanja.relative_to(KORIJEN).as_posix())

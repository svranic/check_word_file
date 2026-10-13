"""R-NASLOVI, numeracija i podebljanje u modelu (priča 2.2)."""

import io
import re
from pathlib import Path

import docx as python_docx
import pytest
from alati.generator import NUM_ID_NASLOVA, VARIJANTE, ispravan, zamijeni_u_zipu
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from provjera.jezgra import model, profil
from provjera.jezgra.model.numeracija import Brojaci, Numeracija, formatiraj_broj
from provjera.jezgra.model.uloge import UlogaOdlomka, ima_rucni_broj
from provjera.jezgra.nalaz import KontekstProvjere, Status, VrstaIzvora
from provjera.jezgra.poruke import KATALOG, tekst_nalaza
from provjera.jezgra.pravila import REGISTAR
from provjera.jezgra.provjera import provjeri, status_pravila

KORIJEN = Path(__file__).resolve().parents[1]
PROFIL = profil.ucitaj((KORIJEN / "profili" / "vjezba_1.toml").read_bytes())
KONTEKST = KontekstProvjere("Ana_Horvat_GiG_OG_Vjezba_1.docx", 54, None, None)
UZORAK = KORIJEN / "samples" / "Lucija_Šimek_GiG_OG_Vjezba_1.docx"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _izmijeni(promjena, sadrzaj: bytes | None = None) -> bytes:
    dokument = python_docx.Document(io.BytesIO(sadrzaj or ispravan(54, PROFIL)))
    promjena(dokument)
    spremnik = io.BytesIO()
    dokument.save(spremnik)
    return spremnik.getvalue()


def _nalazi(sadrzaj: bytes, p=PROFIL):
    return [n for n in provjeri(sadrzaj, p, KONTEKST).nalazi if n.pravilo_id == "R-NASLOVI"]


def _kljucevi(sadrzaj: bytes) -> set[str]:
    return {n.poruka_kljuc.removeprefix("R-NASLOVI.") for n in _nalazi(sadrzaj)}


def _naslovi(m):
    return [o for o in m.odlomci if o.uloga in (UlogaOdlomka.NASLOV_POGLAVLJA, UlogaOdlomka.NASLOV_LITERATURE)]


def _apstraktni(d):
    numbering = d.part.numbering_part.element
    return next(a for a in numbering.iterchildren(qn("w:abstractNum"))
                if a.get(qn("w:abstractNumId")) == NUM_ID_NASLOVA)


# --- matrica iz specifikacije ---------------------------------------------------


def test_ispravan_rad_daje_jedan_prolaz():
    [nalaz] = _nalazi(ispravan(54, PROFIL))
    assert (nalaz.status, nalaz.poruka_kljuc) == (Status.PROSLO, "R-NASLOVI.prolazi")
    m = model.procitaj(ispravan(54, PROFIL), PROFIL.uloge)
    assert [o.numeracija.vrijednost.oznaka for o in _naslovi(m)] == ["1.", "2.", "3.", "4."]
    assert all(o.numeracija.izvor.vrsta is VrstaIzvora.STIL for o in _naslovi(m))


def test_rucni_broj_zamjenjuje_bez_numeracije():
    nalazi = _nalazi(VARIJANTE["naslovi_rucni_broj"](54, PROFIL))
    assert {n.poruka_kljuc for n in nalazi} == {"R-NASLOVI.rucni_broj"}
    assert len(nalazi) == 4 and all(n.status is Status.NE_PROLAZI for n in nalazi)
    tekst = tekst_nalaza(nalazi[0])
    assert "Početna (Home) → Višerazinski popis (Multilevel List)" in tekst and "1. Uvod" in tekst


def test_bez_numeracije_je_jedan_sazeti_nalaz():
    [nalaz] = _nalazi(VARIJANTE["naslovi_bez_numeracije"](54, PROFIL))
    assert (nalaz.status, nalaz.poruka_kljuc, nalaz.pojave) == (Status.NE_PROLAZI, "R-NASLOVI.bez_numeracije", 4)
    assert "Višerazinski popis (Multilevel List)" in tekst_nalaza(nalaz)


def test_drugi_stil_objasnjava_stil_naspram_rucnog_oblikovanja():
    sadrzaj = VARIJANTE["naslovi_drugi_stil"](54, PROFIL)
    assert _kljucevi(sadrzaj) == {"drugi_stil"}
    nalazi = _nalazi(sadrzaj)
    [nalaz] = [n for n in nalazi if n.poruka_kljuc == "R-NASLOVI.drugi_stil"]
    assert nalaz.status is Status.NE_PROLAZI and nalaz.pronadeno.iznos == "Normal"
    assert nalaz.lokacija.isjecak == "Zaključak"
    tekst = tekst_nalaza(nalaz)
    assert "Ručno oblikovanje" in tekst and "Heading 1" in tekst and "{" not in tekst


def test_nedostaje_navodi_naziv():
    sadrzaj = VARIJANTE["naslovi_nedostaje"](54, PROFIL)
    assert _kljucevi(sadrzaj) == {"nedostaje"}
    nalazi = _nalazi(sadrzaj)
    [nalaz] = [n for n in nalazi if n.poruka_kljuc == "R-NASLOVI.nedostaje"]
    assert nalaz.parametri["naziv"] == "Zaključak"
    assert "„Zaključak“" in tekst_nalaza(nalaz)


def test_dva_nedostajuca_poglavlja_daju_dva_nalaza():
    def obrisi(d):
        for o in d.paragraphs:
            if o.text in ("Uvod", "Zaključak"):
                o._p.getparent().remove(o._p)

    nalazi = [n for n in _nalazi(_izmijeni(obrisi)) if n.poruka_kljuc == "R-NASLOVI.nedostaje"]
    assert sorted(n.parametri["naziv"] for n in nalazi) == ["Uvod", "Zaključak"]


@pytest.mark.parametrize(
    ("format_", "uzorak", "start", "oznake"),
    [
        ("lowerLetter", "%1)", "1", ["a)", "b)", "c)", "d)"]),
        ("decimal", "%1.", "2", ["2.", "3.", "4.", "5."]),
        ("upperRoman", "%1.", "1", ["I.", "II.", "III.", "IV."]),
        ("decimal", "%1", "1", ["1", "2", "3", "4"]),
    ],
)
def test_kriva_oznaka(format_, uzorak, start, oznake):
    def promijeni(d):
        lvl = _apstraktni(d).find(qn("w:lvl"))
        lvl.find(qn("w:numFmt")).set(qn("w:val"), format_)
        lvl.find(qn("w:lvlText")).set(qn("w:val"), uzorak)
        lvl.find(qn("w:start")).set(qn("w:val"), start)

    sadrzaj = _izmijeni(promijeni)
    assert [o.numeracija.vrijednost.oznaka for o in _naslovi(model.procitaj(sadrzaj, PROFIL.uloge))] == oznake
    nalazi = _nalazi(sadrzaj)
    assert {n.poruka_kljuc for n in nalazi} == {"R-NASLOVI.kriva_oznaka"}
    assert [(n.ocekivano.iznos, n.pronadeno.iznos) for n in nalazi] == list(zip(["1.", "2.", "3.", "4."], oznake))


def test_numeriran_sazetak_ne_prolazi():
    sadrzaj = VARIJANTE["naslovi_sazetak_numeriran"](54, PROFIL)
    assert _kljucevi(sadrzaj) == {"sazetak_numeriran"}
    nalazi = _nalazi(sadrzaj)
    [nalaz] = [n for n in nalazi if n.poruka_kljuc == "R-NASLOVI.sazetak_numeriran"]
    assert nalaz.lokacija.isjecak == "Sažetak" and nalaz.pronadeno.iznos == "1."
    assert "{" not in tekst_nalaza(nalaz)


def test_sazetak_u_heading1_bez_numeracije_prolazi():
    def sazetak(d):
        o = next(o for o in d.paragraphs if o.text == "Sažetak")
        o.style = d.styles["Heading 1"]
        o._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="0"/></w:numPr>'))

    assert status_pravila(provjeri(_izmijeni(sazetak), PROFIL, KONTEKST).nalazi, "R-NASLOVI") is Status.PROSLO


def _hrvatski_word(sadrzaj: bytes) -> bytes:
    def preimenuj(xml: bytes) -> bytes:
        return xml.replace(b'"Heading1"', b'"Naslov1"')

    return zamijeni_u_zipu(sadrzaj, {
        "word/styles.xml": preimenuj,
        "word/document.xml": preimenuj,
        "word/numbering.xml": preimenuj,
    })


def test_hrvatski_word_naslov1_s_pstyle_u_razini():
    def druga_razina_prva(d):
        # Razina povezana s Naslov1 nije prva: razina se mora naći preko pStyle.
        apstraktni = _apstraktni(d)
        lvl0, lvl1 = apstraktni.findall(qn("w:lvl"))
        lvl0.find(qn("w:pStyle")).set(qn("w:val"), "Heading2")
        lvl1.find(qn("w:pStyle")).set(qn("w:val"), "Heading1")
        lvl1.find(qn("w:lvlText")).set(qn("w:val"), "%2.")

    sadrzaj = _hrvatski_word(_izmijeni(druga_razina_prva))
    m = model.procitaj(sadrzaj, PROFIL.uloge)
    naslovi = _naslovi(m)
    assert {o.stil.kanonski_id for o in naslovi} == {"Heading1"}
    assert naslovi[0].numeracija.vrijednost == Numeracija(1, "decimal", "%2.", "1.")
    assert status_pravila(provjeri(sadrzaj, PROFIL, KONTEKST).nalazi, "R-NASLOVI") is Status.PROSLO


# --- oblikovanje -----------------------------------------------------------------


def test_font_naslova_iz_stila_je_jedan_sazeti_nalaz():
    [nalaz] = _nalazi(VARIJANTE["naslovi_font"](54, PROFIL))
    assert (nalaz.poruka_kljuc, nalaz.pronadeno.iznos, nalaz.pojave) == ("R-NASLOVI.font", "Times New Roman", 4)
    assert nalaz.izvor.vrsta is VrstaIzvora.STIL
    assert "Početna (Home) → Stilovi (Styles) → Heading 1 → Izmijeni (Modify)" in tekst_nalaza(nalaz)


def test_font_naslova_iz_teme():
    def tema(d):
        rfonts = d.styles["Heading 1"].element.rPr.rFonts
        rfonts.set(qn("w:asciiTheme"), "majorHAnsi")

    def aptos(xml: bytes) -> bytes:
        return re.sub(rb'(<a:majorFont>\s*<a:latin typeface=")[^"]*', rb"\g<1>Aptos Display", xml)

    [nalaz] = _nalazi(zamijeni_u_zipu(_izmijeni(tema), {"word/theme/theme1.xml": aptos}))
    assert nalaz.poruka_kljuc == "R-NASLOVI.font_tema" and nalaz.pronadeno.iznos == "Aptos Display"
    assert nalaz.izvor.vrsta is VrstaIzvora.TEMA


@pytest.mark.parametrize(
    ("promjena", "kljuc"),
    [
        (lambda s: setattr(s.font, "bold", False), "podebljano"),
        (lambda s: setattr(s.font, "size", 28 * 6350), "velicina"),
        (lambda s: setattr(s.paragraph_format, "alignment", 1), "poravnanje"),
        (lambda s: setattr(s.paragraph_format, "space_before", 24 * 12700), "razmak_ispred"),
        (lambda s: setattr(s.paragraph_format, "space_after", 0), "razmak_iza"),
    ],
)
def test_odstupanje_oblika_u_stilu(promjena, kljuc):
    [nalaz] = _nalazi(_izmijeni(lambda d: promjena(d.styles["Heading 1"])))
    assert (nalaz.poruka_kljuc, nalaz.pojave) == (f"R-NASLOVI.{kljuc}", 4)
    assert "{" not in tekst_nalaza(nalaz)


def test_iskljuceno_pravilo_nema_nalaza():
    naslovi = PROFIL.pravila["R-NASLOVI"].model_copy(update={"ukljuceno": False})
    iskljucen = PROFIL.model_copy(update={"pravila": {**PROFIL.pravila, "R-NASLOVI": naslovi}})
    sadrzaj = VARIJANTE["naslovi_bez_numeracije"](54, PROFIL)
    assert status_pravila(provjeri(sadrzaj, iskljucen, KONTEKST).nalazi, "R-NASLOVI") is None


def test_iskljuceno_pravilo_ne_treba_ciljeve():
    tekst = (KORIJEN / "profili" / "vjezba_1.toml").read_text(encoding="utf-8")
    pocetak = tekst.index("[pravila.R-NASLOVI]")
    kraj = tekst.find("\n[", pocetak + 1)
    bez = tekst[:pocetak] + "[pravila.R-NASLOVI]\nukljuceno = false\n" + (tekst[kraj + 1:] if kraj != -1 else "")
    assert profil.ucitaj(bez.encode("utf-8")).pravila["R-NASLOVI"].ukljuceno is False
    with pytest.raises(profil.GreskaProfila):
        profil.ucitaj(bez.replace("ukljuceno = false", "ukljuceno = true").encode("utf-8"))


def test_oznaka_bez_n_nije_dopustena():
    tekst = (KORIJEN / "profili" / "vjezba_1.toml").read_text(encoding="utf-8")
    with pytest.raises(profil.GreskaProfila):
        profil.ucitaj(tekst.replace('oznaka = "{n}."', 'oznaka = "1."').encode("utf-8"))


def test_kljucevi_r_naslovi_u_katalogu_i_redoslijed():
    for kljuc in REGISTAR["R-NASLOVI"].poruke:
        assert kljuc in KATALOG
    assert list(REGISTAR).index("R-NASLOVI") == list(REGISTAR).index("R-TEKST") + 1


def test_ima_rucni_broj():
    assert ima_rucni_broj("1. Uvod") and ima_rucni_broj("2.1 Podnaslov")
    assert not ima_rucni_broj("Uvod") and not ima_rucni_broj("Zaključak")
    assert ima_rucni_broj("1 Uvod") and ima_rucni_broj("1.Uvod") and ima_rucni_broj(" 3. Zaključak")
    assert not ima_rucni_broj("3D modeliranje") and not ima_rucni_broj("2024godina")


def test_sazetak_s_upisanim_brojem_ne_prolazi():
    def broj(d):
        o = next(o for o in d.paragraphs if o.text == "Sažetak")
        o.runs[0].text = "1. Sažetak"

    nalazi = _nalazi(_izmijeni(broj))
    assert {n.poruka_kljuc for n in nalazi} == {"R-NASLOVI.sazetak_numeriran"}
    assert nalazi[0].pronadeno.iznos == "1. Sažetak"


def test_drugi_stil_s_upisanim_brojem():
    def normal(d):
        o = next(o for o in d.paragraphs if o.text == "Zaključak")
        o.style = d.styles["Normal"]
        o.runs[0].text = "3. Zaključak"

    assert _kljucevi(_izmijeni(normal)) == {"drugi_stil", "rucni_broj"}


def test_font_teme_bez_teme_je_neprovjerljiv():
    def tema(d):
        d.styles["Heading 1"].element.rPr.rFonts.set(qn("w:asciiTheme"), "majorHAnsi")

    def bez_veze(xml: bytes) -> bytes:
        return re.sub(rb"<Relationship [^>]*/theme\"[^>]*/>", b"", xml)

    sadrzaj = zamijeni_u_zipu(_izmijeni(tema), {
        "word/theme/theme1.xml": lambda _: None,
        "word/_rels/document.xml.rels": bez_veze,
    })
    [nalaz] = _nalazi(sadrzaj)
    assert (nalaz.status, nalaz.poruka_kljuc, nalaz.pojave) == (
        Status.NEPROVJERLJIVO, "R-NASLOVI.font_nepoznat", 4,
    )


def test_dodatni_naslov_broji_se_redom_dokumenta():
    def dodatak(d):
        o = next(o for o in d.paragraphs if o.text == "Zaključak")
        o.insert_paragraph_before("Dodatak", style="Heading 1")

    sadrzaj = _izmijeni(dodatak)
    m = model.procitaj(sadrzaj, PROFIL.uloge)
    assert [o.numeracija.vrijednost.oznaka for o in _naslovi(m)] == ["1.", "2.", "3.", "4.", "5."]
    # Oznake odgovaraju poretku u dokumentu; nedostaje ništa, pa pravilo prolazi.
    assert _kljucevi(sadrzaj) == {"prolazi"}


# --- podebljanje (toggle) ------------------------------------------------------------


def _model_runova(promjena):
    m = model.procitaj(_izmijeni(promjena), PROFIL.uloge)
    return next(o for o in m.odlomci if o.tekst.startswith("Ovo je osnovni tekst poglavlja Uvod")).runovi


def _stil_znaka(d, bold):
    stil = d.styles.add_style("Naglasak", WD_STYLE_TYPE.CHARACTER)
    stil.font.bold = bold
    return stil


def test_podebljanje_je_toggle_stila_odlomka_i_znaka():
    def toggle(d):
        d.styles["Normal"].font.bold = True
        odlomak = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst poglavlja Uvod"))
        odlomak.runs[0].style = _stil_znaka(d, True)
        odlomak.add_run(" drugi")

    prvi, drugi = _model_runova(toggle)
    assert prvi.podebljano.vrijednost is False and prvi.podebljano.izvor.stil == "Naglasak"
    assert drugi.podebljano.vrijednost is True and drugi.podebljano.izvor.stil == "Normal"


def test_izravno_podebljanje_je_apsolutno():
    def izravno(d):
        d.styles["Normal"].font.bold = True
        odlomak = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst poglavlja Uvod"))
        odlomak.runs[0].style = _stil_znaka(d, True)
        odlomak.runs[0].bold = True
        odlomak.add_run(" ne").bold = False

    prvi, drugi = _model_runova(izravno)
    assert (prvi.podebljano.vrijednost, prvi.podebljano.izvor.vrsta) == (True, VrstaIzvora.IZRAVNO)
    assert (drugi.podebljano.vrijednost, drugi.podebljano.izvor.vrsta) == (False, VrstaIzvora.IZRAVNO)


def test_nepodebljano_po_zadanom():
    [run] = _model_runova(lambda d: None)
    assert (run.podebljano.vrijednost, run.podebljano.izvor.vrsta) == (False, VrstaIzvora.ZADANO)


# --- brojači numeracije ----------------------------------------------------------------


def _numbering(tijelo: str):
    return parse_xml(f'<w:numbering xmlns:w="{W}">{tijelo}</w:numbering>')


_DVIJE_RAZINE = (
    '<w:abstractNum w:abstractNumId="0">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
    '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%1.%2)"/></w:lvl>'
    '<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="none"/><w:lvlText w:val=""/></w:lvl>'
    "</w:abstractNum>"
)


def _oznake(brojaci: Brojaci, pozivi) -> list[str | None]:
    van = []
    for num_id, ilvl in pozivi:
        n = brojaci.sljedeca(num_id, ilvl, None)
        van.append(n.oznaka if n else None)
    return van


def test_reset_dublje_razine():
    brojaci = Brojaci(_numbering(_DVIJE_RAZINE + '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'))
    assert _oznake(brojaci, [("1", 0), ("1", 1), ("1", 1), ("1", 0), ("1", 1)]) == [
        "1.", "1.a)", "1.b)", "2.", "2.a)",
    ]


def test_vise_numid_na_istom_apstraktnom_dijeli_brojace():
    brojaci = Brojaci(_numbering(
        _DVIJE_RAZINE
        + '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
        + '<w:num w:numId="2"><w:abstractNumId w:val="0"/></w:num>'
    ))
    assert _oznake(brojaci, [("1", 0), ("2", 0), ("1", 0)]) == ["1.", "2.", "3."]


def test_start_override_pri_prvoj_upotrebi_numid():
    brojaci = Brojaci(_numbering(
        _DVIJE_RAZINE
        + '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
        + '<w:num w:numId="2"><w:abstractNumId w:val="0"/>'
        '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'
    ))
    assert _oznake(brojaci, [("1", 0), ("1", 0), ("2", 0), ("2", 0), ("1", 0)]) == ["1.", "2.", "1.", "2.", "3."]


def test_bez_numeracije():
    brojaci = Brojaci(_numbering(_DVIJE_RAZINE + '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'))
    assert _oznake(brojaci, [("0", 0), ("7", 0), (None, 0), ("1", 2), ("1", 5)]) == [None] * 5
    assert Brojaci(None).sljedeca("1", 0, None) is None


def test_razina_bez_ilvl_preko_pstyle_inace_nula():
    brojaci = Brojaci(_numbering(
        '<w:abstractNum w:abstractNumId="0">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Naslov2"/>'
        '<w:lvlText w:val="%1.%2."/></w:lvl></w:abstractNum>'
        '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
    ))
    assert brojaci.sljedeca("1", None, "Naslov1").oznaka == "1."
    assert brojaci.sljedeca("1", None, "Naslov2") == Numeracija(1, "decimal", "%1.%2.", "1.1.")


@pytest.mark.parametrize(
    ("broj", "oblik", "oznaka"),
    [
        (3, "decimal", "3"), (28, "lowerLetter", "bb"), (2, "upperLetter", "B"),
        (14, "lowerRoman", "xiv"), (9, "upperRoman", "IX"), (4, "ordinal", "4"), (1, "decimalZero", "01"), (12, "decimalZero", "12"),
    ],
)
def test_formatiraj_broj(broj, oblik, oznaka):
    assert formatiraj_broj(broj, oblik) == oznaka


def test_numid_nula_u_odlomku_ponistava_numeraciju_stila():
    def ponisti(d):
        uvod = next(o for o in d.paragraphs if o.text == "Uvod")
        uvod._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="0"/></w:numPr>'))

    m = model.procitaj(_izmijeni(ponisti), PROFIL.uloge)
    uvod, studij = _naslovi(m)[:2]
    assert uvod.numeracija.vrijednost is None and uvod.numeracija.izvor.vrsta is VrstaIzvora.IZRAVNO
    assert studij.numeracija.vrijednost.oznaka == "1."


# --- stvarni uzorak -------------------------------------------------------------------


@pytest.mark.skipif(not UZORAK.exists(), reason="stvarni uzorak postoji samo lokalno")
def test_stvarni_uzorak_ne_prolazi_naslove_zbog_oznake_bez_tocke():
    """Odluka korisnika (2026-10-13, ispravljena): uzorak ima numeraciju „%1“ kroz stil Naslov1."""
    izvjestaj = provjeri(UZORAK.read_bytes(), PROFIL, KontekstProvjere(UZORAK.name, 54, None, None))
    nalazi = [n for n in izvjestaj.nalazi if n.pravilo_id == "R-NASLOVI"]
    assert status_pravila(izvjestaj.nalazi, "R-NASLOVI") is Status.NE_PROLAZI
    oznake = [n for n in nalazi if n.poruka_kljuc == "R-NASLOVI.kriva_oznaka"]
    assert [(n.ocekivano.iznos, n.pronadeno.iznos) for n in oznake] == [("1.", "1"), ("2.", "2"), ("3.", "3"), ("4.", "4")]


def test_lvl_override_s_lvl_nasljeduje_osnovnu_razinu():
    brojaci = Brojaci(_numbering(
        _DVIJE_RAZINE
        + '<w:num w:numId="1"><w:abstractNumId w:val="0"/>'
        '<w:lvlOverride w:ilvl="0"><w:lvl w:ilvl="0"><w:lvlText w:val="(%1)"/></w:lvl></w:lvlOverride></w:num>'
    ))
    assert brojaci.sljedeca("1", 0, None) == Numeracija(0, "decimal", "(%1)", "(1)")


def test_num_style_link_koristi_razine_povezanog_popisa():
    brojaci = Brojaci(_numbering(
        '<w:abstractNum w:abstractNumId="0"><w:styleLink w:val="Naslovi"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
        "</w:abstractNum>"
        '<w:abstractNum w:abstractNumId="1"><w:numStyleLink w:val="Naslovi"/></w:abstractNum>'
        '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
        '<w:num w:numId="2"><w:abstractNumId w:val="1"/></w:num>'
    ))
    assert _oznake(brojaci, [("2", 0), ("2", 0), ("1", 0)]) == ["1.", "2.", "3."]


@pytest.mark.parametrize("oznaka", ["{n[0]}.", "{n.x}.", "{{n}}.", "1."])
def test_neispravna_oznaka_u_profilu(oznaka):
    tekst = (KORIJEN / "profili" / "vjezba_1.toml").read_text(encoding="utf-8")
    with pytest.raises(profil.GreskaProfila):
        profil.ucitaj(tekst.replace('oznaka = "{n}."', f'oznaka = "{oznaka}"').encode("utf-8"))


def test_naslov_bez_numeracije_medu_poglavljima_se_ne_broji():
    def ponisti(d):
        studij = next(o for o in d.paragraphs if o.text.startswith("Studij"))
        studij._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="0"/></w:numPr>'))

    sadrzaj = _izmijeni(ponisti)
    m = model.procitaj(sadrzaj, PROFIL.uloge)
    assert [o.numeracija.vrijednost and o.numeracija.vrijednost.oznaka for o in _naslovi(m)] == ["1.", None, "2.", "3."]
    nalazi = _nalazi(sadrzaj)
    assert [(n.poruka_kljuc, n.lokacija.isjecak) for n in nalazi] == [
        ("R-NASLOVI.bez_numeracije", "Studij geodezije i geomatike u Varaždinu"),
    ]


def test_rucni_broj_uz_wordovu_numeraciju():
    def broj(d):
        uvod = next(o for o in d.paragraphs if o.text == "Uvod")
        uvod.runs[0].text = "1. Uvod"

    nalazi = _nalazi(_izmijeni(broj))
    assert [n.poruka_kljuc for n in nalazi] == ["R-NASLOVI.rucni_broj_uz_numeraciju"]
    tekst = tekst_nalaza(nalazi[0])
    assert "1. Uvod" in tekst and "Word već numerira" in tekst


def test_podnaslov_heading2_na_istom_popisu():
    def podnaslov(d):
        tekst = next(o for o in d.paragraphs if o.text.startswith("Ovo je osnovni tekst poglavlja Uvod"))
        pod = tekst.insert_paragraph_before("Podnaslov", style="Heading 2")
        pod._p.get_or_add_pPr().append(parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="1"/><w:numId w:val="{NUM_ID_NASLOVA}"/></w:numPr>'
        ))

    sadrzaj = _izmijeni(podnaslov)
    m = model.procitaj(sadrzaj, PROFIL.uloge)
    assert [o.numeracija.vrijednost.oznaka for o in _naslovi(m)] == ["1.", "2.", "3.", "4."]
    [pod] = [o for o in m.odlomci if o.tekst == "Podnaslov"]
    assert pod.numeracija.vrijednost.oznaka == "1.1."
    assert _kljucevi(sadrzaj) == {"prolazi"}


def test_upute_za_numeraciju_ne_vode_na_galeriju_bez_tocke():
    for kljuc in ("rucni_broj", "bez_numeracije", "kriva_oznaka"):
        tekst = KATALOG[f"R-NASLOVI.{kljuc}"]
        assert "Definiraj novi višerazinski popis (Define New Multilevel List)" in tekst
        assert "Poveži razinu sa stilom (Link level to style)" in tekst
    assert "1., 2., 3." not in KATALOG["R-NASLOVI.kriva_oznaka"]
